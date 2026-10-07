# ruff: noqa: E501
"""V5 edits: 5.3 Commands, 5.4 Device Twin, 5.5 Offline Behavior, 5.6 Backend Event Processing: V4 P396-P479, tables T27-T29.

Addresses V4's original numbering through ctx (see common.py). Every edit: ctx.note(...)."""
from docx.oxml.ns import qn

from .common import (EditError, add_row, delete_para, delete_row, insert_after,  # noqa: F401
                     insert_lines_after, insert_table_after, replace, replace_in_cell,
                     set_cell, set_row_font_size, set_text, zebra)


# ---------------------------------------------------------------- 5.3 Commands

def _envelope(ctx):
    # T27 R3 (id): s5-10.
    set_cell(ctx.cell(27, 3, 2, "Recommended"), "Yes (APP-FR-056)")
    set_cell(ctx.cell(27, 3, 3, "Set this so the hub returns cmd_ack"),
             "Correlation ID: a string of at most 63 characters (a UUID fits). The hub returns it as data.id "
             "in the cmd_ack. The hub acks every eflostop.cmd command, but without an id the ack has no "
             "data.id and cannot be matched. Use a new id for every send.")
    ctx.note("s5-10: T27 id row: required (APP-FR-056), at most 63 characters; the hub acks every envelope command, an id is needed to match the ack")

    # Envelope rules after T27: s5-09, worded as canon C19 (legacy keyword scan), C12 (id, size), C15 (no id dedupe).
    t27 = ctx.T(27, "Field")
    label = insert_after(t27, "Envelope rules:", like=ctx.P(407, "Complete command reference"))
    insert_lines_after(label, [
        'The hub runs a C2D message as a command only when it is a JSON object whose schema is exactly '
        '"eflostop.cmd" (or the legacy "eflostop.cmd.v1") and whose cmd is a non-empty string. Schema and '
        'command names are lower case and case-sensitive. The hub ignores envelope fields that are not in '
        'this table.',
        'Any other message gets no cmd_ack. The hub then searches its whole text, in any letter case, for the '
        'old plain-text commands DECOMMISSION_VALVE, DECOMMISSION_LORA:, DECOMMISSION_BLE:, DECOMMISSION_ALL, '
        'VALVE_OPEN, VALVE_CLOSE, RULES_CONFIG:, SENSOR_META:, LEAK_RESET and OVERRIDE_CANCEL, and runs the '
        'one it finds, again with no cmd_ack. So a command with a misspelt schema and cmd "valve_open" opens '
        'the valve, and DECOMMISSION_ALL anywhere in the text (an id, a label) resets the hub to its factory '
        'state. A valid JSON object with no cmd string and none of these words is applied as a provision '
        'payload. This applies to hub firmware 2.1.3 and 2.1.4.',
        'The backend shall build every command with a JSON library as an exact eflostop.cmd envelope, and '
        'shall never send user text, or anything else, as a C2D body outside such an envelope.',
        'Send one command per C2D message, at most 8,192 bytes. A longer message is not run: the hub answers '
        'it with a cmd_ack error (5.2.7).',
        'id is a string of at most 63 characters. A longer id comes back cut and will not match.',
        'The hub runs commands one at a time, in arrival order, and does not de-duplicate them by id: a '
        'command delivered twice runs twice and is acked twice. For commands sent while the hub is offline, '
        'and the expiry the backend sets on them, see 5.5.',
    ], like=ctx.P(460, "Up to 16 event messages"))
    ctx.note("s5-09: new 'Envelope rules' after T27: exact case-sensitive schema and cmd, 8,192-byte limit, 63-character id, no cmd_ack and the legacy keyword scan for anything that is not a valid envelope (canon C19 wording), no de-duplication by id")


def _command_table(ctx):
    # Note under "Complete command reference:" (decision: error.detail texts live once, in 5.2.7).
    p407 = ctx.P(407, "Complete command reference")
    insert_after(p407,
                 'The Error Scenarios column lists when the hub answers with cmd_ack status "error". The exact '
                 'error.detail texts, and which of them to show to users as they are, are in the table in 5.2.7.',
                 like=ctx.P(397, "Send commands using this JSON envelope"))
    ctx.note("decision 'Where the cmd_ack error.detail list lives': note before T28; its Error Scenarios column names the cases and refers to the 5.2.7 table for the exact texts")

    T = 28
    ctx.T(T, "Error Scenarios")

    # R1 valve_open: s5-11 (conditions only, texts in 5.2.7), canon C4, C12, C23.
    set_cell(ctx.cell(T, 1, 2, "Hub connects BLE and sends open command"),
             'Open the valve. The hub checks the command, connects to the valve if needed and queues the open. '
             'An "ok" cmd_ack means queued, not opened: valve_state_changed and the next snapshot confirm it. '
             'If the valve is not linked, the hub holds the open and writes it when the valve links again, with '
             'no time limit; a newer valve command replaces it, and a hub restart drops it. A valve that is '
             'already open still acks "ok".')
    set_cell(ctx.cell(T, 1, 3, "Valve already open"),
             'cmd_ack error, checked in this order: no valve is set up (since 2.1.4); the valve reports RMLEAK '
             'set, or (since 2.1.4) the hub has a leak incident latched and no override window is open, also '
             'while the valve is disconnected (2.1.3 refuses only while the valve reports RMLEAK set); the '
             'valve\'s last battery reading is 10 % or less (since 2.1.4; a 2.1.3 hub acks "ok" and the valve '
             'itself stays shut); the hub\'s valve command queue is full (since 2.1.4; 2.1.3 acks "ok").\n'
             'A 2.1.3 hub with no valve set up acks "ok" and can then link to and drive any nearby eFloStop '
             'valve: offer valve controls only when the hub has a valve set up (twin reported valve_id not null).')
    ctx.note("s5-11: T28 valve_open: ok means queued; a command for an unlinked valve is held (canon C23); the refusals in check order, each new 2.1.4 case with its 2.1.3 behaviour (RMLEAK refusal on both, latched-incident refusal since 2.1.4); 'already open' and 'BLE not connected' are no longer listed as errors; the 2.1.3 no-valve hazard")

    # R2 valve_close: s5-12.
    set_cell(ctx.cell(T, 2, 2, "Close the valve."),
             'Close the valve. As for valve_open, "ok" means queued, and a close for a valve that is not linked '
             'is held until it links. A manual close does not set RMLEAK or start a leak incident.')
    set_cell(ctx.cell(T, 2, 3, "Valve already closed"),
             'cmd_ack error when no valve is set up or the hub\'s valve command queue is full (both since 2.1.4; '
             '2.1.3 always acks "ok"). A close is never refused for RMLEAK or battery. A valve that is already '
             'closed or not linked acks "ok".')
    ctx.note("s5-12: T28 valve_close: ok means queued and held for an unlinked valve; the two 2.1.4 refusals; never refused for RMLEAK or battery")

    # R3 valve_set_state: s5-13.
    set_cell(ctx.cell(T, 3, 1, '{"state":"open"}'),
             '{"state":"open"} or {"state":"closed"} (exact lower-case values)')
    set_cell(ctx.cell(T, 3, 3, "Missing or invalid state field"),
             'cmd_ack error when state is missing or not a string, or has any other value. State "open" also '
             'gets every valve_open error, and state "closed" every valve_close error.')
    ctx.note("s5-13: T28 valve_set_state: exact lower-case values; the open and close refusals apply as for valve_open and valve_close")

    # R4 leak_reset: s5-14.
    set_cell(ctx.cell(T, 4, 2, "Clear the leak incident latch and RMLEAK"),
             'Clear the leak incident latch and RMLEAK, and end any active 24-hour override window. Does NOT '
             'open the valve. The hub sends rmleak_cleared (with override_cancelled true when it ended a window) '
             'only when there was something to clear; otherwise it acks "ok" with no event. The valve then '
             'reports rmleak false (valve.rmleak in the snapshot).')
    set_cell(ctx.cell(T, 4, 3, "Refused while any leak source is still wet"),
             'cmd_ack error while any provisioned leak source (a leak sensor or the valve flood probe) is still '
             'wet, even one whose trigger-mask bit is clear and even with auto_close_enabled false. The hub also '
             'refuses when it is briefly busy; a resend after a few seconds then succeeds. No active leak '
             'incident is not an error.')
    ctx.note("s5-14: T28 leak_reset: also ends the override window (rmleak_cleared with override_cancelled true); no event when nothing to clear; wet-source refusal counts sources whose bit is clear; no active incident is not an error")

    # R5 override_cancel: s5-15 + tm-28.
    set_cell(ctx.cell(T, 5, 2, "Cancel the 24-hour override window"),
             'Cancel the 24-hour override window and re-enable auto-close at once; the hub sends '
             'auto_close_reenabled. If auto_close_enabled is true and any provisioned leak source is still wet, '
             'the hub asserts RMLEAK and closes the valve at once, whatever the trigger mask, and sends no '
             'auto_close event (4.8.1).')
    set_cell(ctx.cell(T, 5, 3, "Override not active"),
             'No error when no window is active: the hub acks "ok", sends no event and changes nothing. cmd_ack '
             'error only when the hub is briefly busy: resend.')
    ctx.note("s5-15, tm-28: T28 override_cancel: closes the valve at once for any wet source while auto-close is on, mask not checked; 'override not active' acks ok; error only when the hub is busy")

    # R6 decommission: s5-16 (canon C3.7, C5, C8, C20).
    set_cell(ctx.cell(T, 6, 1, '"target":"valve|ble_leak_sensor|all"'),
             '{"target":"valve"}, {"target":"ble_leak_sensor","sensor_id":"00:80:E1:27:B4:96"} or '
             '{"target":"all"}. Targets are not case-sensitive; "ble" and "ble_leak" are also accepted for '
             'ble_leak_sensor. {"target":"lora","sensor_id":"0x754A6237"} removes a LoRa leak sensor (not used '
             'by the MVP app).')
    set_cell(ctx.cell(T, 6, 2, "Remove a device."),
             'Remove a device and its stored name and location. Since 2.1.4 the other devices keep their '
             'health, battery and RSSI values; on 2.1.3 they show as not heard yet until each reports again. '
             'Since 2.1.4, removing the last device also resets the rules to auto_close_enabled true and '
             'trigger_mask 7 and clears the leak latch and any override window, with no event, and the twin '
             'report that follows shows provisioned false (2.1.3 keeps the last rules until it restarts and '
             'updates the twin only at the next connect).\n'
             'target "all" resets the hub to its factory state: it erases the device lists, the sensor names, '
             'the hub name, the cloud registration cache, the rules (back to true and 7), the leak latch, the '
             'override window, the snapshot interval (back to 300 s) and, since 2.1.4, its kept events (5.5), '
             'and keeps Wi-Fi. It '
             'sends one last snapshot (data.reason "decommission") and restarts, normally 3 to 5 s after the '
             'cmd_ack and at most about 30 s after it.')
    set_cell(ctx.cell(T, 6, 3, "Unknown target"),
             'cmd_ack error when the target is missing or unknown; for target "valve" when no valve is set up '
             '(since 2.1.4; 2.1.3 acks "ok"); for ble_leak_sensor when sensor_id is missing, not a valid MAC or '
             'not provisioned; for lora when that sensor is not provisioned; or when the hub could not save the '
             'change or erase its data.')
    ctx.note("s5-16: T28 decommission: exact target forms (lora documented, not used by the app); since 2.1.4 the remaining devices keep their state and the last removal resets the rules to true/7; what target all erases (kept events only since 2.1.4) and when the hub restarts (3-5 s, at most about 30 s); error cases as the hub checks them")

    # R7 rules_config: tm-29 + s5-17 (canon C3, decision 'How the app writes the trigger mask').
    set_cell(ctx.cell(T, 7, 1, '"trigger_mask":5'),
             '{"trigger_valve_flood":false}. Keys: auto_close_enabled, trigger_ble_leak, trigger_lora and '
             'trigger_valve_flood (JSON booleans) and trigger_mask (number). Send only the keys that change. The '
             'app sends one key per change, and only auto_close_enabled, trigger_ble_leak or trigger_valve_flood '
             '(4.8.1).')
    set_cell(ctx.cell(T, 7, 2, "Update auto-close rules."),
             'Change auto-close and the leak sources that trigger it. Keys left out are unchanged. The new rules '
             'are saved on the hub and apply from the next leak report; the cmd_ack is followed by a twin report '
             'and a snapshot that show them. The trigger mask is not checked when the valve reconnects, on '
             'override_cancel or when the override window ends: then any wet provisioned leak source closes the '
             'valve while auto_close_enabled is true (4.8.1).')
    set_cell(ctx.cell(T, 7, 3, "Invalid JSON payload."),
             'cmd_ack error only when the payload is missing or null, or the hub could not save the change. Not '
             'errors: a payload that changes nothing, unknown keys and values of the wrong type (ignored, acked '
             '"ok"). Confirm the result in data.rules of the next snapshot.')
    ctx.note("tm-29, s5-17: T28 rules_config: example {\"trigger_valve_flood\":false} instead of trigger_mask 5; the per-source keys; the app sends one boolean key per change; merge behaviour; where the mask is not checked; the real error cases")

    # R8 sensor_meta: s5-18.
    set_cell(ctx.cell(T, 8, 2, "Set sensor name and location."),
             'Set a sensor\'s location and label. location_code and label are optional: one left out keeps its '
             'stored value, and label "" clears the label. The hub cuts the label to 31 bytes of UTF-8, so count '
             'bytes, not characters. location_code is one of the A.1 codes (any letter case).')
    set_cell(ctx.cell(T, 8, 3, "Invalid sensor_type, unknown sensor_id."),
             'cmd_ack error when sensor_type is missing or is not ble_leak_sensor or lora ("ble" and "ble_leak" '
             'are also accepted), sensor_id is missing or empty, the hub already holds names for 32 other '
             'sensors, or the hub could not save the change. Not errors: a location_code not in A.1 (ignored; the '
             'stored code is kept) and a sensor_id the hub does not have (stored anyway). Send sensor_meta only '
             'for provisioned sensors.')
    ctx.note("s5-18: T28 sensor_meta: optional fields, label limit in bytes, unknown location ignored, unknown sensor_id stored; the real error cases")

    # R9 provision: s5-19 + tm-31 (canon C3.6, C22 without 'since 2.0.0', APP-FR-097).
    set_cell(ctx.cell(T, 9, 1, '"valve_mac":"..."'),
             '{"valve_id":"00:80:E1:27:9A:E6", "ble_leak_sensors":["00:80:E1:27:B4:96", ...], '
             '"auto_close_enabled":true, "sensor_meta":[{...}]}. Every key is optional; always include valve_id '
             'or a sensor list.')
    set_cell(ctx.cell(T, 9, 2, "Initial device provisioning"),
             'Set the hub\'s devices. Sent by the backend during hub setup and whenever devices are added. Each '
             'list sent replaces that whole list, and a key left out keeps its current value: to add one sensor, '
             'send the complete list (second provision example below). provision takes valve_id; valve_mac is '
             'still accepted but deprecated: send valve_id. A null valve_id is ignored, and a provision cannot '
             'remove the valve (use decommission). Since 2.1.4, a provision whose sensor lists leave the hub '
             'with no device resets the rules to auto_close_enabled true and trigger_mask 7 and clears the leak '
             'latch and any override window, as removing the last device does; provisioned stays true. Up to '
             '2.1.3 such a hub kept its last rules, also across restarts, because it stays provisioned. '
             'sensor_meta entries (the fields of the sensor_meta command) '
             'set names in the same command; an invalid entry is skipped.\n'
             'Top-level auto_close_enabled true turns auto-close on and sets the trigger mask to 7; false turns '
             'it off and leaves the mask unchanged. The backend sends it only in the first provision of a hub, '
             'and only when the user chose a setting (APP-FR-097). An optional rules object '
             '{"auto_close_enabled":bool, "trigger_mask":number} is applied after the top-level key and wins '
             'over it. The trigger keys of rules_config are not read here (4.8.1). A provision that holds only '
             'auto_close_enabled or rules marks a hub with no devices as provisioned.')
    set_cell(ctx.cell(T, 9, 3, "Invalid MAC format"),
             'cmd_ack error when the valve id is not a valid MAC (XX:XX:XX:XX:XX:XX; an empty string fails too, '
             'and then nothing in the payload is applied), the payload holds nothing usable (no valve_id or '
             'valve_mac string, no ble_leak_sensors or lora_sensors array, no JSON boolean auto_close_enabled '
             'and no rules object), or the hub could not save it. Not errors (acked "ok"): an invalid sensor id '
             'is skipped, and entries past 16 in a list are dropped. Since 2.1.4 a repeated sensor id, in any '
             'letter case, is ignored. A 2.1.3 hub stores it twice: the sensor appears twice in the snapshot, '
             'and the copy that is never heard shows as offline. So the backend sends each sensor id once, '
             'compared without regard to letter case. Check the device lists in the next snapshot or the counts '
             'in twin reported.')
    ctx.note("global fix INT-09/XC-03: T28 provision: a 2.1.3 hub emptied by a provision kept its rules across restarts (master provisioning_manager.c:591 marks it provisioned; :151 loads the rules at boot)")
    ctx.note("s5-19, tm-31: T28 provision: valve_id (valve_mac deprecated, no version history); each list replaces the stored list; since 2.1.4 a provision that empties the hub resets the rules and clears the latch and override (2.1.3 kept the rules until restart); setup opt-in auto_close_enabled only in the first provision (APP-FR-097); rules object wins; trigger keys not read; invalid or extra sensors skipped with ok; repeated id ignored since 2.1.4 (2.1.3 stores it twice, so the backend sends each id once); error when nothing in the payload is usable")

    # R10 set_hub_name: s5-20 with canon C16 (the app and backend do not use it).
    set_cell(ctx.cell(T, 10, 2, "Set the user-friendly hub name"),
             'Set the hub name. The app and backend do not use this command: they set the name through Device '
             'Twin desired hub_name (5.4.2). The hub applies desired.hub_name again at every connect, so a name '
             'set only by this command reverts whenever the desired value differs. The hub saves the name and '
             'sends it in gateway.name and reported hub_name (since 2.1.4 the twin report follows the cmd_ack). '
             '"" clears the name. Limit: 31 bytes of UTF-8.')
    set_cell(ctx.cell(T, 10, 3, "Missing or empty name"),
             'cmd_ack error when name is missing or not a string, or longer than 31 bytes. An empty name is not '
             'an error: it clears the name.')
    ctx.note("s5-20, canon C16: T28 set_hub_name: not used by the app or backend (desired.hub_name wins at every connect); 31 bytes; empty name clears it")

    # R11 override_enable: s5-21; canon C22 drops 'Requires gateway.fw >= 1.4.0'.
    set_cell(ctx.cell(T, 11, 2, "Requires gateway.fw >= 1.4.0"),
             'Open the valve during an active leak and start the 24-hour water-access override window (remote '
             'equivalent of the physical valve button). An "ok" cmd_ack means the window started and the open was '
             'queued: water_access_override_enabled follows, and valve_state_changed confirms that the valve '
             'opened. The hub first tries to reach the valve for up to 10 s, so this cmd_ack can take about 10 s, '
             'and commands sent meanwhile wait.')
    set_cell(ctx.cell(T, 11, 3, "No active leak"),
             'cmd_ack error when no valve is set up; the valve does not link within about 10 s; there is no '
             'active leak to override (no leak incident latched, no valve RMLEAK and no window open); the valve '
             'flood probe is wet; or the hub hits an internal error. The hub does not check the valve battery: '
             'at 10 % or less it starts the window and the valve stays shut (APP-FR-043).')
    ctx.note("global fix completeness-08: T28 override_enable: no battery check (rules_engine.c:1440-1515), so at 10 % or less the window starts and the valve stays shut")
    # V4's override_enable row carries no run size (11 pt in a 9 pt table).
    set_row_font_size(ctx.row(T, 11, "override_enable"), 9)
    ctx.note("final check: T28 override_enable row set to the table's 9 pt (V4's R11 carried no size)")
    ctx.note("s5-21, canon C22: T28 override_enable: what ok means, the 10 s valve wait, the error cases; firmware 1.4.0 requirement dropped (every hub this document covers supports it)")

    # After T28: s5-23 (canon C9, C11, C12, C23).
    t28 = ctx.T(T, "Error Scenarios")
    normal = ctx.P(397, "Send commands using this JSON envelope")
    insert_lines_after(t28, [
        'Every eflostop.cmd command gets a cmd_ack (5.2.7). error.code is always the command name, not an '
        'error code: branch on status, and match error.detail exactly against the table in 5.2.7. A cmd not in '
        'this table, or a misspelt one, gets status "error" with error.detail "unknown command".',
        'What "ok" means: for valve_open, valve_close, valve_set_state and override_enable, the hub accepted and '
        'queued the command; valve_state_changed and the snapshot confirm what the valve did, and a snapshot '
        'sent just after the ack can still show the old state. For provision, the configuration was saved; the '
        'snapshots show when each device is first heard. For the other commands, the hub applied the change, '
        'and the twin report or the snapshot shows the result.',
        'What follows the cmd_ack: after every command that succeeds the hub sends a snapshot, normally within '
        '5 s. After a provision that adds devices, more snapshots follow for about 5 minutes while the new '
        'devices are first heard. provision, decommission, rules_config and set_hub_name also produce a twin '
        'report. The cmd_ack comes before the snapshot the command causes (other events can come between them) '
        'and, since 2.1.4, before the twin report it causes. 2.1.3 sent the twin report first for provision, '
        'rules_config and set_hub_name, and sent none after a single-device decommission. Do not depend on the '
        'order of the twin report and the snapshot: the twin report waits while 8 or more messages are queued '
        'on the hub.',
    ], like=normal)
    ctx.note("s5-23: new text after T28: every envelope command is acked, error.code is the command name, unknown command; what ok means per command; the snapshot and twin report that follow and their order since 2.1.4 and on 2.1.3")


def _examples(ctx):
    # P412: canon C2 terminology ("BLE sensor" alone).
    replace(ctx.P(412, "decommission (single BLE sensor)"), "single BLE sensor", "single BLE leak sensor")
    ctx.note("canon C2: P412 example label says 'BLE leak sensor'")

    # P416-P417 rules_config example: s5-v01, tm-30.
    set_text(ctx.P(416, "rules_config:"), "rules_config (one key per change; 4.8.1):")
    set_text(ctx.P(417, '"trigger_mask":5'),
             '{"schema":"eflostop.cmd","ver":1,"id":"rc-004","cmd":"rules_config","payload":{"trigger_valve_flood":false}}')
    ctx.note("s5-v01, tm-30: rules_config example sends {\"trigger_valve_flood\":false} instead of trigger_mask 5, which turned LoRa auto-close off")

    # P420 set_hub_name example label: canon C16.
    set_text(ctx.P(420, "set_hub_name:"),
             "set_hub_name (the hub supports it, but the app and backend set the hub name through Device Twin "
             "desired hub_name, 5.4.2):")
    ctx.note("canon C16: set_hub_name example labelled as not used by the app or backend")

    # Provision example: s5-22 (valve_id), s5-v02 + tm-32 (setup opt-in instead of rules trigger_mask 5).
    set_text(ctx.P(422, "provision (initial setup)"), "provision (first setup of a hub):")
    replace(ctx.P(429, '"valve_mac": "00:80:E1:27:9A:E6"'), '"valve_mac"', '"valve_id"')
    set_text(ctx.P(431, '"rules": {'), '    "auto_close_enabled": true')
    delete_para(ctx.P(432, '"auto_close_enabled": true,'))
    delete_para(ctx.P(433, '"trigger_mask": 5'))
    delete_para(ctx.P(434, "    }"))
    ctx.note("s5-22, s5-v02, tm-32: first provision example uses valve_id and the top-level auto_close_enabled opt-in instead of valve_mac and rules.trigger_mask 5")

    # After P436: tm-32 note + s5-v03 second example.
    p436 = ctx.P(436, "}")
    a = insert_after(p436,
                     "Send auto_close_enabled only in the first provision of a hub, and only when the user chose a "
                     "setting (APP-FR-097). Leave it out when adding or removing devices later: true sets the "
                     "trigger mask back to 7 (4.8.1).",
                     like=ctx.P(397, "Send commands using this JSON envelope"))
    b = insert_after(a, "provision (adding a sensor after setup: send the complete sensor list, and no "
                        "auto_close_enabled or rules):", like=ctx.P(422, "provision"))
    c = ctx.code_after(b, [
        '{"schema":"eflostop.cmd","ver":1,"id":"prov-008","cmd":"provision","payload":{"ble_leak_sensors":'
        '["00:80:E1:27:B4:96","00:80:E1:27:B6:A5","00:80:E1:27:C1:02"],"sensor_meta":[{"sensor_type":'
        '"ble_leak_sensor","sensor_id":"00:80:E1:27:C1:02","location_code":"kitchen","label":"Under sink"}]}}'])
    insert_after(c, "valve_id is left out, so the valve stays as it is, and the rules stay as the user set them. "
                    "A sensor left out of ble_leak_sensors would be removed.",
                 like=ctx.P(397, "Send commands using this JSON envelope"))
    ctx.note("tm-32, s5-v03: after the provision example, the first-provision-only rule for auto_close_enabled and a second example that adds a sensor with the complete list")


# ---------------------------------------------------------------- 5.4 Device Twin

def _twin(ctx):
    normal = ctx.P(397, "Send commands using this JSON envelope")

    # P439: s5-02 + tm-33 (the 'when' part goes after the field table).
    set_text(ctx.P(439, "The hub publishes these reported properties"),
             "The hub writes these reported properties as a twin patch, not in the eflostop.v2 envelope, so "
             "they do not arrive on the telemetry stream (5.6). Example from hub firmware 2.1.4:")

    # P440-P452 example: s5-01 + tm-34.
    set_text(ctx.P(441, '"fw_version": "1.3.0"'), '  "fw_version": "2.1.4",')
    p446 = set_text(ctx.P(446, '"valve_device_id"'), '  "valve_mac": null,')
    ctx.code_after(p446, ['  "valve_device_id": null,',
                          '  "valve_id": "00:80:E1:27:9A:E6",',
                          '  "lora_sensor_count": 0,'])
    set_text(ctx.P(449, '"trigger_mask": 5'), '  "trigger_mask": 7,')
    ctx.code_after(ctx.P(450, '"uptime_s": 971'), ['  "snapshot_interval_s": 300,'])
    set_text(ctx.P(451, '"free_heap": 142320'), '  "free_heap": 38912')
    ctx.note("s5-01, tm-34: twin reported example as 2.1.4 writes it: fw 2.1.4, valve_mac and valve_device_id null, valve_id, lora_sensor_count, trigger_mask 7, snapshot_interval_s, realistic free_heap")

    # Field table after the example: s5-03 (canon C3.9, C8, C20). Template: T26's 3-column layout (read only).
    if ctx.orig_cell(26, 0, 0) != "Field" or ctx.orig_cell(26, 0, 2) != "Description":
        raise EditError("T26 is not the 3-column Field | Type | Description table expected as the template")
    tbl = insert_table_after(ctx.P(452, "}"), ctx.tables[26], ["Property", "Type", "Description"], [
        ["fw_version", "string", "Hub firmware version, the same as gateway.fw."],
        ["gateway_id", "string", "Gateway ID, the same as gateway.id."],
        ["short_id", "string", "Last four hex digits of the Gateway ID. Always present."],
        ["hub_name", "string", "Hub name (5.4.2). Empty string when unset."],
        ["provisioned", "boolean",
         "false before the first provision and after a decommission removes the last device. It can read true "
         "with no devices, after a provision that set only rules or emptied the sensor lists: decide \"no "
         "devices\" from valve_id and the two counts."],
        ["valve_id", "string or null",
         "The provisioned valve's BLE MAC, upper case with colons. null when no valve is set up. Read this key "
         "for the valve."],
        ["valve_mac, valve_device_id", "null",
         "Old names of valve_id. The hub writes them as null in every report, which removes any old value from "
         "the twin. Ignore them."],
        ["lora_sensor_count", "integer", "Provisioned LoRa leak sensors (0-16)."],
        ["ble_leak_sensor_count", "integer", "Provisioned BLE leak sensors (0-16)."],
        ["auto_close_enabled", "boolean",
         "Auto-close master switch (4.8.1). Left out of a report when the hub cannot read its rules within 1 s; "
         "the twin then keeps the previous value."],
        ["trigger_mask", "integer",
         "Auto-close trigger bits, default 7 (4.8.1). Left out in the same case as auto_close_enabled."],
        ["uptime_s", "integer", "Hub uptime in seconds when the report was built. Diagnostic only."],
        ["snapshot_interval_s", "integer",
         "Heartbeat snapshot interval in force, in seconds (60-3600, default 300). Compare it with "
         "desired.snapshot_interval_s to confirm a write (5.4.2)."],
        ["free_heap", "integer", "Free hub memory in bytes. Diagnostic only."],
    ])
    # Banding as in T26-T28: insert_table_after copies T26's first body row (shaded) for every
    # row, so the even body rows lose their cell shading here, leaving them plain like V4's.
    for i, row in enumerate(tbl.rows):
        if i >= 2 and i % 2 == 0:
            for tc in row._tr.findall(qn("w:tc")):
                tcPr = tc.tcPr
                shd = tcPr.find(qn("w:shd")) if tcPr is not None else None
                if shd is not None:
                    tcPr.remove(shd)
    ctx.note("s5-03: new field table for the twin reported properties after the example, banded like T26")

    # P453: s5-04 (its first clause now lives in the table rows hub_name, short_id, valve_id).
    set_text(ctx.P(453, "hub_name is the empty string when unset"),
             "If the hub's device list is busy for more than 1 s while a report is built, that one report can "
             "read provisioned false, valve_id null and both counts 0; the next report corrects it. Do not remove "
             "devices from the backend registry on a twin report alone: use the decommission cmd_ack and the "
             "snapshot.")
    ctx.note("s5-04: P453 warns that one twin report can transiently read empty; the unset and null cases moved into the field table")

    # When the hub writes reported: s5-02 + tm-33 (canon C3.9, C8, C11).
    w = insert_after(ctx.P(453, "hub_name is the empty string when unset"),
                     "The hub writes reported properties:", like=normal)
    w = insert_lines_after(w, [
        "at every MQTT connect, after that connection's lifecycle message, and once more after it has read the "
        "desired properties (5.4.2);",
        "after every provision and every decommission;",
        "after rules_config and set_hub_name;",
        "after every desired-property change.",
    ], like=ctx.P(460, "Up to 16 event messages"))
    insert_after(w,
                 "It writes them only while connected; a change made while the hub is offline is reported at the "
                 "next connect. Since 2.1.4 a command's cmd_ack always goes out before the twin report it causes. "
                 "The report normally follows within a second, but can take several seconds while the hub has many "
                 "messages waiting to be sent. So do not read reported once when the cmd_ack arrives: wait until the "
                 "property equals the value the command set (compare with the value sent, not with the value read "
                 "when the ack arrives), or use the snapshot that follows the command. On a slow or reconnecting link "
                 "an older report can, rarely, land after a newer one; the hub's next report corrects it. Up to "
                 "2.1.3, provision, rules_config and set_hub_name wrote the twin before their cmd_ack, a "
                 "single-device decommission did not update the twin until the next connect, and a hub with no "
                 "devices reported at every connect without a lifecycle message.",
                 like=normal)
    ctx.note("s5-02, tm-33: when the hub writes twin reported (every connect, provision, decommission, rules_config, set_hub_name, desired change); cmd_ack first since 2.1.4; wait for the value to change; the 2.1.3 differences")

    # 5.4.2: s5-06, s5-05, s5-07 (canon C16), s5-08, tm-35.
    set_text(ctx.P(455, "the only desired property the hub handles"),
             "The hub handles these two desired properties and ignores any other:")
    ctx.note("s5-06: P455 says the hub handles two desired properties")

    ctx.T(29, "Property")
    set_cell(ctx.cell(29, 1, 4, "does not persist it across reboot"),
             'How often the hub sends a heartbeat snapshot. It takes effect at once. The hub saves it in flash and '
             'also reads the whole twin at every connect, so the value survives a restart, and a change made while '
             'the hub was offline is applied when it reconnects. The backend does not need to re-apply it. '
             'reported.snapshot_interval_s shows the interval in force. decommission target "all" resets it to '
             '300 s on the hub, but the hub applies desired.snapshot_interval_s again when it reconnects: set it '
             'to null when a hub is reset or changes owner.')
    ctx.note("s5-05: T29 snapshot_interval_s is saved on the hub and re-read at every connect; no re-apply after lifecycle; reported echo; reset and owner change")

    set_cell(ctx.cell(29, 2, 2, "up to 31 chars"), "up to 31 bytes (UTF-8)")
    set_cell(ctx.cell(29, 2, 4, "May also be set via the set_hub_name C2D command"),
             'Hub display name. The backend sets the hub name only through this property ("" clears it). The hub '
             'applies it at once and again at every connect, saves it, reports it in reported.hub_name, and sends '
             'it as gateway.name in every message (left out while the name is empty). A value over 31 bytes is '
             'ignored and reported.hub_name keeps the old name; null removes the desired value and the hub keeps '
             'its current name. The app and backend do not use the set_hub_name command: the hub applies '
             'desired.hub_name again at every connect, so a name set only by command reverts whenever the desired '
             'value differs. decommission target "all" clears the name on the hub; when a hub is reset or changes '
             'owner, the backend sets desired.hub_name and desired.snapshot_interval_s to null, or the hub applies '
             'them again when it reconnects.')
    ctx.note("s5-07, canon C16: T29 hub_name: 31 bytes; the only way the backend sets the name; applied again at every connect; set_hub_name not used; null both desired properties at reset or owner change")

    p456 = set_text(ctx.P(456, "Values outside the 60-3600 range"),
                    "A snapshot_interval_s outside 60-3600, or a hub_name over 31 bytes, is rejected: the hub keeps "
                    "its current value and still writes reported. So desired.snapshot_interval_s 30 with "
                    "reported.snapshot_interval_s 300 means the write was rejected. Desired properties get no "
                    "cmd_ack; the reported value is the only confirmation.")
    ctx.note("s5-08: P456 says how the backend detects a rejected desired value (the reported echo)")

    insert_after(p456, "auto_close_enabled and trigger_mask are reported properties only. The hub ignores them in "
                       "desired properties. Change them with rules_config (4.8.1).")
    ctx.note("tm-35: after P456, auto_close_enabled and trigger_mask cannot be set through desired properties")
    return tbl


# ---------------------------------------------------------------- 5.5 Offline Behavior

def _offline(ctx):
    bullet = ctx.P(460, "Up to 16 event messages")
    label_like = ctx.P(459, "When MQTT is offline:")

    # P458: s5-24 with canon C21 (no BLE-scanning statements for Wi-Fi setup or router retries).
    p458 = set_text(ctx.P(458, "The hub keeps operating when the internet is down"),
                    "The hub keeps operating when the router or the internet is down. While Wi-Fi is still saved, "
                    "it keeps watching its leak sensors and its valve on its own, closes the valve on a leak when "
                    "auto-close is on, and keeps up to 16 events to send when it reconnects. The rules engine and "
                    "health engine do not depend on the cloud connection. Since 2.1.4 this also holds when the hub "
                    "restarts while the router is down. A 2.1.3 hub that restarts while Wi-Fi is down does not act "
                    "on leaks until it is back on Wi-Fi and its cloud start-up has run.")
    insert_after(p458, "While the hub's Wi-Fi setup network is open, for first setup or after the 10-second Wi-Fi "
                       "reset, the hub is offline to the cloud: it sends nothing and runs no commands. It reaches the "
                       "cloud only after it has joined the router and, since 2.1.4, after its setup network has "
                       "closed, 15 to 60 seconds after it joins (at most 75 s).")
    ctx.note("s5-24, canon C21: P458 local protection while the router or internet is down, since 2.1.4 also after a restart, and the 2.1.3 behaviour; new paragraph: the hub is offline to the cloud while its Wi-Fi setup network is open. The finding's BLE-listening exceptions are not carried (WP8 changes them within 2.1.4)")

    # P460-P464: s5-25, s5-26, s5-27, s5-28 (canon C13, C11).
    set_text(ctx.P(460, "Up to 16 event messages"),
             'Up to 16 events are kept in the hub\'s flash. They survive a hub restart. Since 2.1.4 decommission '
             'target "all" clears them. A 2.1.3 hub keeps them through the reset and sends them once it is '
             'provisioned again, each with its ts from before the reset. The backend keeps any event whose ts is '
             'older than the cmd_ack of that decommission out of the new setup\'s history.')
    set_text(ctx.P(461, "Each buffered event is capped to 512 bytes"),
             "An event longer than 512 bytes is not kept (since 2.1.4); it is lost if it cannot be sent at once. "
             "A 2.1.3 hub cuts such an event to 512 bytes and later sends it as invalid JSON: the backend logs a "
             "message it cannot parse and moves on, without failing.")
    set_text(ctx.P(462, "Only event messages are buffered"),
             "Only events are kept (leak, valve, rules and health events, and cmd_ack). Snapshots, lifecycle "
             "messages and twin reports are not kept: at every connect the hub sends a new lifecycle message and "
             "twin report, and a snapshot soon after.")
    p464 = set_text(ctx.P(464, "On reconnect, buffered events are drained"),
                    "On reconnect, kept events are sent oldest first, each with its original ts, before that "
                    "connection's lifecycle message; new events follow it. A cmd_ack for a command that arrives "
                    "during this replay is sent at once and can come before the lifecycle and before older kept "
                    "events. Since 2.1.4 an event the hub's MQTT client refuses while connected is also kept and "
                    "retried every 10 s, in order (2.1.3 dropped it).")
    p464b = insert_after(p464,
                         "Since 2.1.4 the hub also keeps events raised before its clock is first set after a power cut "
                         "(for example when the router is down at boot), and sends them after the first connect with "
                         "ts corrected from its uptime; one left over from a restart before the sync is dropped. "
                         "2.1.3 discarded every event raised before the clock sync.", like=bullet)
    ctx.note("s5-25, s5-26, s5-27, s5-28: offline buffer bullets: kept in flash across restarts; cleared by decommission all since 2.1.4 (2.1.3 keeps them and sends them after the next provision with their old ts); over 512 bytes not kept since 2.1.4 (2.1.3 sent invalid JSON); cmd_ack kept, twin reports not; replay order and the cmd_ack exception; an event the MQTT client refuses is kept since 2.1.4 (2.1.3 dropped it); pre-sync events kept since 2.1.4 (canon C13)")

    # Delivery, duplicates and order: s5-29 (canon C10, C11; APP-NF-013, APP-NF-014).
    # The label lands between the new bullet and P465, so P465 becomes the first bullet under it.
    insert_after(p464b, "Delivery, duplicates and order (APP-NF-013, APP-NF-014):", like=label_like)
    p465 = ctx.P(465, "out-of-order and potentially duplicate events")
    # Integrator (s3 review R15, s3 cross-module): the dedupe and order rules are stated once, in 5.2
    # ("Delivery and duplicates", "Dedupe rule (APP-NF-013)", "Order (APP-NF-014)"). 5.5 points there
    # and keeps only what 5.2 does not say: how an event can still be lost.
    set_text(p465,
             "Delivery is at least once: the same message (event, cmd_ack, snapshot or lifecycle) can arrive "
             "twice, and a copy or a kept event can arrive after newer messages (more often since 2.1.4). The "
             "rules for this are stated once, in section 5.2: \"Delivery and duplicates\", \"Dedupe rule "
             "(APP-NF-013)\" (drop only byte-identical copies, keep the first) and \"Order (APP-NF-014)\" (order "
             "by ts, never by arrival).")
    last = insert_after(p465,
        "At-least-once delivery does not mean that no event is ever lost. An event can still be lost when more than 16 events wait while the hub "
        "cannot send (the oldest is overwritten), when it is longer than 512 bytes and cannot be sent at once, "
        "when the hub already has 24 messages waiting during a long stall, when the connection dies silently "
        "just after the hub handed the event to it, or when the hub restarts again before its clock is first "
        "set. After an outage the first snapshot is the source of truth.", like=bullet)
    ctx.note("s5-29: P465 becomes a 'Delivery, duplicates and order' list: at least once (more often since 2.1.4) and how an event can still be lost")
    ctx.note("integrator (s3 review R15): the 5.5 copy of the dedupe and order rules is cut to a pointer to the 5.2 paragraphs 'Delivery and duplicates', 'Dedupe rule (APP-NF-013)' and 'Order (APP-NF-014)', so the rule is stated once; the list of how an event can still be lost stays here")

    # Commands sent while the hub is offline: s5-30 with canon C15.
    lab2 = insert_after(last, "Commands sent while the hub is offline:", like=label_like)
    insert_lines_after(lab2, [
        "The hub does not check a command's age and does not de-duplicate commands by id: it runs every C2D "
        "message it receives, in arrival order, when it receives it, and sends its cmd_ack then. A command still "
        "queued in IoT Hub runs when the hub reconnects. The app blocks commands while the hub is shown offline "
        "(APP-NF-011), but a command sent just before the hub dropped can still run late.",
        "So the backend shall set a 30-second expiry (the message's absolute expiry time, ExpiryTimeUtc in the "
        "IoT Hub service SDK) on valve_open, valve_close, valve_set_state and override_enable. 30 s is the app's "
        "20 s ack timeout (APP-FR-053) plus 10 s for the backend and IoT Hub: a command the app is still waiting "
        "for does not expire first, and one it reported as timed out can still run for at most about 10 s after "
        "that. Every other command (leak_reset, override_cancel, rules_config, sensor_meta, provision, "
        "decommission, set_hub_name) keeps the IoT Hub default time-to-live (1 hour unless the IoT Hub's "
        "cloud-to-device settings change it). Every send, retries included, uses a new correlation id.",
    ], like=bullet)
    ctx.note("s5-30, canon C15: new 'Commands sent while the hub is offline': no age check or id de-duplication, queued commands run at reconnect; the backend sets a 30 s expiry on the four valve-actuating commands only")


# ---------------------------------------------------------------- 5.6 Backend Event Processing

def _backend(ctx):
    set_text(ctx.P(469, 'check the "type" field'),
             'On each message, parse the body as UTF-8 JSON and check its "type" field: snapshot, lifecycle or '
             'event (for an event, data.event names it). The hub sets no message properties (no content type, no '
             'content encoding, no application properties), so IoT Hub message routing cannot filter on the body '
             'or on type: route all device messages to one endpoint and branch in the consumer.')
    ctx.note("s5-v04: P469 parse the body; the hub sets no message properties, so IoT Hub routing cannot filter on the body")

    set_text(ctx.P(470, "For snapshots: upsert into a single current-state record"),
             "For snapshots: upsert into a single current-state record per gateway (overwrite in place; do not "
             "retain snapshot history), and skip a snapshot whose ts is older than the stored one. data.reason "
             "says why the hub sent it (heartbeat, event, boot, fast, commission or decommission; 5.2.2); it is "
             "not state. A decommission snapshot is the last one before a factory-reset restart. The app reads "
             "this record on load.")
    ctx.note("s5-31: P470 skip an older snapshot (ts guard); data.reason is not state; the decommission snapshot")

    set_text(ctx.P(471, "For lifecycle: log it."),
             "For lifecycle: log it, and update the device registry with provisioned, valve_id (absent when there "
             "is no valve), ble_leak_sensor_count, lora_sensor_count and rules. To decide that a hub has no "
             "devices, use valve_id and the two counts, not provisioned. Lifecycle comes at every MQTT connect, "
             "not only at boot: tell a restart from a reconnect with ts minus gateway.uptime_s (section 5.2.1); "
             "data.reset_reason gives the cause of the last restart and repeats unchanged on every reconnect. Since "
             "2.1.4 a hub with no devices also sends lifecycle (provisioned normally false), twin reports and "
             "snapshots. A 2.1.3 hub with no devices sends "
             "no lifecycle and no snapshots (only cmd_acks and twin reports) until it is provisioned, and sends "
             "its lifecycle after the first provision. So do not wait for a lifecycle to confirm a decommission: "
             "use its cmd_ack.")
    ctx.note("global fix INT-08, completeness-09: 5.6 lifecycle bullet tells a restart from a reconnect with ts minus uptime_s (5.2.1), not a small uptime_s; provisioned is normally false on an empty hub (a rules-only or emptying provision leaves it true)")
    ctx.note("s5-32: P471 lifecycle fields for the registry, sent at every connect, restart detection; empty hub since 2.1.4 and on 2.1.3; confirm a decommission from its cmd_ack")

    set_text(ctx.P(472, "For events: append to event history store"),
             "For events: drop byte-identical duplicates (APP-NF-013), append the rest to the event history store "
             "ordered by ts (APP-NF-014), and run the push notification rules. Ignore device_offline and "
             "device_recovered for a device the hub no longer has: one held on the hub can arrive after the "
             "cmd_ack of the decommission that removed the device (APP-NF-017).")
    ctx.note("s5-33: P472 dedupe, order by ts, ignore health events for removed devices (APP-NF-013, APP-NF-014, APP-NF-017)")

    set_text(ctx.P(473, "match the id field to a pending command"),
             'For cmd_ack events: match data.id (and data.cmd) to the pending command and resolve it with status '
             '"ok" or "error". "ok" on valve_open, valve_close, valve_set_state and override_enable means accepted '
             'and queued, not done: the valve\'s state comes from valve_state_changed and the snapshot. A missing '
             'ack means unknown, not failed: an ack can be late (kept on the hub while it was offline) or lost. '
             'Record a late ack even after the app\'s 20 s timeout (APP-NF-018). The first ack for an id resolves '
             'it.')
    ctx.note("s5-34: P473 ok on valve commands means queued; a missing ack is unknown; record late acks (APP-NF-018); the first ack for an id resolves it")

    insert_after(ctx.P(475, "Send C2D commands via the Azure IoT Hub service SDK"),
                 "Twin reported changes do not arrive on the telemetry stream. Read them from the device twin with "
                 "the service SDK, or route twin change notifications. After a setting change, wait until the "
                 "reported value equals the value sent instead of reading it once when the cmd_ack arrives (5.4.1).")
    ctx.note("s5-35: new bullet after P475: read twin reported from the twin, not the telemetry stream")
    ctx.note("global fix XC-16: 5.4.1 and 5.6 compare twin reported with the value sent: a 2.1.3 hub writes the twin before the cmd_ack (master app_iothub.c:1015, :1072, :1093), so a change from the value read at the ack never comes")

    set_text(ctx.P(478, "Snapshots are the single source of truth"),
             "On app startup, fetch the latest stored snapshot from the backend and show it right away. Then "
             "subscribe to live updates. When an event arrives, apply it to the UI state, then let the next "
             "snapshot confirm it: the hub sends one shortly after every leak, valve, rules or health event and "
             "after every command that succeeds (normally within 5 s), and a heartbeat at the snapshot interval. "
             "Snapshots are the single source of truth.")
    ctx.note("s5-36: P478 says when the confirming snapshot comes")


def apply(ctx):
    _envelope(ctx)
    _command_table(ctx)
    _examples(ctx)
    _twin(ctx)
    _offline(ctx)
    _backend(ctx)
    zebra(ctx.T(28, "Error Scenarios"), ctx.body_shd)
    zebra(ctx.T(29, "Property"), ctx.body_shd)
    ctx.note("global fix INT-06: T28 (commands) and T29 (desired properties) re-banded; V4's own banding was broken there")
