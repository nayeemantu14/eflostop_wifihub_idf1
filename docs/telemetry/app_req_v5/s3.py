# ruff: noqa: E501
"""V5 edits: 5.1 MQTT topics, 5.2 envelope, 5.2.1 lifecycle, 5.2.2 snapshot: V4 P179-P268, tables T20-T23.

Addresses V4's original numbering through ctx (see common.py). Every edit: ctx.note(...)."""
import copy

from docx.oxml.ns import qn
from docx.shared import Pt

from .common import (EditError, add_row, delete_para, delete_row, insert_after,  # noqa: F401
                     insert_lines_after, insert_table_after, replace, replace_in_cell,
                     set_cell, set_text, unique_cells)


# ---------------------------------------------------------------- texts

P180_TEXT = (
    "This section is the ground truth for message formats. Every field name, value and JSON example "
    "below was checked against hub firmware 2.1.4, the release the app ships against. Hubs in the "
    "field run 2.1.3 until they are updated. Where 2.1.4 behaves differently in a way the cloud can "
    "see, the text says \"since 2.1.4\", gives the 2.1.3 behavior, and says what the app and backend "
    "do for both. The running version is gateway.fw in every message, and fw_version in twin "
    "reported. If you spot a difference "
    "between this document and the hub behavior, the hub behavior wins and this document should be "
    "updated."
)

T20_D2C = (
    "All telemetry from hub to cloud: lifecycle, snapshots and events (cmd_ack included), schema "
    "eflostop.v2, published at QoS 1. {device_id} is the hub's IoT Hub device ID, which is its "
    "Gateway ID (for example GW-34B7DA6AAD54): the hub registers with DPS under it. The hub sets no "
    "message properties (no content type or content encoding), so IoT Hub routing queries on body "
    "fields do not match these messages: route the whole stream and parse the JSON from the message "
    "body."
)
T20_C2D = (
    "Commands from cloud to hub (eflostop.cmd, section 5.3). The hub subscribes at QoS 1. It does not "
    "run a C2D message larger than 8192 bytes; it answers with cmd_ack status \"error\" (error.detail "
    "in section 5.2.7)."
)
T20_REPORTED = (
    "Hub reports its identity, firmware version, provisioning state, device counts and settings "
    "(section 5.4.1). Sent after every MQTT connect (also by a hub with no devices), after every "
    "provision, after a decommission that removes a single device (since 2.1.4; 2.1.3 sent none "
    "then), after rules_config and set_hub_name, and after every desired-property change."
)
T20_DESIRED = (
    "Backend pushes settings: snapshot_interval_s and hub_name (section 5.4.2). The hub applies a "
    "change when it arrives, keeps it across restarts, and applies the desired values again at every "
    "connect."
)
T20_GET = (
    "The hub reads its full twin once per MQTT connect and applies the desired properties in it. A "
    "desired value changed while the hub was offline is applied at the next connect; the backend "
    "does not need to send it again."
)
T20_RES = "IoT Hub's answers to the twin GET and to each reported PATCH."

P183_TEXT = (
    "All D2C telemetry messages use this top-level envelope. The hub writes the keys in the order "
    "shown; parse them by name, not by position. gateway.name is present only when a hub name is set. "
    "The event examples below leave out gateway.name to save space."
)

T21_TS = (
    "Unix epoch seconds (UTC) at which the hub built the message, never below 1704067200 "
    "(2024-01-01). A message sent again (from the hub's offline buffer, or by MQTT's own resend) "
    "keeps its original ts, so order events by ts, never by arrival or IoT Hub enqueued time. "
    "Snapshots and lifecycle messages are sent only after the hub's clock is set. Since 2.1.4 an "
    "event raised before the hub's clock is first set after a power cut is kept and sent after the "
    "first connect, with ts corrected from the hub's uptime; 2.1.3 discarded such events."
)
T21_ID = (
    "Gateway ID: \"GW-\" followed by 12 uppercase hex digits, for example \"GW-34B7DA6AAD54\". Derived "
    "from the hub's Wi-Fi MAC; it never changes."
)
T21_FW = (
    "Hub firmware version, MAJOR.MINOR.PATCH, for example \"2.1.4\". Compare versions numerically, "
    "part by part, not as strings."
)
T21_UPTIME = (
    "Whole seconds since the hub last started. It starts again from 0 at every restart. ts minus "
    "uptime_s is when the hub started. A lifecycle follows a restart when that start time is at or "
    "after the previous lifecycle's ts (allow 1 s for rounding): the hub restarted after sending "
    "that lifecycle. After a reconnect without a restart, the start time equals the previous "
    "lifecycle's within a few seconds (section 5.2.1). Do not compare uptime_s values alone: a "
    "restart soon after an earlier restart can show a larger one."
)
T21_NAME = (
    "Present only when a hub name is set. The user-assigned hub name, up to 31 bytes of UTF-8 (31 "
    "characters of plain ASCII, fewer for accented or non-Latin text). The backend sets it through "
    "the Device Twin desired property hub_name (section 5.4.2)."
)
T21_TYPE = (
    "\"lifecycle\", \"snapshot\" or \"event\". Every event, cmd_ack included, has type \"event\"; "
    "data.event names it (for example \"leak_detected\" or \"cmd_ack\")."
)

P196_OLD = "that is not described in this section."
P196_NEW = (
    "that is not described in this section, and any value of type, data.event, data.reason or "
    "data.reset_reason that they do not know. They shall accept null wherever this section allows "
    "it, and an empty data.valve object ({})."
)

DELIVERY = [
    "**Delivery and duplicates.** The hub delivers every D2C message (events, cmd_ack, snapshots, "
    "lifecycle) at least once: it publishes at QoS 1. Since 2.1.4 an event whose publish meets a "
    "dropping connection is kept and sent again after the reconnect, and the MQTT client can deliver "
    "its own copy too, so the same message can arrive twice, and the second copy can arrive after "
    "newer messages. The copies are byte-identical, ts included. A 2.1.3 hub can also deliver a "
    "message twice, but loses an event whose publish meets a dropping connection.",

    "**Dedupe rule (APP-NF-013).** The backend drops a message only when its body is byte-identical "
    "to one already received (compare the raw body, or its SHA-256) and keeps the first copy. It "
    "never de-duplicates on a subset of fields such as gateway.id + ts + event + device id: ts has "
    "one-second resolution, so such a key merges different messages (two cmd_acks in one second, two "
    "valve_state_changed for one valve in one second, two auto_close for one sensor in one second). "
    "History rows and pushes come from the first copy only. One harmless exception: an event raised "
    "before the hub's clock was first set gets its real ts at the sync; if writing that ts to flash "
    "failed, two copies can differ in ts by about 1 s and in nothing else. The backend may treat two "
    "messages identical except for a ts at most 1 s apart as one. Two real messages can also be "
    "byte-identical: a sensor that goes wet, dry and wet again within one second with the same "
    "battery and RSSI loses its second leak_detected; the snapshot that follows within about 5 s has "
    "the true state.",

    "**Order (APP-NF-014).** Order events by ts (when the hub built the message), never by arrival or "
    "IoT Hub enqueued time. After a reconnect, the events the hub kept while offline go out oldest "
    "first, before that connection's lifecycle; new messages follow it. A cmd_ack for a command that "
    "arrives during that replay is sent at once and can come before the lifecycle and before older "
    "kept events. Snapshots and lifecycle messages are never kept: the hub builds them again after "
    "it reconnects. An event older than the latest snapshot goes into history only; it does not "
    "change the displayed state. At-least-once is not \"never lost\": an event accepted into a "
    "connection that is dying silently can still be lost, and the hub keeps at most 16 unsent "
    "events (section 5.5). After an outage the first snapshot is the source of truth.",
]

P198_TEXT = (
    "Lifecycle is sent once per MQTT connection: after every restart and after every reconnect "
    "without a restart. It tells the backend: \"I just came online, here is my config.\" The hub "
    "sends it after the events it kept while offline; new events follow it, but a cmd_ack for a "
    "command that arrives during that replay can come first (see Order above). "
    "Every hub reconnects about every 18 hours to renew its 24-hour SAS token, so a lifecycle "
    "without a restart is normal. On such a reconnect, reset_reason repeats the value from the last "
    "start. To tell a restart from a reconnect, take ts minus gateway.uptime_s, when the hub "
    "started: after a restart it is at or after the previous lifecycle's ts (allow 1 s for "
    "rounding); after a reconnect it equals the previous lifecycle's start time within a few "
    "seconds. "
    "Since 2.1.4 a hub with no devices sends a lifecycle too; a 2.1.3 hub that is not provisioned "
    "(provisioned false) sends no lifecycle and no snapshot, only twin reported (section 5.2.2). The "
    "backend should log "
    "each lifecycle and update its device registry, but take the device list from the snapshot "
    "(section 5.2.2): the lifecycle carries only counts."
)

T22_RESET = (
    "Why the hub last started: power_on, software, panic, watchdog, brownout, deep_sleep or "
    "unknown. software is a restart the hub made itself, for example after decommission target "
    "\"all\" or the 10-second Wi-Fi reset. A lifecycle sent after a reconnect without a restart "
    "repeats the value from the last start."
)
T22_PROV = (
    "The hub's stored provisioning state, not a device count. false on a new hub, after decommission "
    "target \"all\", and after a decommission removes the last device. true after any successful "
    "provision, even one that sets only rules or leaves no device. So do not use it to decide "
    "whether the hub has devices: use valve_id and the two counts here, or the device lists in the "
    "snapshot (section 5.2.2). Rarely, when the hub is busy for over 1 s, one lifecycle reads "
    "provisioned false with no valve_id and counts of 0; the next snapshot is authoritative."
)
T22_VALVE = (
    "The provisioned valve's BLE MAC, uppercase with colons (AA:BB:CC:DD:EE:FF). Present only when a "
    "valve is provisioned."
)
T22_LORA = (
    "Number of provisioned LoRa leak sensors (0-16). Always present. The app shows LoRa leak sensors "
    "read-only (section 2.4)."
)
T22_ACE = (
    "Auto-close master switch: while it is false, the hub never closes the valve on its own, "
    "whatever trigger_mask says. Default true. See section 4.8.1. The rules object is in every "
    "lifecycle unless the hub could not read its rules within 1 s: then keep the last known values. "
    "The hub goes back to the defaults (auto_close_enabled true, trigger_mask 7) after decommission "
    "target \"all\" and, since 2.1.4, when a decommission or a provision leaves it with no device. Up "
    "to 2.1.3 such a hub kept its last rules: until its next restart when a decommission emptied it, and "
    "across restarts when a provision emptied it (it stays provisioned)."
)
T22_MASK = (
    "Which leak sources can make the hub close the valve on its own (auto-close): bit 0 (value 1) BLE "
    "leak sensors, bit 1 (value 2) LoRa leak sensors, bit 2 (value 4) the valve flood probe. Other "
    "bits are not used. Default 7 (all three). A set bit acts only while auto_close_enabled is true "
    "and no override window is active. "
    "The hub does not range-check the value and reports it as stored. See section 4.8.1."
)

P218_TEXT = (
    "Snapshot is the full system state. The hub sends one after changes and as a heartbeat; "
    "data.reason says which (table below). The heartbeat comes when no other snapshot was sent for "
    "one snapshot interval: 300 s (5 minutes) by default, 60 to 3600 s through the Device Twin "
    "desired property snapshot_interval_s (section 5.4.2). The hub keeps the interval across "
    "restarts. A snapshot with data.reason \"event\" or \"heartbeat\" comes at least 5 s after the "
    "previous snapshot; the other reasons are sent without that spacing. Snapshots are sent only "
    "while the hub "
    "is connected and its clock is set; they are never kept offline. Since 2.1.4 a hub with no "
    "devices sends snapshots too (example after the field reference); a 2.1.3 hub with no devices "
    "sends none. The backend shall store only a single current-state record per hub - the \"current "
    "state\" - and overwrite (upsert) it with each new snapshot whose ts is not older than the stored "
    "one, because a resent snapshot can arrive late. Snapshot history is not retained. Event history "
    "is unaffected: events are appended separately. The app uses snapshots to build and refresh its "
    "entire UI."
)

REASON_INTRO = (
    "data.reason says why the hub sent the snapshot. It is not state: the app renders every snapshot "
    "the same way, whatever its data.reason, and uses data.reason only for logs and diagnostics. "
    "2.1.3 and 2.1.4 use the same six values:"
)
REASON_ROWS = [
    ["\"heartbeat\"",
     "No other snapshot was sent for one snapshot interval (snapshot_interval_s, default 300 s)."],
    ["\"event\"",
     "A few seconds after a change, and at least 5 s after the previous snapshot. The changes: a "
     "command that succeeded (after its cmd_ack; after a valve command this snapshot can still show "
     "the old position, and the valve's own report then brings another snapshot), leak_detected or leak_cleared, "
     "valve_state_changed, a rules event, a health event (device_offline or device_recovered), and "
     "the valve link dropping or coming back. Since 2.1.4 also a change of the valve's RMLEAK, a "
     "change of the valve's battery rating, and a change of the hub's rating to, from or between "
     "warning and critical. Also every 30 s and on each sensor report for 5 minutes after a "
     "provision that adds devices (at most 40 snapshots)."],
    ["\"fast\"",
     "The first snapshot after an MQTT connect, sent as soon as the valve link is ready (at the "
     "latest 150 s after the connect) if the \"boot\" snapshot has not gone out yet. Sensors not "
     "heard yet show connected false and null values."],
    ["\"boot\"",
     "The first complete snapshot after an MQTT connect, or after a provision that adds devices: "
     "once every device has been heard, or when the sync window ends (3 minutes after a restart, "
     "2.5 minutes after a provision). After a reconnect without a restart, and on a hub with no "
     "devices, it follows the lifecycle within seconds."],
    ["\"commission\"",
     "Within 6 minutes after a \"fast\" snapshot or a provision that adds devices: each time a "
     "device is heard for the first time."],
    ["\"decommission\"",
     "The last snapshot before the hub restarts after decommission target \"all\". It shows the "
     "emptied hub."],
]
EXAMPLE_INTRO = (
    "Example: a heartbeat snapshot from a hub with a valve and two BLE leak sensors, one of them "
    "offline."
)

SNAPSHOT_LINES = [
    '{',
    '  "schema": "eflostop.v2",',
    '  "ts": 1770589401,',
    '  "gateway": { "id": "GW-34B7DA6AAD54", "short_id": "AD54", "name": "Lobby Hub", "fw": "2.1.4", "uptime_s": 971 },',
    '  "type": "snapshot",',
    '  "data": {',
    '    "reason": "heartbeat",',
    '    "system_health": {',
    '      "rating": "critical",',
    '      "reason": "1 sensor offline"',
    '    },',
    '    "valve": {',
    '      "valve_id": "00:80:E1:27:9A:E6",',
    '      "state": "open",',
    '      "battery": 18,',
    '      "leak_state": false,',
    '      "rmleak": false,',
    '      "connected": true,',
    '      "fw_version": "2.2.0",',
    '      "rating": "warning",',
    '      "last_seen_age_s": 0',
    '    },',
    '    "lora_sensors": [],',
    '    "ble_leak_sensors": [',
    '      {',
    '        "sensor_id": "00:80:E1:27:B4:96",',
    '        "connected": true,',
    '        "rating": "excellent",',
    '        "last_seen_age_s": 12,',
    '        "battery": 78,',
    '        "rssi": -72,',
    '        "leak_state": false,',
    '        "fw_version": "1.1.0",',
    '        "location": { "code": "laundry", "label": "Behind washer" }',
    '      },',
    '      {',
    '        "sensor_id": "00:80:E1:27:B6:A5",',
    '        "connected": false,',
    '        "rating": "critical",',
    '        "last_seen_age_s": 655,',
    '        "battery": 64,',
    '        "rssi": -86,',
    '        "leak_state": false,',
    '        "fw_version": null,',
    '        "location": { "code": "bathroom", "label": "" }',
    '      }',
    '    ],',
    '    "rules": {',
    '      "auto_close_enabled": true,',
    '      "trigger_mask": 7',
    '    },',
    '    "override_active": false',
    '  }',
    '}',
]
EXAMPLE_NOTE = (
    "In this example the second sensor was last heard 655 s ago, so it is offline: its rating, and "
    "so the hub rating, are critical. It keeps its last battery and RSSI; its fw_version is null "
    "while it is offline. The valve battery (18 %) rates the valve warning, but \"Valve battery low\" "
    "is not in system_health.reason, because the reason names only the causes at the hub's rating "
    "level."
)

# T23: (row, column, expect, new text, finding ids) for in-place cell edits.
T23_REASON_ROW = [
    "data.reason", "string",
    "Why the hub sent this snapshot: \"heartbeat\", \"event\", \"fast\", \"boot\", \"commission\" or "
    "\"decommission\" (see the table above). Always present.",
]
T23_RATING = (
    "\"excellent\", \"good\", \"warning\" or \"critical\". The worst rating among the provisioned "
    "devices, with three exceptions. A device not heard since the hub started, or since a provision "
    "added it, is left out for its first 10 minutes (a sensor) or 3 minutes (the valve; 2.5 minutes "
    "when a provision added it), although its own entry already reads critical and connected false "
    "(APP-FR-086). While a leak incident is latched, the rating is at least \"warning\" (reason "
    "\"Leak interlock latched\"). A hub with no devices reads \"excellent\" (since 2.1.4; a 2.1.3 hub "
    "with no devices sends no snapshot)."
)
T23_HREASON = (
    "Human-readable text for display. Show it as sent; do not parse it. It names only the causes at "
    "the hub's rating level, most urgent first, so a lesser issue (for example \"Valve battery low\" "
    "while a sensor is offline) is not in it: read each device's rating for those. Possible texts: "
    "\"All devices healthy\"; \"Syncing - waiting for N devices\"; \"No devices provisioned\" (since "
    "2.1.4); \"Degraded\"; or a list joined by \", \" of: \"Leak detected: <name>\" (the sensor's label, else its "
    "sensor_id, or \"valve\"), \"N leaks detected\", \"Leak interlock latched\", \"Valve offline\", \"Valve "
    "disconnected\", \"Valve battery critical\" (since 2.1.4), \"Valve battery low\", \"N sensors "
    "offline\", \"N sensors battery low\", \"N sensors signal weak\", \"syncing N devices\". With N = 1 "
    "the noun is singular, for example \"1 sensor offline\" or \"Syncing - waiting for 1 device\". "
    "2.1.3 can also send \"Health data unavailable\"; 2.1.4 never does."
)
T23_VALVE_OBJ = [
    "data.valve", "object",
    "The provisioned valve. Since 2.1.4 an empty object {} when no valve is provisioned. A 2.1.3 hub "
    "with no valve sends {\"state\": \"disconnected\", \"connected\": false} with no valve_id or, while "
    "it is linked to a nearby valve (section 2.4), that valve's valve_id and readings. Treat data.valve "
    "as \"no valve\" when it has no valve_id and, on 2.1.3, also while twin reported valve_id is null. "
    "While the valve is not linked, the object has only valve_id, state "
    "(\"disconnected\"), connected (false), rating and last_seen_age_s; battery, leak_state, rmleak "
    "and fw_version are left out. Example: {\"valve_id\": \"00:80:E1:27:9A:E6\", \"state\": "
    "\"disconnected\", \"connected\": false, \"rating\": \"critical\", \"last_seen_age_s\": 245}",
]
T23_VALVE_ID = (
    "The provisioned valve's BLE MAC, uppercase with colons (AA:BB:CC:DD:EE:FF). Present whenever a "
    "valve is provisioned. Since 2.1.4 it is always the provisioned valve's MAC, and the live fields "
    "below come only from a link to that valve; 2.1.3 could show the MAC and readings of whichever "
    "valve was linked."
)
T23_STATE = (
    "\"open\", \"closed\", \"disconnected\" or \"unknown\". \"disconnected\": the hub has no link to the "
    "valve. \"unknown\": the link is up but the valve's readings are not in yet, or the valve reported "
    "no valid position. Since 2.1.4 it reads \"unknown\" until every reading is in. A 2.1.3 hub also "
    "sends \"unknown\", but only until it has read the position; for a moment after that its battery "
    "reads 0 and rmleak false. While state is \"unknown\", ignore leak_state and rmleak: just after the "
    "link comes up they read false, even on a valve locked after a leak, and battery is null (0 from a "
    "2.1.3 hub)."
)
T23_BATT = (
    "Valve battery % (0-100). Present only while the valve is linked; null while state is "
    "\"unknown\" (linked, not read yet). Valve bands since 2.1.4: 10 % or less is critical (\"Valve "
    "battery critical\"; the hub then refuses valve_open and valve_set_state \"open\"), 11-20 % is "
    "warning (\"Valve battery low\"), above 20 % is excellent (the valve has no \"good\" band). 2.1.3 "
    "rated the valve battery like a sensor's (20 % or less warning, 21-35 % good), had no critical "
    "band, and sent 0 when it had not read the battery: treat a valve battery of 0 from a 2.1.3 hub "
    "as unknown."
)
T23_VLEAK = (
    "true when the valve flood probe is wet. Present only while the valve is linked; ignore it while "
    "state is \"unknown\"."
)
T23_RMLEAK = (
    "true while the valve's RMLEAK interlock is set: the valve is locked closed after an auto-close "
    "and cannot be opened until RMLEAK is cleared (leak_reset, the hub's automatic clear about 10 s "
    "after every leak source is dry, or the 24-hour override; section 4.4). Present only while the "
    "valve is linked; ignore it while state is \"unknown\". When it is missing, the lock state is "
    "unknown: never read it as false. Since 2.1.4 the hub also refuses valve_open and "
    "valve_set_state \"open\" while a leak incident is latched and no override window is active, "
    "whether the valve is linked or not; "
    "2.1.3 checked only the valve's own RMLEAK."
)
T23_VCONN = (
    "true while the hub has a BLE link to the provisioned valve. The readings may not be in yet: "
    "check that state is not \"unknown\". Absent when data.valve is {} (no valve provisioned). 2.1.3 "
    "also reported true for a link to a valve other than the provisioned one."
)
T23_VRATING = (
    "The valve's health rating. critical: its flood probe is wet, it has been unlinked for 3 minutes "
    "or more (\"Valve offline\"), it has not linked since the hub started or since it was added, or "
    "(since 2.1.4) its battery is 10 % or less. warning: unlinked for less than 3 minutes (\"Valve "
    "disconnected\"), or battery 11-20 %. Otherwise excellent: since 2.1.4 the valve has no \"good\" "
    "band. 2.1.3: battery 20 % or less warning, 21-35 % good, no battery critical."
)
T23_VSEEN = (
    "0 while the valve is linked; seconds since the link dropped once it is unlinked; null if the "
    "valve has not linked since the hub started or since it was added. 2.1.3 counted from the "
    "valve's last changed value, so a steady linked valve showed hundreds of seconds."
)
T23_VFW = (
    "Valve firmware version, for example \"2.2.0\". Present only while the valve is linked. null "
    "until the hub has read it, or if the valve does not report one. Left out (not null) while the "
    "valve is not linked."
)
T23_LORA_ROW = [
    "data.lora_sensors[]", "array",
    "One entry per provisioned LoRa leak sensor. Always present; [] when there are none. The app "
    "lists these sensors and their leaks read-only (section 2.4). Each entry has these keys, in this "
    "order: sensor_id (\"0x\" and 8 uppercase hex digits, for example \"0x1A2B3C4D\"), connected, "
    "rating, last_seen_age_s, battery, rssi, leak_state, snr, location. snr (number/null) is the "
    "signal-to-noise ratio in dB of the last packet heard, up to 10 minutes old; it is null while "
    "connected is false. The other keys mean the same as in a ble_leak_sensors entry. There is no "
    "fw_version.",
]
T23_SID = (
    "The sensor's BLE MAC, uppercase with colons, for example \"00:80:E1:27:B4:96\". One entry per "
    "provisioned BLE leak sensor. Key entries by sensor_id, not by array position."
)
T23_SCONN = (
    "true while the hub has heard the sensor within the last 10 minutes, false after that. The "
    "sensor's rating turns critical, and device_offline is sent, at the hub's next 30 s check, so 10 "
    "to 10.5 minutes after the last report; a snapshot sent in between can show connected false next "
    "to the previous rating."
)
T23_SRATING = (
    "The sensor's health rating. critical: wet, not heard for 10 minutes, or not heard since the hub "
    "started or since it was added. warning: battery 20 % or less, or RSSI -90 dBm or weaker. good: "
    "battery 21-35 %, or RSSI -89 to -80 dBm. Otherwise excellent. A sensor's battery never makes it "
    "critical."
)
T23_SSEEN = (
    "Seconds since the hub last heard the sensor. null if it has not been heard since the hub "
    "started or since it was added. Since 2.1.4, adding or removing other devices does not reset it. "
    "Up to 2.1.3, every provision or decommission returned every remaining sensor to this not-heard "
    "state (null values, connected false, rating critical) until it was heard again."
)
T23_SBATT = (
    "Last battery % heard (0-100), kept while the sensor is offline. null until the sensor is first "
    "heard since the hub started or since it was added."
)
T23_SLEAK = (
    "true while the hub holds the sensor as wet. Always true or false, never null. A sensor that goes "
    "offline while wet stays true (with connected false) until it reports dry."
)
T23_SRSSI = (
    "Last signal strength heard, in dBm, kept while the sensor is offline. null until the sensor is "
    "first heard since the hub started or since it was added."
)
T23_SCODE = (
    "The sensor's location: one of the 13 codes in Appendix A.1, for example \"laundry\". "
    "\"unknown\" until a location is set."
)
T23_SLABEL = (
    "User-defined label, up to 31 bytes of UTF-8 (31 characters of plain ASCII, fewer for accented "
    "or non-Latin text). Empty string if not set. The hub does not reject a longer label: it keeps "
    "only the first 31 bytes, which can cut a character in half and leave invalid text. So the app "
    "limits label input by UTF-8 byte length (APP-FR-068)."
)
T23_SFW = (
    "Sensor firmware version, for example \"1.1.0\". null while the sensor is offline, until it is "
    "heard, or if its firmware does not report a version."
)
T23_RULES_ROWS = [
    ["data.rules.auto_close_enabled", "boolean",
     "Auto-close master switch. Default true. Same object as in the lifecycle (section 5.2.1). The "
     "Settings screen reads both rules values from the latest snapshot (section 4.8.1)."],
    ["data.rules.trigger_mask", "integer",
     "Which leak sources can make the hub close the valve on its own: bit 0 (value 1) BLE leak "
     "sensors, bit 1 (value 2) LoRa leak sensors, bit 2 (value 4) the valve flood probe. Default 7. "
     "data.rules is left out of a snapshot only if the hub could not read its rules within 1 s: keep "
     "the last known values."],
]
T23_OVR = (
    "true while the 24-hour override window is active: leaks are still reported, but the hub does "
    "not close the valve on its own. The window ends at expiry, with override_cancel, with an "
    "accepted leak_reset, with decommission target \"all\" and, since 2.1.4, when the hub's last "
    "device is removed. When override_active turns false, the window has ended, whatever the reason. "
    "Rarely, when the hub is busy for over 1 s, one snapshot reads override_active false during a "
    "window; the next snapshot shows it again."
)
T23_OVR_REM = (
    "Seconds left in the 24-hour override window. Present only while override_active is true. Drive "
    "the countdown from this value."
)
T23_EXPIRES_ROW = [
    "data.expires_ts", "integer",
    "Unix epoch seconds at which the override window ends. Present only while override_active is "
    "true and the end is known on the synced clock. Since 2.1.4 it is left out for up to about 30 s "
    "after the first clock sync when the window started before it; use override_remaining_s.",
]

EMPTY_INTRO = (
    "**Hub with no devices.** Since 2.1.4 a hub with no devices still sends a lifecycle at each "
    "connect, twin reported, and snapshots (one soon after each connect, then heartbeats). Its "
    "snapshot looks like this:"
)
EMPTY_LINES = [
    '{',
    '  "schema": "eflostop.v2",',
    '  "ts": 1770589401,',
    '  "gateway": { "id": "GW-34B7DA6A1F20", "short_id": "1F20", "fw": "2.1.4", "uptime_s": 52 },',
    '  "type": "snapshot",',
    '  "data": {',
    '    "reason": "boot",',
    '    "system_health": { "rating": "excellent", "reason": "No devices provisioned" },',
    '    "valve": {},',
    '    "lora_sensors": [],',
    '    "ble_leak_sensors": [],',
    '    "rules": { "auto_close_enabled": true, "trigger_mask": 7 },',
    '    "override_active": false',
    '  }',
    '}',
]
EMPTY_AFTER = [
    "data.reason varies as in the table above, and rules shows the stored values, normally the "
    "defaults. Removing the last device clears the leak latch and the override window and resets "
    "the rules, with no event. Decide \"no devices\" from the device lists: data.valve has no "
    "valve_id and both sensor arrays are empty (in lifecycle and twin reported: no valve_id or "
    "valve_id null, and both counts 0). Do not use the provisioned flag: it can read true with no "
    "devices after a provision that set only rules or emptied the sensor lists. The app shows \"No "
    "devices set up yet\", not \"All Clear\". A hub with sensors but no valve sends \"valve\": {} too.",

    "A 2.1.3 hub that is not provisioned (provisioned false), such as a new hub or one emptied by a "
    "decommission, sends no lifecycle and no snapshot, only twin reported. While IoT Hub reports it "
    "connected, the app shows \"No devices set up yet\" and uses only the IoT Hub connection state for "
    "online/offline. A 2.1.3 hub left with no devices by a provision stays provisioned and keeps sending "
    "lifecycle and snapshots, with \"valve\": {\"state\": \"disconnected\", \"connected\": false} and "
    "empty sensor arrays; the device-list rule above still identifies it. A 2.1.3 hub with sensors but "
    "no valve also sends \"valve\": {\"state\": \"disconnected\", \"connected\": false}.",
]
DEVSET_TEXT = (
    "**After a provision or decommission.** The snapshot that follows the command's cmd_ack already "
    "shows the new device set. Devices that stay keep their position and status. A removed device "
    "disappears from its array. A newly added device is appended at the end of its array and reads "
    "not heard (connected false, null values, rating critical) until the hub hears it (APP-FR-086). "
    "Up to 2.1.3, every provision or decommission also returned the remaining devices to the "
    "not-heard state until each was heard again."
)


def _two_col_template(ctx):
    """A 2-column V4 table to copy formatting from (T5, 2.6 Hub States). Read-only use."""
    if ctx.orig_cell(5, 0, 0) != "State" or len(ctx.tables[5].columns) != 2:
        raise EditError("s3: T5 is not the expected 2-column template (State | Meaning for the App)")
    return ctx.tables[5]


def _size9(row):
    """Give every run in a row the 9 pt size of V4's table body text (V4's T23 R11 and R21 runs
    carry no size and print at the document default)."""
    for c in unique_cells(row):
        for p in c.paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)


# tcPr children that must follow w:shd (ECMA-376 CT_TcPr sequence).
_AFTER_SHD = {qn(t) for t in ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign",
                              "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge",
                              "w:tcPrChange")}


def _body_shading(ctx):
    """V4's body-row shading element (<w:shd w:fill="F5F5F5"/>), copied from T20 R1 before any edit."""
    tc = ctx.row(20, 1, "D2C (telemetry)")._tr.findall(qn("w:tc"))[0]
    tcpr = tc.find(qn("w:tcPr"))
    shd = tcpr.find(qn("w:shd")) if tcpr is not None else None
    if shd is None or shd.get(qn("w:fill")) != "F5F5F5":
        raise EditError("s3: T20 R1 has no F5F5F5 body shading to copy")
    return copy.deepcopy(shd)


def _reband(table, shd_tpl):
    """V4's zebra banding: body rows 1, 3, 5 ... shaded F5F5F5, rows 2, 4, 6 ... unshaded.
    The header row (row 0) is left as it is."""
    for i, row in enumerate(table.rows):
        if i == 0:
            continue
        for tc in row._tr.findall(qn("w:tc")):
            tcpr = tc.get_or_add_tcPr()
            for old in tcpr.findall(qn("w:shd")):
                tcpr.remove(old)
            if i % 2 == 0:
                continue
            shd = copy.deepcopy(shd_tpl)
            nxt = next((ch for ch in tcpr if ch.tag in _AFTER_SHD), None)
            if nxt is not None:
                nxt.addprevious(shd)
            else:
                tcpr.append(shd)


def apply(ctx):
    shd_tpl = _body_shading(ctx)   # read before any edit; used by the re-banding at the end

    # ------------------------------------------------------------ 5. intro
    set_text(ctx.P(180, "This section is the ground truth for message formats"), P180_TEXT)
    ctx.note("s3-01: 5 intro names hub firmware 2.1.4 as the reference and states the 2.1.3 version policy (canon C1).")

    # ------------------------------------------------------------ 5.1 T20
    ctx.T(20, "MQTT Topic")
    set_cell(ctx.cell(20, 1, 2, "All telemetry from hub to cloud"), T20_D2C)
    ctx.note("s3-02: T20 D2C row: QoS 1, {device_id} is the Gateway ID, no message properties (parse the body; $body routing does not match).")
    set_cell(ctx.cell(20, 2, 2, "Commands from cloud to hub"), T20_C2D)
    ctx.note("s3-03: T20 C2D row: QoS 1 subscription and the 8192-byte limit with its cmd_ack error.")
    set_cell(ctx.cell(20, 3, 2, "Hub reports state"), T20_REPORTED)
    ctx.note("s3-04: T20 twin reported row: when the hub reports; since 2.1.4 also after a single-device decommission.")
    r4 = ctx.row(20, 4, "Twin desired")
    set_cell(ctx.cell(20, 4, 2, "Backend pushes config"), T20_DESIRED)
    ctx.note("s3-05: T20 twin desired row: snapshot_interval_s and hub_name, kept across restarts, applied again at every connect.")
    r = add_row(ctx.T(20, "MQTT Topic"), ["Twin GET", "$iothub/twin/GET/?$rid={n}", T20_GET], like_row=r4, after_row=r4)
    add_row(ctx.T(20, "MQTT Topic"), ["Twin responses", "$iothub/twin/res/#", T20_RES], like_row=r4, after_row=r)
    ctx.note("s3-06: T20 adds the twin GET (once per connect) and twin response topics.")

    # ------------------------------------------------------------ 5.2 envelope
    set_text(ctx.P(183, "All D2C telemetry messages use this top-level envelope"), P183_TEXT)
    p188 = ctx.P(188, '"id": "GW-34B7DA6AAD54"')
    insert_lines_after(p188, ['    "short_id": "AD54",', '    "name": "Lobby Hub",'], like=p188)
    replace(ctx.P(189, '"fw": "1.9.0"'), "1.9.0", "2.1.4")
    ctx.note("s3-07: envelope example adds gateway.short_id and gateway.name, fw 2.1.4; key order and optional name stated.")

    ctx.T(21, "Field | Type | Description")
    set_cell(ctx.cell(21, 2, 2, "Unix epoch timestamp"), T21_TS)
    ctx.note("s3-08: T21 ts: build time, kept on resends, order by ts; since 2.1.4 pre-sync events are kept (2.1.3 discarded them).")
    set_cell(ctx.cell(21, 3, 2, "Gateway ID like"), T21_ID)
    ctx.note("s3-09: T21 gateway.id: format and case.")
    set_cell(ctx.cell(21, 4, 2, "Firmware version string"), T21_FW)
    ctx.note("s3-10: T21 gateway.fw: example 2.1.4, compare numerically.")
    set_cell(ctx.cell(21, 5, 1, "number"), "integer")
    set_cell(ctx.cell(21, 5, 2, "Hub uptime in seconds"), T21_UPTIME)
    ctx.note("s3-11: T21 gateway.uptime_s: integer; the start time (ts minus uptime_s) against the previous lifecycle tells a restart from a reconnect.")
    set_cell(ctx.cell(21, 7, 2, "User-assigned hub name"), T21_NAME)
    ctx.note("s3-12: T21 gateway.name: limit is 31 bytes of UTF-8; set through desired hub_name.")
    set_cell(ctx.cell(21, 8, 2, '"snapshot", "lifecycle", or "event"'), T21_TYPE)
    ctx.note("s3-v04: T21 type: every event, cmd_ack included, is type \"event\", named by data.event.")

    p196 = ctx.P(196, "Forward compatibility:")
    replace(p196, P196_OLD, P196_NEW)
    ctx.note("s3-13: forward compatibility also covers unknown enum values, nulls and an empty data.valve.")
    insert_lines_after(p196, DELIVERY, like=p196, rich=True)
    ctx.note("s3-14: 5.2 adds delivery, dedupe (APP-NF-013) and order (APP-NF-014) rules per canon C10/C11.")

    # ------------------------------------------------------------ 5.2.1 lifecycle
    set_text(ctx.P(198, "Lifecycle is sent every time the hub connects"), P198_TEXT)
    ctx.note("s3-15: lifecycle: once per connection, after the kept events; reset_reason repeats on a reconnect; 2.1.3 vs 2.1.4 for a hub with no devices; device list from the snapshot.")

    replace(ctx.P(201, '"ts": 1770589401'), "1770589401", "1770588454")
    set_text(ctx.P(202, '"gateway": { "id": "GW-34B7DA6AAD54", "fw": "1.9.0", "uptime_s": 3 }'),
             '  "gateway": { "id": "GW-34B7DA6AAD54", "short_id": "AD54", "name": "Lobby Hub", "fw": "2.1.4", "uptime_s": 24 },')
    p208 = ctx.P(208, '"valve_device_id": "00:80:E1:27:9A:E6"')
    set_text(p208, '    "valve_id": "00:80:E1:27:9A:E6",')
    insert_after(p208, '    "lora_sensor_count": 0,', like=p208)
    replace(ctx.P(209, '_leak_sensor_count": 3,'), '": 3,', '": 2,')
    replace(ctx.P(212, '_mask": 5'), '": 5', '": 7')
    ctx.note("s3-16, tm-18: lifecycle example: valve_id (not valve_device_id), lora_sensor_count, short_id and name, fw 2.1.4, 2 BLE sensors to match the snapshot, trigger_mask 7.")

    ctx.T(22, "Field | Type | Description")
    set_cell(ctx.cell(22, 2, 2, "Why the hub rebooted"), T22_RESET)
    ctx.note("s3-17: T22 reset_reason: reason for the last restart, repeated on a reconnect; what software covers.")
    set_cell(ctx.cell(22, 3, 2, "true if hub has at least one device assigned"), T22_PROV)
    ctx.note("s3-18: T22 provisioned: stored provisioning state, not a device count; decide devices from valve_id and the counts.")
    row4 = ctx.row(22, 4, "data.valve_device_id")
    set_cell(ctx.cell(22, 4, 0, "data.valve_device_id"), "data.valve_id")
    set_cell(ctx.cell(22, 4, 2, "Valve BLE MAC, uppercase"), T22_VALVE)
    ctx.note("s3-19: T22 data.valve_device_id becomes data.valve_id (the key every hub sends); old-key history left to section 2.5 (canon C22).")
    add_row(ctx.T(22, "Field | Type | Description"), ["data.lora_sensor_count", "integer", T22_LORA], like_row=row4, after_row=row4)
    ctx.note("s3-20: T22 adds data.lora_sensor_count (LoRa wording per canon C20).")
    set_cell(ctx.cell(22, 6, 2, "Whether auto-close is turned on"), T22_ACE)
    ctx.note("s3-21, tm-19: T22 rules.auto_close_enabled: master switch, default, reset to defaults, left out when the hub is busy.")
    set_cell(ctx.cell(22, 7, 2, "Bitmask of which sensor types trigger auto-close"), T22_MASK)
    ctx.note("s3-v01, tm-19: T22 rules.trigger_mask: bits 0/1/2 with values, default 7, no range check.")

    # ------------------------------------------------------------ 5.2.2 snapshot
    p218 = ctx.P(218, "Snapshot is the full system state sent periodically")
    set_text(p218, P218_TEXT)
    ctx.note("s3-22: snapshot intro: change-driven plus heartbeat, 5 s spacing, never kept offline, interval persisted, upsert only with a ts not older than the stored one.")

    pa = insert_after(p218, REASON_INTRO, like=p218)
    tbl = insert_table_after(pa, _two_col_template(ctx), ["data.reason", "When the hub sends it"], REASON_ROWS)
    insert_after(tbl, EXAMPLE_INTRO, like=p218)   # the example (P219...) follows it
    ctx.note("s3-23: snapshot adds the data.reason table (canon C9; \"boot\" without the Wi-Fi setup clause).")

    # Snapshot example: reuse V4's 48 code paragraphs P219-P266, then append the rest.
    ctx.P(219, "{")
    last = ctx.P(266, "}")
    ctx.P(222, '"fw": "1.9.0", "uptime_s": 971')
    ctx.P(230, '"device_id": "00:80:E1:27:9A:E6"')
    n_old = 266 - 219 + 1
    if len(SNAPSHOT_LINES) < n_old:
        raise EditError("s3: snapshot example shorter than V4's; adjust the reuse logic")
    for i, line in enumerate(SNAPSHOT_LINES[:n_old]):
        set_text(ctx.P(219 + i), line)
    tail = ctx.code_after(last, SNAPSHOT_LINES[n_old:])
    insert_after(tail, EXAMPLE_NOTE, like=p218)
    ctx.note("s3-24, tm-20: snapshot example rewritten: data.reason, valve_id and sensor_id keys, lora_sensors, rules, wire key order, consistent ratings and reason; with an explanation.")

    # Field reference T23 (V4 rows R0-R23).
    ctx.T(23, "Path | Type | Notes")
    add_row(ctx.T(23, "Path | Type | Notes"), T23_REASON_ROW,
            like_row=ctx.row(23, 1, "data.system_health.rating"), after_row=ctx.row(23, 0, "Path"))
    ctx.note("s3-27: T23 adds data.reason.")
    set_cell(ctx.cell(23, 1, 2, "Worst rating across all devices"), T23_RATING)
    ctx.note("s3-25: T23 system_health.rating: the three roll-up exceptions (not heard yet, interlock floor, empty hub).")
    set_cell(ctx.cell(23, 2, 2, "Human-readable reason"), T23_HREASON)
    ctx.note("s3-26: T23 system_health.reason: only causes at the hub rating, the full list of texts, show as sent.")
    add_row(ctx.T(23, "Path | Type | Notes"), T23_VALVE_OBJ,
            like_row=ctx.row(23, 2, "data.system_health.reason"),
            after_row=ctx.row(23, 2, "data.system_health.reason"))
    ctx.note("s3-28: T23 adds the data.valve object shapes (empty {}, unlinked, 2.1.3 no-valve shape).")
    set_cell(ctx.cell(23, 3, 0, "data.valve.device_id"), "data.valve.valve_id")
    set_cell(ctx.cell(23, 3, 2, "Valve BLE MAC address"), T23_VALVE_ID)
    ctx.note("s3-29: T23 data.valve.device_id becomes data.valve.valve_id; always the provisioned valve since 2.1.4; key history left to 2.5 (canon C22).")
    set_cell(ctx.cell(23, 4, 2, '"open", "closed", "disconnected", or "unknown"'), T23_STATE)
    ctx.note("s3-30: T23 valve.state: meaning of disconnected and unknown; ignore leak_state and rmleak while unknown.")
    set_cell(ctx.cell(23, 5, 1, "integer"), "integer/null")
    set_cell(ctx.cell(23, 5, 2, "Battery percentage"), T23_BATT)
    ctx.note("s3-31: T23 valve.battery: integer/null, valve bands since 2.1.4, open refused at 10 % or less, 2.1.3 sent 0 when unread.")
    set_cell(ctx.cell(23, 6, 2, "Valve built-in flood probe status"), T23_VLEAK)
    ctx.note("s3-32: T23 valve.leak_state: present only while linked.")
    set_cell(ctx.cell(23, 7, 2, "RMLEAK latch is asserted"), T23_RMLEAK)
    ctx.note("s3-33: T23 valve.rmleak: missing means unknown, never false; since 2.1.4 valve_open refused while an incident is latched.")
    set_cell(ctx.cell(23, 8, 2, "true if BLE connection is active"), T23_VCONN)
    ctx.note("s3-v02: T23 valve.connected: link to the provisioned valve only; readings may not be in yet; absent for {}.")
    set_cell(ctx.cell(23, 9, 2, "Health rating for the valve device"), T23_VRATING)
    ctx.note("s3-34: T23 valve.rating: valve rating rules, 2.1.4 battery bands, 2.1.3 bands.")
    set_cell(ctx.cell(23, 10, 1, "number/null"), "integer/null")
    set_cell(ctx.cell(23, 10, 2, "Seconds since last BLE contact"), T23_VSEEN)
    ctx.note("s3-35: T23 valve.last_seen_age_s: 0 while linked since 2.1.4; 2.1.3 meaning.")
    set_cell(ctx.cell(23, 11, 2, "Valve firmware version"), T23_VFW)
    _size9(ctx.row(23, 11, "data.valve.fw_version"))   # V4's runs here carry no size (9 pt like the rest)
    ctx.note("s3-36: T23 valve.fw_version: left out (not null) while unlinked; example 2.2.0; row text set to the table's 9 pt.")
    # like_row R10: V4's R11 has no run size, so a copy of it would print larger than the table.
    add_row(ctx.T(23, "Path | Type | Notes"), T23_LORA_ROW,
            like_row=ctx.row(23, 10, "data.valve.last_seen_age_s"),
            after_row=ctx.row(23, 11, "data.valve.fw_version"))
    ctx.note("s3-37: T23 adds data.lora_sensors[] (shown read-only per canon C20).")
    set_cell(ctx.cell(23, 12, 0, "data.ble_leak_sensors[].device_id"), "data.ble_leak_sensors[].sensor_id")
    set_cell(ctx.cell(23, 12, 2, "BLE MAC address string"), T23_SID)
    ctx.note("s3-38: T23 ble_leak_sensors[].device_id becomes sensor_id; key entries by sensor_id.")
    set_cell(ctx.cell(23, 13, 2, "advertisement was received within the timeout"), T23_SCONN)
    ctx.note("s3-39: T23 sensor connected: 10-minute timeout, checked every 30 s.")
    set_cell(ctx.cell(23, 14, 2, "Per-sensor health rating"), T23_SRATING)
    ctx.note("s3-40: T23 sensor rating: the rating rules.")
    set_cell(ctx.cell(23, 15, 1, "number/null"), "integer/null")
    set_cell(ctx.cell(23, 15, 2, "Seconds since last BLE advertisement"), T23_SSEEN)
    ctx.note("s3-41: T23 sensor last_seen_age_s: null means not heard since start or since added; 2.1.3 reset all sensors on any provision or decommission.")
    set_cell(ctx.cell(23, 16, 2, "Battery %"), T23_SBATT)
    ctx.note("s3-42: T23 sensor battery: last heard, kept while offline.")
    set_cell(ctx.cell(23, 17, 2, "true if leak detected"), T23_SLEAK)
    ctx.note("s3-43: T23 sensor leak_state: never null; stays true while offline and wet.")
    set_cell(ctx.cell(23, 18, 2, "Signal strength in dBm"), T23_SRSSI)
    ctx.note("s3-44: T23 sensor rssi: last heard, kept while offline, null until heard.")
    set_cell(ctx.cell(23, 19, 2, "Location code from sensor_meta"), T23_SCODE)
    ctx.note("s3-v03: T23 location.code: the 13 codes of A.1, \"unknown\" until set.")
    set_cell(ctx.cell(23, 20, 2, "User-defined label (max 31 chars)"), T23_SLABEL)
    ctx.note("s3-45: T23 location.label: 31 bytes of UTF-8, longer labels cut silently.")
    set_cell(ctx.cell(23, 21, 2, "Sensor firmware version"), T23_SFW)
    _size9(ctx.row(23, 21, "data.ble_leak_sensors[].fw_version"))   # V4's runs here carry no size
    ctx.note("s3-46: T23 sensor fw_version: null while offline; example 1.1.0; advertisement byte offsets removed; row text set to the table's 9 pt.")
    r = ctx.row(23, 21, "data.ble_leak_sensors[].fw_version")
    for vals in T23_RULES_ROWS:
        # like_row R20: V4's R21 has no run size (see the LoRa row above).
        r = add_row(ctx.T(23, "Path | Type | Notes"), vals,
                    like_row=ctx.row(23, 20, "data.ble_leak_sensors[].location.label"), after_row=r)
    ctx.note("s3-47, tm-21: T23 adds data.rules.auto_close_enabled and data.rules.trigger_mask.")
    set_cell(ctx.cell(23, 22, 2, "24-hour override window is active"), T23_OVR)
    ctx.note("s3-v05: T23 override_active: what the window means and every way it ends (since 2.1.4 also when the last device is removed).")
    r23 = ctx.row(23, 23, "data.override_remaining_s")
    set_cell(ctx.cell(23, 23, 2, "Seconds remaining in override window"), T23_OVR_REM)
    add_row(ctx.T(23, "Path | Type | Notes"), T23_EXPIRES_ROW, like_row=r23, after_row=r23)
    ctx.note("s3-48: T23 override_remaining_s reworded; adds data.expires_ts (left out briefly after the first clock sync since 2.1.4).")

    # After T23: the hub with no devices, and device-set changes.
    p_rich = ctx.P(196, "Forward compatibility:")   # bold lead-in template (formatting only)
    t23 = ctx.T(23, "Path | Type | Notes")
    pe = insert_after(t23, EMPTY_INTRO, like=p_rich, rich=True)
    last_code = ctx.code_after(pe, EMPTY_LINES)
    pz = insert_lines_after(last_code, EMPTY_AFTER, like=p218)
    insert_after(pz, DEVSET_TEXT, like=p_rich, rich=True)
    ctx.note("global fix XC-11: 5.2.1 and 5.2.2: a 2.1.3 hub goes silent when it is not provisioned; one emptied by a provision keeps sending (master app_iothub.c:2418, provisioning_manager.c:591)")
    ctx.note("global fix INT-17: the event examples leave out gateway.name (said in 5.2); the no-devices snapshot is another hub (GW-34B7DA6A1F20), not the populated hub at the same ts")
    ctx.note("global fix INT-09/XC-03: T22 rules.auto_close_enabled: 2.1.3 kept its rules until restart after an emptying decommission, across restarts after an emptying provision")
    ctx.note("global fix XC-05: T23 data.valve: a 2.1.3 hub with no valve can carry a nearby valve's valve_id; use twin reported valve_id on 2.1.3")
    ctx.note("global fix XC-04/completeness-05: T23 valve.state: a 2.1.3 hub also sends \"unknown\" until it reads the position (master app_ble_valve.c:1409-1411, :1464-1466 reset state to -1; telemetry_v2.c:708-712), then battery 0 and rmleak false for a moment (setup reads state, flood, rmleak, battery)")
    ctx.note("s3-49: 5.2.2 adds the snapshot of a hub with no devices (since 2.1.4; 2.1.3 behavior per canon C8) and how the arrays change after a provision or decommission.")

    # Zebra banding of V4's tables, broken by the inserted rows (and in V4's own T21 and T23).
    for t in (ctx.T(20, "MQTT Topic"), ctx.T(21, "Field | Type | Description"),
              ctx.T(22, "Field | Type | Description"), t23, tbl):
        _reband(t, shd_tpl)
    ctx.note("s3: T20-T23 and the new data.reason table re-banded (odd body rows shaded F5F5F5, even rows plain).")
