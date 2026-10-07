# ruff: noqa: E501
"""V5 edits: 4.1-4.7 and 4.9-4.11 Functional Requirements: V4 P131-P171 and P174-P178, tables T6-T14 and T17-T19 (4.8 and 4.8.1 belong to tm).

Addresses V4's original numbering through ctx (see common.py). Every edit: ctx.note(...)."""
import copy

from docx.oxml.ns import qn
from docx.shared import Pt

from .common import (EditError, add_row, colour_priority, delete_para, delete_row,  # noqa: F401
                     insert_after, insert_lines_after, insert_table_after, replace, replace_in_cell,
                     set_cell, set_row_font_size, set_text, unique_cells, zebra)


def _req(ctx, k, r, rid, c2=None, c3=None, c2_expect=None, c3_expect=None):
    """Set C2 (Requirement) and/or C3 (Acceptance) of requirement row R<r> of T<k>.

    The row must carry requirement `rid` in C0 (checked against V4)."""
    ctx.cell(k, r, 0, rid)
    if c2 is not None:
        set_cell(ctx.cell(k, r, 2, c2_expect), c2)
    if c3 is not None:
        set_cell(ctx.cell(k, r, 3, c3_expect), c3)


def _drop_row(table, row):
    """Delete V4's original row `row` (a _Row) from `table`, wherever it now sits."""
    trs = [r._tr for r in table.rows]
    if row._tr not in trs:
        raise EditError("row to delete is not in the table")
    delete_row(table, trs.index(row._tr))


def apply(ctx):
    _onboarding(ctx)
    _dashboard(ctx)
    _leak_flow(ctx)
    _leak_reset_and_override(ctx)
    _leak_table(ctx)
    _remote_override(ctx)
    _valve_control(ctx)
    _sensors(ctx)
    _health(ctx)
    _history(ctx)
    _push(ctx)
    _device_mgmt(ctx)
    zebra(ctx.T(6, "Req ID"), ctx.body_shd)
    ctx.note("global fix INT-06: T6 (4.1 Account and Login) re-banded; V4's own banding was broken there")


# ---------------------------------------------------------------- 4.2 Hub Onboarding (T7)

def _onboarding(ctx):
    ctx.T(7, "Req ID")
    for r, old, new in ((1, "APP-FR-04", "APP-FR-004"), (2, "APP-FR-05", "APP-FR-005"),
                        (3, "APP-FR-06", "APP-FR-006"), (4, "APP-FR-07", "APP-FR-007"),
                        (5, "APP-FR-08", "APP-FR-008"), (6, "APP-FR-09", "APP-FR-009")):
        set_cell(ctx.cell(7, r, 0, old), new)
    ctx.note("s2-01: 4.2 requirement IDs APP-FR-04..09 written as APP-FR-004..009 (format fix; the trace range is s7's)")

    _req(ctx, 7, 3, "APP-FR-06",
         c2="The app shall validate the Gateway ID format (GW- followed by 12 hex characters: GW-[0-9A-Fa-f]{12}) and convert it to upper case before sending it to the backend. The hub's ID is always upper case, and IoT Hub device IDs are case-sensitive.",
         c3="Invalid format shows inline error. gw-34b7da6aad54 is sent as GW-34B7DA6AAD54.",
         c2_expect="validate the Gateway ID format", c3_expect="Invalid format shows inline error")
    ctx.note("s2-02: APP-FR-006 converts the Gateway ID to upper case before sending it")

    _req(ctx, 7, 4, "APP-FR-07",
         c2="If the backend cannot find the device in IoT Hub, the app shall show an error message explaining that the hub may not be online yet, and asking the user to try again in two minutes: \"Hub not found. Make sure the hub is powered on and connected to Wi-Fi. A hub that has just joined Wi-Fi can take up to two minutes to appear: try again then.\" A hub is in the IoT Hub registry only after its first DPS registration. Since 2.1.4 it reaches the cloud only after its Wi-Fi setup network has closed, 15 to 60 seconds after it joins the router (at most 75 s), and a first DPS registration can take up to a minute more (section 6.4).",
         c3="Error message is shown within 10 seconds, with the advice to try again in two minutes.",
         c2_expect="cannot find the device in IoT Hub", c3_expect="Error message is shown within 10 seconds")
    ctx.note("integrator (s6 cross-module, s6-06): APP-FR-007 message tells the user to try again in two minutes, as in 3.8 step 7 and 6.4 (canon C21)")

    _req(ctx, 7, 5, "APP-FR-08",
         c2="After successful hub linking, the app shall navigate to the dashboard. If the backend already holds a snapshot for this hub, the app shall show it at once. Otherwise it shall show a loading state until the first snapshot arrives. This can take up to the snapshot interval (default 5 minutes), or up to about 150 seconds after the hub restarts or reconnects. When the hub has no devices (data.valve has no valve_id and both ble_leak_sensors and lora_sensors are empty; system_health.reason reads \"No devices provisioned\"), the app shall show the empty state \"No devices set up yet\". A hub on firmware 2.1.3 with no devices sends no snapshot: the app shows the same empty state when twin reported has no valve_id (or valve_id null) and both sensor counts are 0.",
         c3="Dashboard shows \"Waiting for first update...\" then populates. A hub with no devices shows \"No devices set up yet\", on 2.1.4 and on 2.1.3.",
         c2_expect="After successful hub linking", c3_expect="Waiting for first update")
    ctx.note("s2-03: APP-FR-008 shows a stored snapshot at once, gives the first-snapshot timing and adds the \"No devices set up yet\" empty state (2.1.3: from twin reported device lists)")


# ---------------------------------------------------------------- 4.3 Dashboard (T8)

def _dashboard(ctx):
    ctx.T(8, "Req ID")
    _req(ctx, 8, 1, "APP-FR-021",
         c2="The dashboard shall show a status banner. The first match wins: \"Leak Detected\" (red) when any device reports a leak in the latest state (valve.leak_state, a disconnected valve's last known leak state (APP-FR-036), or leak_state in any entry of ble_leak_sensors or lora_sensors), also while the hub is offline, then with \"Hub offline - last known state\"; otherwise \"Hub Offline\" (grey) when the hub is offline (APP-FR-025); otherwise \"No devices set up yet\" when the hub has no devices (APP-FR-008); otherwise \"Needs Attention\" (orange) when system_health.rating is \"warning\" or \"critical\"; otherwise \"All Clear\" (green).",
         c3="Banner changes within 5 seconds of receiving a new event. A grey \"Hub Offline\" banner never hides an active leak.",
         c2_expect="status banner", c3_expect="Banner changes within 5 seconds")
    ctx.note("s2-04: APP-FR-021 banner in the canonical first-match order, with \"Needs Attention\" and \"No devices set up yet\"; a leak shows also while the hub is offline")

    _req(ctx, 8, 2, "APP-FR-022",
         c2="The dashboard shall show a valve card with: current state (open, closed, unknown or disconnected), battery %, and the RMLEAK lock icon. When valve.state is \"unknown\" (linked, readings not in yet), the app shall show \"Connecting...\" and ignore valve.rmleak and valve.leak_state. When the valve is disconnected, the snapshot leaves out battery, rmleak, leak_state and fw_version: the app shall show \"--\" for the battery and keep the last known lock state (APP-FR-038) and leak state (APP-FR-036). A null battery means unknown: show \"--\", never 0 %. A 2.1.3 hub sends battery 0 for a valve battery it has not read yet: treat 0 from a 2.1.3 hub as unknown. When data.valve has no valve_id, no valve is set up (\"valve\":{} since 2.1.4; {\"state\":\"disconnected\",\"connected\":false} on 2.1.3): the app shall show \"No valve set up\" in place of the valve card, with no valve controls. On a 2.1.3 hub it shall do the same while twin reported valve_id is null, even when data.valve carries a valve_id: such a hub can link to a nearby valve that is not its own (section 2.4).",
         c3="Valve card matches snapshot valve.state, valve.battery and valve.rmleak. A missing or null battery shows \"--\". A valve block without valve_id, or a 2.1.3 hub whose twin reported valve_id is null, shows \"No valve set up\" and no valve controls.",
         c2_expect="valve card with", c3_expect="Valve card matches snapshot")
    ctx.note("global fix INT-02/XC-05: APP-FR-022 shows \"No valve set up\" (not a hidden card) and adds the 2.1.3 guard: twin reported valve_id null means no valve even when data.valve has a valve_id (master app_ble_valve.c:1316 links to any valve by name; telemetry_v2.c:702-704)")
    ctx.note("s2-05: APP-FR-022 covers valve state \"unknown\", the fields left out while disconnected, null (and 2.1.3 zero) battery and the valve block without valve_id")

    _req(ctx, 8, 3, "APP-FR-023",
         c2="The dashboard shall show a sensor list with one row per provisioned sensor, from both ble_leak_sensors and lora_sensors (lora_sensors is empty on a hub without LoRa leak sensors; LoRa leak sensors are listed read-only), showing: location name (location.label; when the label is empty, location.code shown as a word, e.g. \"Kitchen\"; when the code is also \"unknown\", \"Unnamed sensor\"), leak status icon, battery icon, and connected/disconnected badge.",
         c2_expect="sensor list with one row per provisioned sensor")
    ctx.note("s2-06: APP-FR-023 lists lora_sensors too (read-only) and gives the location-name fallback")

    _req(ctx, 8, 5, "APP-FR-025",
         c2="The dashboard shall determine hub online/offline state primarily from the Azure IoT Hub device-connection state reported by the backend. The backend shall ignore a disconnect followed by a reconnect within 60 s: every hub closes and reopens its connection about every 18 hours to renew its SAS token. As a secondary gate, if the backend has received neither a snapshot nor a lifecycle message for longer than the larger of 2 x snapshot_interval_s and 5 minutes, the dashboard shall also treat the hub as offline. snapshot_interval_s is the value in force in twin reported (300 s when unknown), so the default is 10 minutes. When offline, show \"Hub Offline\" and grey out live data; an active leak still shows (APP-FR-021).",
         c3="Hub shown offline when IoT Hub reports the device disconnected, and also after ~10 min (default 2x 300s) without a snapshot or lifecycle message. Hub shown online only when both gates agree. On a 2.1.3 hub with no devices (it sends no snapshot and no lifecycle), only the IoT Hub connection state applies.",
         c2_expect="primarily from the Azure IoT Hub device-connection state", c3_expect="Hub shown offline when IoT Hub reports")
    ctx.note("s2-07: APP-FR-025 secondary gate is the larger of 2 x snapshot_interval_s (twin reported) and 5 minutes, restarted by a snapshot or a lifecycle; 60 s reconnect debounce; 2.1.3 empty hub uses the IoT Hub state only")


# ---------------------------------------------------------------- 4.4 Leak flow (P143-P153)

def _leak_flow(ctx):
    set_text(ctx.P(141, "Figure 8. Leak incident end-to-end flow"),
             "Figure 8. Leak incident end-to-end flow. From leak detection (a BLE or LoRa leak sensor, or the valve flood probe) through auto-close, the RMLEAK clear (automatic, or Leak Reset), and valve re-open. The alternate path shows the 24-hour override window.")
    ctx.note("R-s2-13: Figure 8 caption names the three leak sources and the automatic RMLEAK clear (figure to be redrawn)")

    set_text(ctx.P(143, "1. Sensor reports leak"),
             "1. A BLE or LoRa leak sensor, or the valve flood probe, reports water.")
    set_text(ctx.P(144, "2. Hub publishes leak_detected"),
             "2. The hub publishes a leak_detected event. data.source_type names the source: \"ble_leak_sensor\", \"lora\" or \"valve\".")
    set_text(ctx.P(145, "3. If auto-close is enabled"),
             "3. If auto_close_enabled is true, the source's bit is set in the trigger mask (section 4.8.1) and no 24-hour override window is active, the hub latches the incident, asserts RMLEAK and then closes the valve. If the valve is not connected, the hub holds both commands and sends them when the valve reconnects; if every leak source is dry by then, it drops them and the valve does not close. Otherwise the hub does not close the valve, and the leak is still reported. When only the override window stops the close, the hub sends auto_close_blocked_override instead (at most once a minute).")
    set_text(ctx.P(146, "4. Hub publishes auto_close"),
             "4. The hub publishes an auto_close event after the leak_detected event (data.rmleak_asserted is false when the valve was not connected). When the valve reports the close, the hub publishes valve_state_changed with valve_state \"closed\" and rmleak true. A further leak while the valve is already closed and locked sends no new auto_close. Since 2.1.4 a hub with no valve sends no auto_close; 2.1.3 sent one with rmleak_asserted false.")
    ctx.note("s2-16, tm-09: leak flow steps 1-4 name all three leak sources, the master switch, the trigger mask and the override gate, RMLEAK before the close, the held commands for an unreachable valve (dropped when every source dries first), auto_close_blocked_override only when the window alone stops the close (R-s2-03, R-s2-09), and no auto_close without a valve since 2.1.4")

    set_text(ctx.P(147, "5. Backend sends push"),
             "5. The backend sends the leak push notification (APP-FR-111) and, on auto_close, the valve push (APP-FR-112).")
    set_text(ctx.P(148, "6. App shows leak status"),
             "6. The app shows the leak. When the hub closed the valve (auto_close event, valve.rmleak true), it also shows the valve as LOCKED and the Leak Reset button, disabled while any device still reports leak_state true (the hub refuses a Leak Reset then). When auto-close is off, the source's trigger-mask bit is clear or an override window is active, the hub does not close the valve for this report and only the leak is shown; with auto-close on, a wet source whose bit is clear still closes the valve at the next valve reconnect, override_cancel or end of an override window (section 4.8.1).")
    ctx.note("global fix XC-15: step 6 says a wet source with its bit clear still closes the valve at a valve reconnect, override_cancel or override end (rules_engine.c:1586, :1376, :1863)")
    ctx.note("s2-v01, tm-v02: steps 5-6 point to the push requirements and show LOCKED and Leak Reset only when the hub closed the valve")

    set_text(ctx.P(150, "8. User taps Leak Reset"),
             "8. When every provisioned leak source has been dry for 10 seconds, the hub clears RMLEAK by itself and publishes rmleak_auto_cleared, in practice 10 to 12 seconds after the last source reports dry (up to 2.1.3 the wait was 30 seconds, so 30 to 60 seconds). The valve stays closed. The Leak Reset button then disappears without the user tapping it.")
    set_text(ctx.P(151, "9. Hub clears RMLEAK"),
             "9. The user can also tap Leak Reset once every leak source is dry. The app sends leak_reset; the hub clears RMLEAK and publishes rmleak_cleared.")
    ctx.note("s2-17: steps 8-9 add the automatic RMLEAK clear (10 s since 2.1.4, 30 s up to 2.1.3) before the manual Leak Reset")

    set_text(ctx.P(152, "10. User taps Open Valve"),
             "10. The user taps Open Valve. The app sends valve_open (or valve_set_state \"open\"). The hub queues the command and returns cmd_ack \"ok\", or an error (APP-FR-057).")
    p153 = ctx.P(153, "11. Hub opens valve")
    set_text(p153, "11. The valve opens and reports its new position. The hub publishes valve_state_changed with valve_state \"open\", and a snapshot follows. Water is restored.")
    ctx.note("s2-v02: steps 10-11 say the hub queues the open and the valve confirms it")

    insert_after(p153, "The hub handles a leak by itself, with or without an internet connection. While it cannot reach the cloud it still closes the valve when auto-close is on, keeps up to 16 events and sends them after it reconnects, oldest first, each with ts set to when it happened (section 5.5). The backend and the app shall order events by ts and expect leak events and their pushes to arrive late (APP-NF-014, APP-FR-111). Since 2.1.4 this also holds when the hub restarts while the router is down, and events raised before the hub's clock is set after a power cut are kept and sent with their real time. A 2.1.3 hub that restarts while Wi-Fi is down does not act on leaks until it is back on Wi-Fi and its cloud start-up has run, and it discards events raised before its clock is set.")
    ctx.note("s2-v04: new paragraph after the leak flow: the hub acts on leaks without the cloud and sends events late; order by ts (APP-NF-014); 2.1.3 differences")


# ---------------------------------------------------------------- 4.4.1 / 4.4.2 (P155, P157, P159)

def _leak_reset_and_override(ctx):
    set_text(ctx.P(155, "Leak Reset clears the latch and clears RMLEAK"),
             "Leak Reset clears the latch and clears RMLEAK. It does NOT open the valve. The app shall show Leak Reset as step one, then Open Valve as step two. This prevents accidental flooding if the leak is not actually fixed. Leak Reset is refused while any provisioned leak source is still wet (a BLE or LoRa leak sensor, or the valve flood probe), including a source whose trigger-mask bit is clear, whatever the auto-close setting. The hub then returns the cmd_ack error \"A leak is still active. Fix the leak first, or use override to open the valve during a leak.\" This prevents clearing the interlock and re-opening onto a live leak. To open the valve during an active leak, use the 24-hour override. A Leak Reset also ends an active 24-hour override window; its rmleak_cleared event then carries override_cancelled true. A Leak Reset with nothing to clear succeeds (status \"ok\") and sends no event. The hub also clears RMLEAK by itself about 10 seconds after every leak source is dry (rmleak_auto_cleared; 30 to 60 seconds up to 2.1.3), so a Leak Reset is often not needed.")
    ctx.note("s2-18, tm-10: 4.4.1 states the refusal for any wet provisioned leak source (mask bit clear included), the exact error text, the override end with override_cancelled, the no-op ack and the automatic clear")

    p157 = ctx.P(157, "If the user physically overrides the valve")
    set_text(p157, "If the user presses the override button on the valve after an auto-close, the valve clears RMLEAK and opens, and the hub starts a 24-hour override window. During this window, leaks are still detected and reported, but auto-close is blocked: a leak that would have closed the valve sends auto_close_blocked_override (at most once a minute) instead, and the valve stays open. The snapshot fields override_active, override_remaining_s and expires_ts tell the app when this is happening. Since 2.1.4 expires_ts is left out when the window started before the hub's clock was set (snapshots then carry it from up to about 30 s after the clock sync): count down from override_remaining_s then (APP-FR-041). The window survives a hub restart. The same 24-hour window can also be started remotely from the app via the override_enable command (see section 4.4.3); the only difference is the event's trigger field (\"button\" vs \"c2d_command\"). The hub also reports trigger \"button\" when the valve reconnects open with RMLEAK clear while the hub still has the leak latched: it infers a press made while it could not see the valve. In rare cases this inference reports a press that did not happen.")
    insert_lines_after(p157, [
        "The window ends in one of these ways: it expires (water_access_override_expired); the app sends override_cancel (auto_close_reenabled); the app sends leak_reset once every leak source is dry (rmleak_cleared with override_cancelled true); decommission target \"all\" resets the hub; or, since 2.1.4, the last device is removed from the hub (no event; the next snapshot shows override_active false).",
        "When the window ends by expiry or by override_cancel while any provisioned leak source is still wet and auto_close_enabled is true, the hub asserts RMLEAK and closes the valve at once, even when that source's trigger-mask bit is clear (section 4.8.1). No auto_close event is sent for this close: the app sees water_access_override_expired (auto_close_resumed true) or auto_close_reenabled, then valve_state_changed. Hub firmware 2.1.3 and 2.1.4 both behave this way.",
    ])
    ctx.note("global fix INT-15: 4.4.2, APP-FR-041 and APP-FR-045 say when expires_ts is missing in one wording: since 2.1.4, a window started before the clock was set, up to about 30 s after the sync in snapshots (rules_engine.c:2066-2075)")
    ctx.note("s2-19, tm-43: 4.4.2 adds auto_close_blocked_override, expires_ts, persistence and the inferred button trigger, then two new paragraphs: every way the window ends, and the immediate close (no auto_close, mask not checked) when it ends during a leak")

    set_text(ctx.P(159, "Figure 9. 24-hour override window flow"),
             "Figure 9. 24-hour override window flow. Shows the trigger (valve button or override_enable), the blocked auto-close behavior, and how the window ends: expiry, override_cancel, or leak_reset once every leak source is dry. If a leak source is still wet when the window ends by expiry or override_cancel and auto-close is on, the hub closes the valve at once.")
    ctx.note("tm-v05: Figure 9 caption names both triggers, the three user-visible ends and the immediate close (figure to be redrawn)")


# ---------------------------------------------------------------- 4.4 requirements (T9)

def _leak_table(ctx):
    ctx.T(9, "Req ID")
    _req(ctx, 9, 1, "APP-FR-036",
         c2="When a leak_detected event arrives (data.source_type \"ble_leak_sensor\", \"lora\" or \"valve\"), the app shall show a prominent leak banner on the dashboard and valve screen. An event whose ts is older than the latest snapshot's ts goes into history only; it does not change the live banner (APP-NF-014). The banner clears when a snapshot shows no device with leak_state true. A valve entry with no leak_state (valve disconnected), or with state \"unknown\", keeps the valve's last known leak state: a valve flood probe that was wet when the link dropped still counts as a leak until leak_cleared with source_type \"valve\" arrives or a snapshot shows valve.leak_state false with valve.state \"open\" or \"closed\". system_health.reason then names both the leak and \"Valve offline\" (for example \"Leak detected: valve, Valve offline\"), and no device_offline is sent for the valve.",
         c3="Red banner appears within 5 seconds of event receipt. A replayed old leak_detected does not re-raise the banner. A disconnected valve whose flood probe was wet keeps the banner.",
         c2_expect="When a leak_detected event arrives", c3_expect="Red banner appears")
    ctx.note("global fix completeness-03: APP-FR-036 keeps a disconnected valve's last known leak state in the banner; the hub keeps such a valve critical as a leak and sends no device_offline (app_ble_valve.c:1902-1923, health_engine.c:383-387, telemetry_v2.c:620-635; same on 2.1.3)")
    ctx.note("s2-25: APP-FR-036 names the three sources, ignores stale replays for the live banner and clears from the snapshot")

    _req(ctx, 9, 2, "APP-FR-037",
         c2="When an auto_close event arrives, the app shall show the valve as \"Closing (Locked)\" with the RMLEAK lock icon. It shall show \"Closed (Locked)\" when valve_state_changed reports valve_state \"closed\" with rmleak true, or when a later snapshot shows it. When data.rmleak_asserted is false the valve was out of reach: show \"Will close when the valve reconnects\". If the leak dries before the valve reconnects, the hub drops the held close: take the valve state from the next snapshot. An auto_close with data.cause \"reconnect\" means the valve reconnected during a leak.",
         c3="Valve card shows \"Closed (Locked)\" after the valve's own report, not after auto_close alone.",
         c2_expect="When an auto_close event arrives", c3_expect="Closed (Locked)")
    ctx.note("s2-20: APP-FR-037 shows \"Closing (Locked)\" on auto_close and confirms from valve_state_changed or the snapshot; rmleak_asserted false (the held close is dropped if the leak dries first, R-s2-03) and cause \"reconnect\"")

    _req(ctx, 9, 3, "APP-FR-038",
         c2="The app shall show a Leak Reset button when valve.rmleak is true in the snapshot. While valve.state is \"unknown\" or the valve is disconnected (rmleak absent), the app shall keep the last known value. On rmleak_cleared or rmleak_auto_cleared it shall hide the button. While any device reports leak_state true, the button shall be disabled with the note \"Dry the leak first\", because the hub refuses a Leak Reset then.",
         c3="Button visible when rmleak=true, hidden when rmleak=false, unchanged while the state is unknown or the valve is disconnected. Disabled while any leak_state is true.",
         c2_expect="Leak Reset button when valve.rmleak is true", c3_expect="Button visible when rmleak=true")
    ctx.note("s2-21: APP-FR-038 keeps the last rmleak through unknown/disconnected, hides on rmleak_cleared or rmleak_auto_cleared, disables while a leak is active")

    _req(ctx, 9, 4, "APP-FR-039",
         c2="After a leak_reset cmd_ack with status \"ok\", the app shall hide the RMLEAK lock icon and enable the Open Valve button. The valve confirms the clear a few seconds later (when the valve is out of reach, at its reconnect), so a snapshot received just after the ack can still show valve.rmleak true. If a snapshot received 15 seconds or more after the ack still shows valve.rmleak true while the valve is connected, the app shall show the lock icon again.",
         c3="RMLEAK icon disappears. Open Valve becomes tappable. Since 2.1.4 a snapshot within 15 seconds shows valve.rmleak false; a 2.1.3 hub shows it only in its next periodic snapshot, and the icon stays hidden until then.",
         c2_expect="After a successful leak_reset cmd_ack", c3_expect="RMLEAK icon disappears")
    ctx.note("s2-22: APP-FR-039 reconciles the lock icon only from snapshots 15 s or more after the ack; acceptance gives the 2.1.3 timing (R-s2-04)")

    _req(ctx, 9, 5, "APP-FR-040",
         c2="When override_active is true in the snapshot, the app shall show a warning banner: \"24h override active. Automatic shutoff is paused. [Cancel Override]\". While any device reports leak_state true and data.rules.auto_close_enabled is true, the Cancel Override confirmation shall warn: \"A leak is still active. Cancelling the override will close the valve now.\"",
         c3="Warning banner visible. Cancel Override button sends override_cancel. The leak warning shows in the confirmation while a leak is active.",
         c2_expect="When override_active is true in the snapshot", c3_expect="Warning banner visible")
    ctx.note("tm-11, s2-23: APP-FR-040 banner reads \"Automatic shutoff is paused\"; the cancel confirmation warns that the valve closes while a leak is active")

    _req(ctx, 9, 6, "APP-FR-041",
         c2="The app shall show the time left in the override as a human-readable countdown (e.g., \"23h 45m remaining\"). It shall count down locally from expires_ts, or from the snapshot ts plus override_remaining_s when expires_ts is absent (since 2.1.4, when the window started before the hub's clock was set), and correct it at each snapshot.",
         c3="Countdown decreases every minute and matches the snapshot within a minute.",
         c2_expect="override_remaining_s as a human-readable countdown", c3_expect="Countdown updates with each snapshot")
    ctx.note("s2-24: APP-FR-041 counts down locally from expires_ts (or ts + override_remaining_s)")

    _req(ctx, 9, 7, "APP-FR-042",
         c2="After a successful override_cancel cmd_ack, the override banner shall disappear. The hub sends auto_close_reenabled when a window was active (no event when there was nothing to cancel). If any leak source is still wet and auto-close is enabled, the hub asserts RMLEAK and closes the valve at once, even when that source's trigger-mask bit is clear (section 4.8.1).",
         c3="Banner gone after next snapshot shows override_active=false. With a leak active and auto-close on, the valve card then shows \"Closed (Locked)\".",
         c2_expect="After a successful override_cancel cmd_ack", c3_expect="Banner gone after next snapshot")
    ctx.note("s2-23: APP-FR-042 adds auto_close_reenabled and the immediate close when cancelling during a leak")


# ---------------------------------------------------------------- 4.4.3 (P161, T10)

def _remote_override(ctx):
    p161 = ctx.P(161, "The 24-hour override window can also be started from the app")
    replace(p161,
            "Available on hubs running firmware 1.4.0 or later; the backend should offer it only when gateway.fw is 1.4.0 or higher.",
            "Every hub this document covers (2.1.3 and 2.1.4) supports it; no firmware check is needed. The hub can take up to about 10 seconds to ack it while it reconnects to the valve, which is within the 20-second command timeout (APP-FR-053).")
    ctx.note("s2-26: 4.4.3 drops the gateway.fw 1.4.0 gate and gives the 10 s ack delay")

    insert_after(p161, "The hub refuses override_enable with an error.detail to show verbatim: \"No active leak to override. Use the normal Open Valve control.\" when there is nothing to override; \"Water detected at the valve. It can't be opened remotely until the valve area is dry.\" when the valve flood probe is wet; \"The valve isn't responding. Check its power and connection, then try again.\" when the valve does not link within about 10 seconds; \"No valve is set up for this hub.\" when no valve is provisioned; and \"Something went wrong applying the override. Your water state is unchanged. Try again.\" on an internal error. On success the hub clears the valve's RMLEAK and opens the valve. A second override_enable during an active window restarts it at 24 hours.")
    ctx.note("s2-v03: new paragraph after 4.4.3 lists every override_enable refusal text and the success effect")

    ctx.T(10, "Req ID")
    _req(ctx, 10, 1, "APP-FR-043",
         c2="When valve.rmleak is true, in addition to the disabled \"Open Valve\" button (APP-FR-057), the app shall present a distinct \"Open with 24h Override\" action that sends the override_enable command. While valve.leak_state is true the action shall be disabled with the note \"Water detected at the valve\", because the hub refuses it then. While the valve's last known battery is 10 % or less (a 0 from a 2.1.3 hub is unknown and does not count), the action shall also be disabled, with the note \"Valve battery critical. Replace the batteries.\": the hub does not refuse override_enable for battery, so it would start the 24-hour window and pause automatic shutoff while the valve stays shut.",
         c3="Action visible only while valve.rmleak=true; disabled while valve.leak_state=true or the valve battery is 10 % or less; sends override_enable with a correlation id.",
         c2_expect="Open with 24h Override", c3_expect="Action visible only while valve.rmleak=true")
    ctx.note("s2-27: APP-FR-043 disables the override action while the valve flood probe is wet")
    ctx.note("global fix completeness-08: APP-FR-043 also disables the override at a valve battery of 10 % or less (rules_engine.c:1440-1515 has no battery check; the valve refuses to open at Critical)")

    ctx.cell(10, 3, 0, "APP-FR-045")
    replace_in_cell(ctx.cell(10, 3, 2, "the countdown by remaining_s / override_remaining_s"),
                    "and the countdown by remaining_s / override_remaining_s.",
                    "and the countdown by expires_ts when present, else by remaining_s / override_remaining_s (since 2.1.4 expires_ts is left out when the window started before the hub's clock was set).")
    ctx.note("s2-28: APP-FR-045 countdown from expires_ts, falling back to remaining_s")

    _house_style_t10(ctx)


def _house_style_t10(ctx):
    """V4's T10 (APP-FR-043..045) lacks the house style every other requirement table has. Copy T9's
    header (navy 1A237E fill; bold white 8 pt runs), give every body run 8 pt, band the body rows and
    colour the priorities."""
    hdr = ctx.row(9, 0, "Req ID")
    hdr_tc = hdr._tr.findall(qn("w:tc"))[0]
    hdr_shd = hdr_tc.find(qn("w:tcPr") + "/" + qn("w:shd"))
    hdr_rpr = hdr_tc.find(".//" + qn("w:r") + "/" + qn("w:rPr"))
    if hdr_shd is None or hdr_shd.get(qn("w:fill")) != "1A237E" or hdr_rpr is None:
        raise EditError("T9 R0 does not carry the navy header style to copy")
    t10 = ctx.T(10, "Req ID")
    for tc in t10.rows[0]._tr.findall(qn("w:tc")):
        tcpr = tc.get_or_add_tcPr()
        for old in tcpr.findall(qn("w:shd")):
            tcpr.remove(old)
        tcpr.append(copy.deepcopy(hdr_shd))
        for r in tc.iter(qn("w:r")):
            for old in r.findall(qn("w:rPr")):
                r.remove(old)
            r.insert(0, copy.deepcopy(hdr_rpr))
    for row in t10.rows[1:]:
        for c in unique_cells(row):
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
        colour_priority(row)
    zebra(t10, ctx.body_shd)
    ctx.note("global fix INT-19: T10 (APP-FR-043..045) gets the house style of the other requirement tables: navy bold white header, 8 pt body, banded rows, P0 in red")


# ---------------------------------------------------------------- 4.5 Valve Control (T11)

def _valve_control(ctx):
    set_text(ctx.P(165, "Figure 10. C2D command lifecycle"),
             "Figure 10. C2D command lifecycle (valve_close example). Shows the round trip from user tap through backend and IoT Hub to the hub, which queues the command and returns cmd_ack, then the valve_state_changed event, sent when the valve reports its new position, that confirms the new state in the UI. The drawing still resolves the command on cmd_ack; the card resolves on valve_state_changed or a snapshot (APP-FR-055).")
    ctx.note("R-s2-13: Figure 10 caption: the hub queues the command and acks it; valve_state_changed confirms the state (figure to be redrawn)")

    ctx.T(11, "Req ID")
    _req(ctx, 11, 2, "APP-FR-052",
         c2="After sending a command, the app shall show a pending/spinner state on the valve card until cmd_ack arrives or the 20-second timeout passes; after an \"ok\" for a valve command, APP-FR-055 applies.",
         c2_expect="until cmd_ack arrives or timeout occurs")
    _req(ctx, 11, 8, "APP-FR-058",
         c2="The app shall not allow sending a second command while a previous command for the same hub is pending. A valve command is pending until its cmd_ack and, after an \"ok\", until the new state is confirmed or 60 s pass (APP-FR-055). A settings change waiting for its confirming snapshot (APP-FR-094) does not block valve commands.",
         c2_expect="second command while a previous command")
    ctx.note("global fix INT-13/XC-09: APP-FR-052 ends the spinner at the ack or the 20 s timeout and hands an ok valve command to APP-FR-055; APP-FR-058 defines pending (valve command until confirmed or 60 s; a settings change does not block valve commands), as 6.5 step 10 does")
    ctx.cell(11, 1, 0, "APP-FR-051")
    replace_in_cell(ctx.cell(11, 1, 2, "valve_open or valve_close command"),
                    "any valve_open or valve_close command", "any valve_open, valve_close or valve_set_state command")
    ctx.note("s2-29: APP-FR-051 confirmation also covers valve_set_state")

    _req(ctx, 11, 3, "APP-FR-053",
         c2="The app shall wait up to 20 seconds for cmd_ack. If no ack, show: \"Command timed out. The hub may be offline.\" override_enable can take up to about 10 seconds to ack while the hub reconnects to the valve. A cmd_ack that arrives after the timeout shall not change the UI; the next snapshot shows the real state (APP-NF-018). A missing ack means the result is unknown, not that the command failed.",
         c3="After 20s without ack, timeout message shown. UI reverts to previous state. A late ack does not change the UI.",
         c2_expect="wait up to 20 seconds for cmd_ack", c3_expect="After 20s without ack")
    ctx.note("s2-30: APP-FR-053 gives the override_enable ack delay and ignores a late ack in the UI (APP-NF-018)")

    _req(ctx, 11, 4, "APP-FR-054",
         c2="If cmd_ack.status is \"error\", the app shall keep the UI state unchanged and show the error. Section 5.2.7 lists every error.detail text and marks each Show or Log. When error.detail is a text marked Show (the sentences written for the user by valve_open, valve_close, valve_set_state, leak_reset and override_enable), the app shall show it verbatim. Since 2.1.4 the valve commands can also return \"No valve is set up for this hub.\", \"Valve battery critical (≤10 %): the valve will not open. Replace the batteries.\" and \"The valve command could not be queued. Try again.\" (2.1.3 acked those cases \"ok\"). The two valve_set_state texts for a missing or invalid state (\"missing 'state' field ...\" and \"invalid state value ...\") come only from an app defect: show \"Something went wrong. Try again.\" and log error.detail. For any other error.detail, show \"The hub could not apply this change. Try again.\" and log error.detail. This includes the message-size texts (\"exceeds the reassembly limit: ...\" and \"out of memory: ...\"), which a valve command can also return. Do not branch on error.code: it holds the command name.",
         c3="Texts marked Show in section 5.2.7 appear verbatim in a toast or banner. Any other error shows the generic message, and error.detail is logged.",
         c2_expect="If cmd_ack.status is \"error\"", c3_expect="Error detail from cmd_ack.error.detail")
    ctx.note("s2-31: APP-FR-054 decides by the error.detail text (Show or Log in 5.2.7), not by command: Show texts verbatim, Log texts behind a generic message; error.code is the command name")

    _req(ctx, 11, 5, "APP-FR-055",
         c2="If cmd_ack.status is \"ok\", the app shall show the valve card as \"Opening...\" or \"Closing...\", or \"Waiting for valve\" while valve.connected is false. An \"ok\" means the hub queued the command, not that the valve moved. While the valve is disconnected the hub holds the command and sends it when the valve links again, with no time limit (a newer valve command replaces it; a hub restart drops it). The app shall show the new state when valve_state_changed reports it or a snapshot shows it. A snapshot sent just after the ack can still show the old state: the app shall not revert the card on it. If neither confirms within 60 seconds of the ack, the app shall show the state from the latest snapshot with the message \"The valve has not confirmed yet. Check that it has power and is in range.\" and keep reading snapshots. The command stays pending, with the valve buttons disabled (APP-FR-058), until the new state is confirmed or the 60 seconds pass.",
         c3="Valve card shows the new state after valve_state_changed, normally within a few seconds of the ack. Without confirmation, the message shows 60 seconds after the ack. Valve buttons stay disabled until then.",
         c2_expect="If cmd_ack.status is \"ok\"", c3_expect="Valve card shows new state immediately on ack")
    ctx.note("s2-32: APP-FR-055 shows a pending state on \"ok\" and confirms from valve_state_changed or a snapshot, with the 60 s rule (canon valve confirmation)")

    _req(ctx, 11, 6, "APP-FR-056",
         c2="The app shall always include a correlation id (\"id\", a JSON string of at most 63 characters; a UUID fits) in every command. The hub acks every eflostop.cmd command and echoes the id, so the backend can match the ack; a longer id comes back cut and will not match. The hub does not de-duplicate commands by id: a command delivered twice runs twice and is acked twice, and the first ack for an id resolves it. Every send, retries included, uses a new id, and the app does not resend automatically after a timeout.",
         c3="All commands include an \"id\" string. Backend matches ack by this id.",
         c2_expect="always include a correlation id", c3_expect="All commands include \"id\" field")
    ctx.note("s2-33: APP-FR-056 gives the id type and 63-character limit, and says the hub does not de-duplicate by id")

    _req(ctx, 11, 7, "APP-FR-057",
         c2="The Open Valve button shall be disabled when valve.rmleak is true. The user must tap Leak Reset first, or wait for the hub to clear RMLEAK about 10 seconds after every leak source is dry (30 to 60 seconds up to 2.1.3). The hub enforces this server-side. It rejects a valve_open or valve_set_state \"open\" with the cmd_ack error \"Valve is locked after a leak (RMLEAK). Clear it with leak_reset first, or use override to open the valve during a leak.\" while the valve's RMLEAK is set and, since 2.1.4, also while the hub has a leak incident latched and no override window is active, even when the valve is disconnected. So the valve never briefly opens, even if a client bypasses the disabled button. Since 2.1.4 the hub also refuses an open when the valve's last battery reading is 10 % or less (\"Valve battery critical (≤10 %): the valve will not open. Replace the batteries.\"; closing is never refused for battery), and refuses any valve command when no valve is provisioned (\"No valve is set up for this hub.\") or when its valve command queue is full (\"The valve command could not be queued. Try again.\"). The checks run in this order: no valve, locked, battery, queue. A 2.1.3 hub checks only the valve's RMLEAK and acks the other cases \"ok\".",
         c3="Open Valve greyed out with tooltip: \"Clear leak first\". A C2D valve_open while rmleak=true returns the RMLEAK cmd_ack error and the valve does not move.",
         c2_expect="The Open Valve button shall be disabled when valve.rmleak is true", c3_expect="Open Valve greyed out")
    ctx.note("s2-34: APP-FR-057 adds the 2.1.4 refusals (latched incident, battery 10 % or less, no valve, queue full), their order and the 2.1.3 behaviour")


# ---------------------------------------------------------------- 4.6 Sensors (P167, T12)

def _sensors(ctx):
    set_text(ctx.P(167, "Sensors need user-friendly names and locations"),
             "Sensors need user-friendly names and locations. The hub stores sensor_meta (a location code and a label) for BLE leak sensors and LoRa leak sensors. The app sends a sensor_meta C2D command to update them for BLE leak sensors; it shows LoRa leak sensors read-only. The valve has no metadata: its leak events always carry location code \"unknown\" with an empty label.")
    ctx.T(12, "Req ID")
    _req(ctx, 12, 1, "APP-FR-066",
         c2="The app shall display all provisioned sensors from the snapshot: the ble_leak_sensors array and the lora_sensors array. The hub always sends both; lora_sensors is empty on a hub without LoRa leak sensors. LoRa leak sensors are shown read-only.",
         c2_expect="ble_leak_sensors array")
    ctx.note("s2-35: 4.6 intro and APP-FR-066 cover lora_sensors (shown read-only) and the valve's fixed location")

    _req(ctx, 12, 2, "APP-FR-067",
         c2="For each sensor, the app shall show: sensor_id, location.code, location.label, connected, rating, battery, leak_state, rssi, last_seen_age_s, and fw_version (BLE leak sensors) or snr (LoRa leak sensors).",
         c3="All fields populated from snapshot. Null values shown as \"--\" (a null last_seen_age_s as in APP-FR-070).",
         c2_expect="device_id, location.code", c3_expect="Null values shown as")
    ctx.note("s2-36: APP-FR-067 identity key is sensor_id (no device_id); adds fw_version and snr")

    _req(ctx, 12, 3, "APP-FR-068",
         c2="The app shall allow the user to edit a BLE leak sensor's label and location code by sending a sensor_meta C2D command (sensor_type \"ble_leak_sensor\", sensor_id, location_code, label). The label is at most 31 bytes of UTF-8. The hub cuts a longer label silently and still acks ok, so the app shall enforce the limit and never cut a character in two. An unknown location_code is ignored (the old code is kept) with an ok ack, so the app shall send only the codes in APP-FR-069.",
         c3="Edit saves successfully. The snapshot that follows the ack (normally within 5 seconds) shows the updated location.",
         c2_expect="edit the sensor label (max 31 characters)", c3_expect="Edit saves successfully")
    ctx.note("s2-37: APP-FR-068 gives the sensor_meta keys, the 31-byte UTF-8 label limit and the silently ignored unknown code (BLE leak sensors only: the MVP app does not edit LoRa sensors)")

    _req(ctx, 12, 5, "APP-FR-070",
         c2="last_seen_age_s shall be shown as a human-readable relative time (e.g., \"2 min ago\", \"1 hour ago\"). Null means the hub has not heard the sensor since the hub last restarted or since the sensor was added: show \"Not heard yet\".",
         c2_expect="Null means \"Never seen\"")
    ctx.note("s2-38: APP-FR-070 null last_seen_age_s reads \"Not heard yet\" (APP-FR-067 acceptance points to it)")


# ---------------------------------------------------------------- 4.7 Health (P169, T13, T14)

TIMEOUT_ROWS = [
    ["BLE leak sensor: no advertisement",
     "10 minutes (checked every 30 s, so 10 to 10.5 minutes)",
     "Sensor rating goes to critical (\"N sensor(s) offline\"). The hub sends device_offline if it had heard the sensor since it started or since the sensor was added, and the sensor was dry when it went silent (a wet sensor stays critical as a leak and gets none)."],
    ["LoRa leak sensor: no packet",
     "10 minutes (checked every 30 s, so 10 to 10.5 minutes)",
     "Same as a BLE leak sensor."],
    ["Valve: BLE link lost",
     "3 minutes",
     "warning (\"Valve disconnected\") for 3 minutes, then critical (\"Valve offline\") and device_offline, if the hub had heard the valve since it started or since the valve was added. A valve whose flood probe was wet when the link dropped stays critical as a leak (\"Leak detected: valve, Valve offline\") and gets no device_offline."],
    ["Device not heard since the hub restarted or since it was added",
     "10 minutes (sensor); 3 minutes (valve; 2.5 minutes when a provision added it)",
     "Its own rating reads critical, with connected false and last_seen_age_s null, but the hub leaves it out of system_health.rating until then: system_health.reason reads \"Syncing - waiting for N device(s)\" or ends with \"syncing N device(s)\". After that it counts as offline. No device_offline is sent for it. The first snapshot after a restart comes once the valve link is ready or every device is heard, at the latest about 150 seconds after the hub connects. Since 2.1.4 a provision or decommission leaves the other devices as they were; up to 2.1.3 it returned every remaining device to this state."],
    ["Any device reports a leak",
     "Immediate",
     "critical. No health event: leak_detected reports it."],
    ["Sensor battery at or below 20 %",
     "Immediate",
     "warning (\"N sensor(s) battery low\")."],
    ["Sensor battery 21-35 %",
     "Immediate",
     "good (if no other issues)."],
    ["Valve battery 11-20 % (since 2.1.4)",
     "Immediate",
     "warning (\"Valve battery low\"). A 2.1.3 hub rates the valve battery like a sensor's: warning at 20 % or less, good at 21-35 %."],
    ["Valve battery at or below 10 % (since 2.1.4)",
     "Immediate",
     "critical (\"Valve battery critical\"). No health event; a snapshot follows within seconds. The hub refuses to open the valve (APP-FR-057). A 2.1.3 hub rates it warning. On any hub version the valve itself closes at 10 % or less and will not open: valve_state_changed \"closed\" arrives with no command and no auto_close."],
    ["Sensor RSSI at or below -90 dBm",
     "Immediate",
     "warning (\"N sensor(s) signal weak\")."],
    ["Sensor RSSI from -89 to -80 dBm",
     "Immediate",
     "good (if no other issues). The valve has no RSSI band."],
    ["Leak incident latched",
     "Until the hub clears the incident (RMLEAK cleared by the hub or by Leak Reset, or an override starts)",
     "Hub rating at least warning (\"Leak interlock latched\")."],
]


def _health(ctx):
    set_text(ctx.P(169, "The hub produces a health rating per device"),
             "The hub rates each device excellent, good, warning or critical. system_health.rating is the worst device rating, with three exceptions. First, a device that has not been heard since the hub restarted, or since it was added, is left out for its first 10 minutes (a sensor) or 3 minutes (the valve; 2.5 minutes after a provision); the reason then says the hub is syncing. Second, while a leak incident is latched, the overall rating is at least warning, with the reason \"Leak interlock latched\", even when every device is healthy. Third, since 2.1.4 a hub with no devices reads excellent with the reason \"No devices provisioned\" (a 2.1.3 hub with no devices sends no snapshot). A device's own rating can therefore be critical while system_health.rating is excellent. The app shows system_health.rating as sent and does not compute it from the device list.")
    ctx.note("s2-08: 4.7 intro gives the three exceptions to \"worst device rating\" (not-heard exclusion, interlock floor, empty hub since 2.1.4)")

    ctx.T(13, "Rating")
    for r, name, expect, text in (
        (1, "excellent", "Everything looks normal", "Everything looks normal. The device is connected, and its battery and signal are fine."),
        (2, "good", "Small issues", "Small issue on a sensor: battery 21-35 %, or signal from -89 to -80 dBm. Since 2.1.4 the valve has no good band (a 2.1.3 hub rates a valve battery of 21-35 % good)."),
        (3, "warning", "battery below 20% or signal below -90 dBm", "Needs attention. A sensor battery at 20 % or less, a sensor signal at -90 dBm or weaker, a valve battery at 11-20 % (2.1.3: 20 % or less), or a valve link lost less than 3 minutes ago. The hub rating is also at least warning while a leak incident is latched."),
        (4, "critical", "device offline past its timeout", "Not safe to ignore. A leak at the device, a device offline past its timeout (sensor 10 minutes, valve 3 minutes), a device not heard since the hub restarted or since it was added (APP-FR-086), or, since 2.1.4, a valve battery at 10 % or less (the valve will not open)."),
    ):
        ctx.cell(13, r, 0, name)
        set_cell(ctx.cell(13, r, 2, expect), text)
    ctx.note("s2-09: rating legend: inclusive bands, the valve's own battery bands since 2.1.4, leak and unheard devices critical, the interlock floor")

    # s2-11 + s2-10: split V4's merged T14 into a timeout table and a requirements table.
    t14 = ctx.T(14, "Condition | Timeout | Result")
    p171 = ctx.P(171, "Known timeout rules used by the hub health engine")
    # Body rows alternate fill like every V4 table: shaded (V4 R1) then plain (V4 R2). The
    # template rows are still in t14 here; the _drop_row loop below runs after.
    new_tbl = insert_table_after(p171, t14, ["Condition", "Timeout", "Result"], TIMEOUT_ROWS[:1])
    shaded = ctx.row(14, 1, "BLE leak sensor: no advertisement")
    plain = ctx.row(14, 2, "Valve: BLE disconnected")
    prev = new_tbl.rows[1]
    for i, vals in enumerate(TIMEOUT_ROWS[1:], start=2):
        prev = add_row(new_tbl, vals, like_row=plain if i % 2 == 0 else shaded, after_row=prev)
    insert_after(new_tbl, "Health and battery requirements:", like=p171)
    for r, expect in ((0, "Condition"), (1, "BLE leak sensor: no advertisement"), (2, "Valve: BLE disconnected"),
                      (3, "Boot sync window"), (4, "Battery at or below 20%"), (5, "Battery at or below 35%"),
                      (6, "RSSI at or below -90 dBm"), (7, "RSSI at or below -80 dBm")):
        _drop_row(t14, ctx.row(14, r, expect))
    ctx.note("s2-11, s2-10: V4's merged table split into a Condition | Timeout | Result table and the APP-FR-081..086 requirements table; timeout rules rewritten (3-minute boot sync, LoRa timeout, unheard devices, leak, valve battery bands since 2.1.4, sensor-only RSSI, interlock floor)")

    _req(ctx, 14, 10, "APP-FR-082",
         c3="Reason text matches snapshot exactly. It is a comma-separated, human-readable list, for example \"All devices healthy\", \"No devices provisioned\", \"Syncing - waiting for 2 devices\", \"Leak detected: Kitchen\", \"2 leaks detected\", \"Leak interlock latched\", \"Valve offline\", \"Valve disconnected\", \"Valve battery critical\", \"Valve battery low\", \"1 sensor offline\", \"2 sensors battery low\", \"1 sensor signal weak\", \"syncing 1 device\" (\"No devices provisioned\" and \"Valve battery critical\" since 2.1.4). The app shall display it and shall not parse it.",
         c3_expect="Reason text matches snapshot exactly")
    ctx.note("s2-15: APP-FR-082 acceptance gives example reasons and forbids parsing them")

    _req(ctx, 14, 12, "APP-FR-084",
         c2="When a health event arrives (data.category \"health\") with data.event \"device_offline\", the backend shall send a push notification (APP-FR-113). The hub sends health events for reachability only: device_offline when a device it has heard since it started (or since it was added) loses its link past its timeout while dry, and device_recovered when that device is heard again. A leak, a battery at any level, a weak signal or a device never heard never produces a health event; the snapshot's ratings and system_health carry them. device_recovered needs no push. Its rating can still be critical (the device came back wet or, since 2.1.4, is a valve at 10 % battery or less), so the app takes the device's state from the snapshot, not from the event. The hub sends at most one health event per device per 60 s. Since 2.1.4 a change inside those 60 s is sent when they end; 2.1.3 dropped it, so a device_recovered due within 60 s of its device_offline never came, and only the next snapshot shows the device connected. There is no health_alert event and no new_rating field.",
         c3="Push received for each device_offline. No push for warning ratings or for device_recovered.",
         c2_expect="health_alert event arrives with new_rating", c3_expect="Push received for critical/warning")
    ctx.note("global fix XC-17: 4.7 timeout rows and APP-FR-084: no device_offline for a device that was wet when it went silent (health_engine.c:320-323 leak first, :593-596 offline needs cause LINK); completeness-04: the valve closes itself at 10 % or less")
    ctx.note("s2-12: APP-FR-084 keys on device_offline (data.category \"health\"); health events report reachability only; device_recovered can be critical, and the 60 s debounce (2.1.3 dropped the change) (R-s2-05)")

    _req(ctx, 14, 13, "APP-FR-085",
         c2="The app shall colour each device's battery icon by the hub's bands: orange at 20 % or less, green above 20 %. For the valve at 10 % or less the icon shall be red, on every hub version, with the message: \"Valve battery critical. The valve will not open. Replace the batteries.\" A null or missing battery shows \"--\" (a valve battery of 0 from a 2.1.3 hub means unknown).",
         c3="Orange icon at 20 % or less, green icon above 20 %. Valve at 10 % or less: red icon and the message.",
         c2_expect="battery below 20% with a red battery icon", c3_expect="Red icon shown")
    ctx.note("s2-13: APP-FR-085 battery icon follows the hub's bands, as in sections 3.1-3.3 and the rating legend: orange at 20 % or less (inclusive), red with the valve-critical message for the valve at 10 % or less")

    r086 = add_row(t14, ["APP-FR-086", "P1",
                         "A device with last_seen_age_s null has not been heard since the hub restarted or since it was added. For its first minutes the hub leaves it out of system_health.rating although its own rating reads critical. The app should show it with a grey \"Not heard yet\" badge instead of its critical rating while gateway.uptime_s is below 600 (a sensor) or 180 (the valve), or within 10 minutes (a sensor) or 2.5 minutes (the valve) of the provision that added it. After that, its critical rating shows.",
                         "After a hub restart, sensors not heard yet show grey, not red. A sensor still not heard 10 minutes after the restart shows red."],
                   like_row=ctx.row(14, 12, "APP-FR-084"), after_row=ctx.row(14, 13, "APP-FR-085"))
    colour_priority(r086)
    ctx.note("s2-14: new APP-FR-086 (P1) grey \"Not heard yet\" badge, after APP-FR-085")


# ---------------------------------------------------------------- 4.9 Event History (T17)

def _history(ctx):
    ctx.T(17, "Req ID")
    _req(ctx, 17, 1, "APP-FR-101",
         c3="History shows at least 7 days. Events sorted newest first by ts, each shown once (APP-NF-013, APP-NF-014).",
         c3_expect="History shows at least 7 days")
    ctx.note("s2-44: APP-FR-101 acceptance sorts by ts and shows each event once, pointing to APP-NF-013 and APP-NF-014 (the rule itself lives in 8.3; no APP-FR-106)")

    _req(ctx, 17, 2, "APP-FR-102",
         c2="Each history row shall show: timestamp (the envelope ts), event type, description, and affected device. The device is data.valve_id when present, otherwise data.sensor_id; show its location.label (or location.code) from the event or the latest snapshot, and \"Valve\" for the valve. rmleak_cleared and rmleak_auto_cleared name the valve (valve_id; none on a hub with no valve). water_access_override_enabled, water_access_override_expired, auto_close_reenabled and cmd_ack name no device: show the hub name (gateway.name, or the Gateway ID when the hub has no name).",
         c2_expect="Each history row shall show")
    ctx.note("s2-42: APP-FR-102 resolves the device from valve_id or sensor_id; events with no device show the hub")

    _req(ctx, 17, 3, "APP-FR-103",
         c2="The following event types shall be displayed: leak_detected and leak_cleared (data.source_type \"ble_leak_sensor\", \"lora\" or \"valve\"), auto_close, auto_close_blocked_override, water_access_override_enabled, water_access_override_expired, auto_close_reenabled, rmleak_cleared, rmleak_auto_cleared, valve_state_changed, device_offline and device_recovered (data.category \"health\"), and cmd_ack.",
         c3="All listed event types appear with appropriate icons. An unknown event name is stored and shown with a generic icon.",
         c2_expect="health_alert, cmd_ack", c3_expect="All listed event types appear")
    ctx.note("s2-43: APP-FR-103 lists the real event set (no health_alert)")


# ---------------------------------------------------------------- 4.10 Push Notifications (T18)

def _push(ctx):
    ctx.T(18, "Req ID")
    _req(ctx, 18, 1, "APP-FR-111",
         c2="The backend shall send a push notification when a leak_detected event arrives. Title: \"Leak Detected\". Body: \"Water leak detected in [location.label, else location.code shown as a word]\". When the label is empty and the code is \"unknown\", the body reads \"Water leak detected at the valve\" for data.source_type \"valve\" and \"Water leak detected by a leak sensor\" otherwise. The backend also pushes for an event that arrives late, for example replayed after an outage: when its ts is more than 5 minutes older than its arrival, the body adds \"at [local time of ts]\", for example \"Water leak detected in Kitchen at 14:05\". It never pushes for a byte-identical copy of a message it already has (APP-NF-013). It sends no push for a leak_detected from a device whose current state is already wet, for example the repeat a wet sensor sends after a hub restart (section 5.2.4).",
         c3="Push arrives within 10 seconds of hub event. A duplicate copy gives no second push. A leak_detected repeated for a device that is still wet gives no second push.",
         c2_expect="when a leak_detected event arrives", c3_expect="Push arrives within 10 seconds")
    ctx.note("s2-45: APP-FR-111 location fallback copy, late-event time and no push for a duplicate")
    ctx.note("global fix XC-10: APP-FR-111 sends no second push for a device already wet (the repeat after a hub restart), as 5.2.4 says")

    _req(ctx, 18, 2, "APP-FR-112",
         c2="The backend shall send a push notification when an auto_close event arrives. Title: \"Valve Closing\". Body when data.rmleak_asserted is true: \"The hub is closing the valve because of a leak in [location].\" When it is false (the valve was out of reach): \"Leak in [location]. The valve will close when it reconnects.\" [location] is location.label, else location.code shown as a word, from the event or, for data.sensor_id, from the latest snapshot (an auto_close with data.cause \"reconnect\" carries no location, only sensor_id or valve_id). When there is none (the valve flood probe, or a sensor with no location set), the body leaves out \"in [location]\": \"The hub is closing the valve because of a leak.\" or \"Leak detected. The valve will close when it reconnects.\" A late event gets \"at [local time of ts]\" as in APP-FR-111. Send no push when the latest snapshot's data.valve has no valve_id (no valve set up; only a 2.1.3 hub sends auto_close then). The hub can repeat auto_close for the same leak, at most every 10 s, at each further wet report until the valve reports closed with RMLEAK set, which cannot happen while the valve is out of reach. The backend shall therefore push for an auto_close only when one of these holds: no auto_close from the same hub arrived in the previous 10 minutes; a valve_state_changed with valve_state \"open\" from that hub arrived after the previous auto_close (the valve was opened again, so this is a new closing); the previous auto_close had rmleak_asserted false and this one has rmleak_asserted true; or this one has data.cause \"reconnect\". Every auto_close still goes into the history.",
         c3="Push arrives. Tapping opens valve screen. Repeated auto_close events for one leak give one push. An auto_close after the valve was opened again gives a new push.",
         c2_expect="when an auto_close event arrives", c3_expect="Push arrives")
    ctx.note("s2-46, tm-17: APP-FR-112 push copy follows rmleak_asserted (\"Valve Closing\"), with a location fallback; no push without a valve")
    ctx.note("integrator (s4 cross-module, s4-13 / tm-22), global fix INT-03/XC-07: APP-FR-112 pushes once per closing: an auto_close pushes only when none came from the hub in the previous 10 minutes, the valve was reported open since the previous one, rmleak_asserted moves from false to true, or cause is reconnect (rules_engine.c:911-949: auto_close repeats at each wet report after the 10 s cooldown until the valve reports closed with RMLEAK set; a new leak after a reopen is a new closing)")

    _req(ctx, 18, 3, "APP-FR-113",
         c2="The backend shall send a push notification when a device_offline event (data.category \"health\") arrives. Title: \"Device Offline\". Body: \"[name] is offline. Check its battery and range.\" [name] is \"The valve\" when data.source_type is \"valve\"; otherwise the location.label (or location.code shown as a word) from the latest snapshot for data.sensor_id, or \"A leak sensor\" when the label is empty and the code is \"unknown\". device_recovered needs no push. A health event for a device that is no longer provisioned gets no push (APP-NF-017).",
         c3="Push received within 10 s of the event. No push for a device the latest snapshot no longer lists.",
         c2_expect="health_alert event arrives with new_rating \"critical\"", c3_expect="Push received for critical transitions")
    ctx.note("s2-47: APP-FR-113 keys on device_offline; copy names the device; removed devices are ignored (APP-NF-017)")

    ctx.cell(18, 6, 0, "APP-FR-116")
    set_cell(ctx.cell(18, 6, 2, "Copy by trigger"),
             "On a water_access_override_enabled event, the backend shall send a push notification. Copy by trigger: \"c2d_command\" -> \"Water override enabled - automatic shutoff is paused for 24 hours.\"; \"button\" -> \"The valve was opened at the device during a leak - automatic shutoff is paused for 24 hours.\"")
    colour_priority(ctx.row(18, 6, "APP-FR-116"))
    ctx.note("s2-49: APP-FR-116 \"button\" copy no longer asserts that someone pressed it; both copies say \"automatic shutoff\" (canon UI term)")

    # New rows keep V4's alternating fill: APP-FR-117 shaded like R5, APP-FR-118 plain like R6.
    t18 = ctx.T(18, "Req ID")
    r116 = ctx.row(18, 6, "APP-FR-116")
    r117 = add_row(t18, ["APP-FR-117", "P1",
                               "The backend should send a push notification when the valve battery first reaches 10 % or less: a snapshot shows valve.battery at 10 or less, and the last non-null valve.battery before it was above 10, or there was none. Title: \"Valve Battery Critical\". Body: \"The valve battery is at [valve.battery]%. The valve closes itself and will not open until the batteries are replaced.\" The hub sends no health event for a valve battery. A hub on firmware 2.1.3 sends valve.battery 0 when it has not read the battery yet: treat that 0 as null.",
                               "One push per drop to 10 % or less. No new push when the valve reconnects still at 10 % or less. No push for a 0 from a 2.1.3 hub."],
                   like_row=ctx.row(18, 5, "APP-FR-115"), after_row=r116)
    r118 = add_row(t18, ["APP-FR-118", "P1",
                        "On a water_access_override_expired event with auto_close_resumed true (a leak was still active when the window ended), the backend should send a push notification. Title: \"Override Ended\". Body: \"The 24-hour override has ended while a leak is still active. If automatic shutoff is on, the valve is closing now.\"",
                        "Push within 10 s of the event."],
            like_row=r116, after_row=r117)
    for r in (r117, r118):
        colour_priority(r)
    # V4's APP-FR-116 row has runs with no size (11 pt in an 8 pt table); R118 copied it.
    for r in (r116, r118):
        set_row_font_size(r, 8)
    ctx.note("final check: APP-FR-116 and APP-FR-118 rows set to the table's 8 pt (V4's R6 carried no size)")
    ctx.note("global fix INT-05: APP-FR-086, APP-FR-116, APP-FR-117 and APP-FR-118 priority text coloured by V4's house style (P0 red C62828, P1 orange E65100); APP-FR-117 copy says the valve closes itself (completeness-04)")
    ctx.note("s2-48, s2-50: new APP-FR-117 (P1) valve battery critical push from the snapshot, with the 2.1.3 zero rule; new APP-FR-118 (P1) override-ended push")


# ---------------------------------------------------------------- 4.11 Advanced Device Management (T19)

def _device_mgmt(ctx):
    ctx.T(19, "Req ID")
    _req(ctx, 19, 1, "APP-FR-121",
         c2="The app should allow decommissioning a single BLE leak sensor via decommission command with target=\"ble_leak_sensor\" and sensor_id. The MVP app does not remove LoRa leak sensors.",
         c3="After cmd_ack ok, the next snapshot no longer lists the sensor, and twin reported ble_leak_sensor_count goes down (a 2.1.3 hub updates twin reported only at its next connect). Since 2.1.4 the other devices keep their status; a 2.1.3 hub shows them as not heard yet until each is heard again. A sensor that is not provisioned returns the error \"ble sensor decommission failed\". A health event for the removed sensor that arrives after the ack is ignored (APP-NF-017).",
         c2_expect="decommissioning a single BLE leak sensor", c3_expect="Sensor removed from provisioned list")
    ctx.note("s2-51: APP-FR-121 confirms the removal from the snapshot and twin, keeps the other devices' status since 2.1.4, gives the error text")

    _req(ctx, 19, 2, "APP-FR-122",
         c3="After cmd_ack ok, the next snapshot has \"valve\":{} (2.1.3: {\"state\":\"disconnected\",\"connected\":false}). Since 2.1.4 the hub then refuses valve commands (\"No valve is set up for this hub.\") and sends no auto_close; leaks are still reported. If no BLE or LoRa leak sensors remain, the hub is unprovisioned. Since 2.1.4 its rules then go back to auto_close_enabled true and trigger_mask 7, any leak latch and override window are cleared, and it keeps sending snapshots with system_health.reason \"No devices provisioned\": the app shows \"No devices set up yet\" and refreshes Settings from data.rules. A 2.1.3 hub emptied this way keeps its last rules until its next restart and sends no snapshot or lifecycle until it is provisioned again. Since 2.1.4 a hub with no valve returns the error \"valve decommission failed\" (2.1.3 acked \"ok\").",
         c3_expect="Valve removed. Hub is unprovisioned if no sensors remain.")
    ctx.note("s2-52, tm-42: APP-FR-122 acceptance gives the 2.1.4 valve-less state, the empty-hub rules reset and the 2.1.3 behaviour")

    _req(ctx, 19, 3, "APP-FR-123",
         c2="The app may allow a software reset via decommission command with target=\"all\". The hub erases its devices, sensor metadata, hub name, rules state (leak latch and override window), DPS cache and, since 2.1.4, its kept offline events (a 2.1.3 hub keeps them and sends them after it is provisioned again; section 5.5), sets its rules back to auto_close_enabled true and trigger_mask 7, and sets the snapshot interval back to 300 s. It sends a last snapshot (data.reason \"decommission\") and restarts. Its Gateway ID does not change. After the restart it rejoins its configured Wi-Fi, registers with DPS again and applies the Device Twin desired properties again (snapshot_interval_s, hub_name). When a hub is reset or changes owner, the backend shall set desired.hub_name and desired.snapshot_interval_s to null; otherwise the hub applies them again when it reconnects. This does not erase Wi-Fi credentials. To bring back the Wi-Fi setup network, hold the hub's reset button for 10 seconds. That erases the Wi-Fi settings only; until Wi-Fi is set up again the hub is offline to the cloud.",
         c3="Hub restarts about 3 to 5 seconds after cmd_ack (up to about 30 seconds if its cloud link is stalled). After the restart the hub sends lifecycle with provisioned=false and a boot snapshot with \"valve\":{} and \"No devices provisioned\"; the app shows \"No devices set up yet\". A 2.1.3 hub sends no lifecycle or snapshot until it is provisioned again; its twin reported shows provisioned false, valve_id null and both sensor counts 0.",
         c2_expect="software reset via decommission command", c3_expect="Hub reboots within 3 seconds of cmd_ack")
    ctx.note("s2-53: APP-FR-123 lists what target \"all\" erases, the desired-twin cleanup, the 10-second Wi-Fi reset hold and the 2.1.4 / 2.1.3 messages after the restart")
    ctx.note("integrator: APP-FR-123 erasing the kept offline events is since 2.1.4, as 5.3 and 5.5 say (master app_iothub.c has no offline_buffer_clear)")
