# ruff: noqa: E501
"""V5 edits: 5.2.3 valve event, 5.2.4 water detection, 5.2.5 rules events, 5.2.6 health alert, 5.2.7 cmd_ack: V4 P269-P395, tables T24-T26.

Addresses V4's original numbering through ctx (see common.py). Every edit: ctx.note(...)."""
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

from .common import (EditError, add_row, delete_para, delete_row, insert_after,  # noqa: F401
                     insert_lines_after, insert_table_after, replace, replace_in_cell,
                     set_cell, set_text, unique_cells, zebra)

LE = "≤"   # the "less than or equal" sign inside the valve-battery error.detail (U+2264)


def _band(table):
    """V4 bands its tables: odd body rows keep the F5F5F5 fill, even body rows have none.
    insert_table_after copies the (filled) body row 1 for every row, so clear the even ones."""
    for i, row in enumerate(table.rows):
        if i == 0 or i % 2 == 1:
            continue
        for c in unique_cells(row):
            tcPr = c._tc.tcPr
            shd = tcPr.find(qn("w:shd")) if tcPr is not None else None
            if shd is not None:
                tcPr.remove(shd)


def _widths(table, twips):
    """Set a new table's column widths (grid and every cell), in twips."""
    cols = table._tbl.tblGrid.findall(qn("w:gridCol"))
    if len(cols) != len(twips):
        raise EditError(f"_widths: {len(twips)} widths for {len(cols)} grid columns")
    for gc, w in zip(cols, twips):
        gc.set(qn("w:w"), str(w))
    for row in table.rows:
        for c, w in zip(unique_cells(row), twips):
            c.width = Twips(w)


def _gateway(ctx, n, uptime):
    """V4 example envelope line: add short_id, fw 1.9.0 -> the documented firmware."""
    p = ctx.P(n, f'"fw": "1.9.0", "uptime_s": {uptime}')
    replace(p, '"fw": "1.9.0"', f'"short_id": "AD54", "fw": "{ctx.FW}"')
    return p


def _blank(ctx, anchor):
    """An empty Normal paragraph, as V4 puts between a code block or table and the next text."""
    if ctx.orig(286) != "":
        raise EditError("P286 is expected to be V4's empty spacer paragraph")
    return insert_after(anchor, "", like=ctx.P(286))


def apply(ctx):
    normal = ctx.P(306, "Event names: leak_detected, leak_cleared.")   # a plain Normal paragraph in range

    # ------------------------------------------------------------------ 5.2.3 Valve Position Event
    set_text(ctx.P(270, "Sent when the valve opens or closes."),
             '''Sent when the valve reports a new position, whatever moved it: an app command, an auto-close, override_enable, the button on the valve, or the valve's own close at 10 % battery or less (on any hub version). valve_state is the position the valve reports it was driven to (the valve has no position sensor). A command to the position the valve already has sends no event; its cmd_ack is still "ok". A position change made while the valve was not linked to the hub sends no event either; the next snapshot shows it. So the app confirms a valve command when this event or a snapshot shows the new position (APP-FR-055). Water at the valve flood probe is not reported here: it is a water-detection event (5.2.4).''')
    ctx.note("global fix completeness-04: 5.2.3 names the valve's own close at 10 % battery or less as a cause of valve_state_changed (valve FW 2.2.0 readBatteryADC)")
    ctx.note("s4-01: 5.2.3 intro says what sends valve_state_changed and when no event comes (same position, change while unlinked); 1.9.0 history dropped")

    _gateway(ctx, 274, 1071)
    replace(ctx.P(279, '"device_id": "00:80:E1:27:9A:E6"'), '"device_id"', '"valve_id"')
    p283 = ctx.P(283, '"rmleak": true')
    set_text(p283, '''    "rmleak": true,''')
    ctx.code_after(p283, ['''    "fw_version": "2.2.0"'''])
    ctx.note(f"s4-02: valve_state_changed example: valve_id (not device_id), fw_version added, gateway.short_id added, fw {ctx.FW}")

    if ctx.orig(286) != "" or ctx.orig(285) != "}":
        raise EditError("P286 is expected to be the spacer after the 5.2.3 example")
    intro = insert_after(ctx.P(286), "**valve_state_changed field reference:**", like=ctx.P(287, "Valve position event names"), rich=True)
    tbl = insert_table_after(intro, ctx.T(25, "Field"), ["Path", "Type", "Notes"], [
        ["data.source_type", "string", '''Always "valve".'''],
        ["data.valve_id", "string", '''The provisioned valve's BLE MAC, upper case. The same value as data.valve.valve_id in the snapshot.'''],
        ["data.valve_state", "string", '''"open" or "closed". Rarely "unknown": the valve reported a value the hub cannot read, or the link dropped as the event was built (the other valve fields are then not reliable either). Take the valve's state from the next snapshot.'''],
        ["data.battery", "integer/null", '''Valve battery % (0-100). Since 2.1.4, null when the hub has no reading on this link yet. A 2.1.3 hub sends 0 instead: treat a valve battery of 0 from a 2.1.3 hub as unknown.'''],
        ["data.leak_state", "boolean", '''The valve flood probe when the event was built (true = wet).'''],
        ["data.rmleak", "boolean", '''RMLEAK as the valve last reported it. Normally true on the close that follows an auto-close, because the hub asserts RMLEAK before it closes the valve.'''],
        ["data.fw_version", "string", '''Valve firmware version, for example "2.2.0". Left out (not null) when the hub could not read it.'''],
    ])
    _band(tbl)
    _blank(ctx, tbl)
    ctx.note("s4-03: new valve_state_changed field reference table after the 5.2.3 example (battery null since 2.1.4 / 0 on 2.1.3, fw_version optional)")

    set_text(ctx.P(287, "Valve position event names: valve_state_changed only."),
             '''valve_state_changed is the only valve position event. valve_flood_detected and valve_flood_cleared are not sent: water at the valve flood probe arrives as leak_detected or leak_cleared with source_type "valve" (5.2.4). There is no valve connect or disconnect event. When the valve link drops, the hub sends a snapshot within a few seconds with data.valve.state "disconnected" and connected false; data.valve.rating is normally "warning" for the first 3 minutes (system_health.reason includes "Valve disconnected"), then "critical" ("Valve offline"). Only if the valve is still unreachable after those 3 minutes does the hub send device_offline (5.2.6), and then device_recovered when it links again. A shorter drop sends no event; the snapshots show it. A valve whose flood probe was wet when the link dropped reads "critical" at once ("Leak detected: valve, Valve offline") and gets no device_offline: keep its last known leak state (APP-FR-036).''')
    ctx.note("s4-04: valve link loss is a snapshot, not an event; device_offline only after 3 minutes unlinked")
    ctx.note("global fix completeness-03: a valve that drops while its flood probe is wet stays critical as a leak and gets no device_offline (app_ble_valve.c:1902-1923)")

    # ------------------------------------------------------------------ 5.2.4 Water Detection Event
    set_text(ctx.P(289, "ONE event family for"),
             '''ONE event family for "water was detected somewhere", from any leak source: a BLE leak sensor, a LoRa leak sensor or the valve flood probe. Use data.source_type to tell the sources apart; a single handler covers all three. The hub sends these events for every provisioned leak source whatever auto_close_enabled and trigger_mask say: those settings decide only whether the hub closes the valve (4.8.1, 5.2.5).''')
    ctx.note("s4-05: leak events are sent whatever auto_close_enabled and trigger_mask say; 1.9.0 history dropped")

    _gateway(ctx, 293, 1072)
    replace(ctx.P(298, '"device_id": "00:80:E1:27:B4:96"'), '"device_id"', '"sensor_id"')
    ctx.note(f"s4-06: BLE leak_detected example: sensor_id (not device_id), gateway.short_id, fw {ctx.FW}")

    set_text(ctx.P(307, "Required core, identical keys in identical order"),
             '''Required core, the same keys in the same order for every source: event, source_type, the device id, leak_state, battery, location. The device id key is named for the device type: sensor_id for a BLE or LoRa leak sensor, valve_id for the valve. It is always the third key, right after source_type. A BLE leak sensor's sensor_id is its MAC in upper case with colons; a LoRa leak sensor's is "0x" and 8 upper-case hex digits, for example "0x1A2B3C4D". Source-specific keys follow the core. source_type is "ble_leak_sensor", "lora" or "valve". BLE and LoRa leak sensors add rssi (dBm, as heard by the hub). The valve instead adds valve_state, rmleak and fw_version, and carries no rssi because it is a GATT link rather than an advertisement; fw_version is left out (not null) when the hub could not read it. battery is the battery % (0-100) or, since 2.1.4, null when the hub has no reading (in practice only the valve, just after it links). A 2.1.3 hub sends the valve battery as 0 when it has not read it: treat a valve battery of 0 from a 2.1.3 hub as unknown. location is always present: the code and label set with sensor_meta, {"code":"unknown","label":""} for a sensor with none set, and always {"code":"unknown","label":""} for the valve, because valve locations cannot be set.''')
    ctx.note("s4-09: leak event core keys: sensor_id / valve_id as the third key, battery null since 2.1.4 (0 on 2.1.3), fw_version optional, location for a sensor with none set")

    set_text(ctx.P(308, "Example: the valve's own flood probe went wet."),
             '''Example: the valve flood probe went wet. The first six keys come in the same order as in the sensor example, with valve_id in place of sensor_id; only the trailing source-specific keys differ. valve_state and rmleak are the valve's values when the water was detected, before any auto-close: the auto_close event and the next snapshot show RMLEAK.''')
    ctx.note("s4-10: valve leak example intro: valve_id in place of sensor_id; valve_state and rmleak are sampled before any auto-close")

    _gateway(ctx, 312, 1074)
    replace(ctx.P(317, '"device_id": "00:80:E1:27:9A:E6"'), '"device_id"', '"valve_id"')
    replace(ctx.P(319, '"battery": 92'), "92", "22")
    ctx.note("global fix INT-17: valve leak_detected example battery 22, as in the auto_close example of the same valve (was 92)")
    ctx.note(f"s4-11: valve leak_detected example: valve_id (not device_id), gateway.short_id, fw {ctx.FW}")

    # Three rules for every event, placed at the end of 5.2.4 (after the valve example) so the
    # key description in P307 stays next to the examples it describes.
    if ctx.orig(325) != "}" or not ctx.orig(326).startswith("5.2.5"):
        raise EditError("P325 is expected to close the valve leak example")
    end = _blank(ctx, ctx.P(325))
    end = insert_after(end, '''**When leak events are sent.** A leak sensor's event goes out when the leak state the hub hears from it changes. The first report the hub hears from a sensor after the hub starts, or after a provision adds the sensor again, sends leak_detected if the sensor is wet and nothing if it is dry. A wet valve flood probe is reported when the valve links, and not again at each reconnect while it stays wet; a probe that dried while the valve was not linked sends leak_cleared when the valve links again. So leak_detected can repeat for one device with no leak_cleared in between (after a hub restart, or after a wet sensor is removed and added again), and a device that dried while the hub was restarting never sends leak_cleared. A sensor that goes silent while wet stays wet in the snapshot until the hub hears it dry. Treat each leak event as the device's current state, send no second leak push (APP-FR-111) for a device whose current state is already wet, and take the final state of the leak banner from leak_state in the snapshot.''',
                       like=normal, rich=True)
    ctx.note("s4-07: new 'When leak events are sent' paragraph (state-change gate, repeats after a restart, no leak_cleared for a device that dried during a restart)")
    end = insert_after(end, '''**Duplicates and order.** Any event in 5.2.3 to 5.2.7 can arrive twice, and the second copy can arrive after newer messages: the hub delivers at least once (5.5). Since 2.1.4 this happens more often, because an event whose send meets a dropping connection is kept and sent again after the reconnect. The copies are byte-identical, ts included. The backend drops a message only when its body is byte-identical to one already received, keeps the first copy, and never de-duplicates on a subset of fields such as gateway.id + ts + event (APP-NF-013). One rare, harmless exception: two copies of an event raised before the hub's clock was first set can differ in ts by about 1 s and in nothing else; the backend may treat them as one. Order events by ts, not by arrival time (APP-NF-014). An event older than the latest snapshot goes into the history only and does not change the displayed state; a leak event also changes a device's state only when its ts is not older than the last leak event already applied for that device. Events raised while the hub was offline arrive after it reconnects, before its lifecycle message, each with the ts of when it happened; since 2.1.4 this includes events raised before its clock was first set after a power cut. A 2.1.3 hub could lose an event whose send met a dropping connection, instead of sending it twice, and discarded every event raised before its clock was set; the same rules work for both.''',
                       like=normal, rich=True)
    ctx.note("s4-08: new 'Duplicates and order' paragraph: byte-identical dedupe (canon C10, APP-NF-013), order by ts (APP-NF-014), late pre-sync events since 2.1.4")
    insert_after(end, '''**Snapshot after each event.** Every event in 5.2.3 to 5.2.6 makes the hub send a snapshot, normally within about 5 s, that shows the state after the event. Since 2.1.4 so does a change of RMLEAK as the valve reports it, which has no event of its own; a 2.1.3 hub shows that change in its next snapshot, which can be the heartbeat. While the hub is offline it sends no snapshot; it sends a fresh one after it reconnects. Apply an event to the UI at once, then take the final state from that snapshot.''',
                 like=normal, rich=True)
    ctx.note("s4-v01: new 'Snapshot after each event' paragraph (snapshot within about 5 s; RMLEAK-change snapshot since 2.1.4)")

    # ------------------------------------------------------------------ 5.2.5 Rules Engine Events
    set_text(ctx.P(327, "Rules engine events are generated by the hub auto-close logic."),
             '''Rules engine events come from the hub's auto-close logic, in the standard envelope with type "event". They carry no category key (health events carry category "health"). Which device they name: auto_close and auto_close_blocked_override carry source_type and the id of the leak source that triggered them (sensor_id, or valve_id for the valve flood probe). rmleak_cleared and rmleak_auto_cleared carry valve_id and no source_type, because RMLEAK belongs to the valve whichever device detected the water. The override window events name no device. valve_id is always the provisioned valve's MAC. It is left out on a hub with no valve, and rarely when the hub was too busy to read it, so treat it as optional. Up to 2.1.3, valve_id named the valve the hub was linked to, which in rare cases was not the hub's own valve: ignore a valve_id that differs from the valve_id in twin reported or in the lifecycle message, which both name the provisioned valve. Do not compare it with the snapshot: while the valve is linked, a 2.1.3 snapshot also names the linked valve. The hub sends leak_detected before the auto_close it causes, but match an auto_close to its leak by sensor_id or valve_id and ts, not by position in the stream: since 2.1.4 a leak report the hub could not evaluate at once is evaluated a moment later, and its auto_close can then arrive after the same sensor's leak_cleared. A 2.1.3 hub skipped such a report, so no auto_close came for it.''')
    ctx.note("s4-12: 5.2.5 intro: which device each rules event names and under which key, valve_id optional (2.1.3 linked-valve MAC), match auto_close by id and ts")

    T = 24
    ctx.T(T, "Event Name")
    # R1 auto_close
    set_cell(ctx.cell(T, 1, 1, "Rules engine closed valve due to leak."),
             '''The hub asserted RMLEAK and closed the valve for a leak: a wet report from a leak source whose trigger-mask bit is set, with auto_close_enabled true and no override window open (4.8.1). If the valve is out of reach, the hub holds both commands and sends them when it links again. Normally once per incident; it can repeat, at most every 10 s, while more wet reports arrive before the valve reports closed with RMLEAK set. Not sent by a hub with no valve (since 2.1.4; 2.1.3 sent it with rmleak_asserted false).
Form with cause "reconnect": the valve linked again while a provisioned leak source was still wet, auto_close_enabled was true and no override window was open, and the hub closed it then (no event when the valve was already closed with RMLEAK set, or within 10 s of the previous reconnect auto_close; the hub still closes the valve). This form does not check the trigger mask: any wet provisioned leak source counts (4.8.1).''')
    set_cell(ctx.cell(T, 1, 2, "source_type, device_id, location"),
             '''source_type, sensor_id (BLE or LoRa leak sensor) or valve_id (valve flood probe), rmleak_asserted, location (only when sensor_meta set one for that sensor; never for the valve).
rmleak_asserted true: the valve was linked and the hub sent RMLEAK and the close. false: the valve was out of reach; the hub holds both until it links again.
Form with cause "reconnect": cause, sensor_id or valve_id (the first leak source still wet), rmleak_asserted, active_leak_count; no source_type and no location.''')
    set_cell(ctx.cell(T, 1, 3, "Send push. Show valve locked."),
             '''Send push (APP-FR-112); data.location can be missing. rmleak_asserted true: show the valve closing and locked until valve_state_changed or a snapshot confirms it (APP-FR-037). rmleak_asserted false: show "Will close when the valve reconnects" (APP-FR-037) until valve_state_changed or a snapshot shows the valve closed; an auto_close with cause "reconnect" can follow when the valve links.''')
    ctx.note("global fix completeness-09: the reconnect auto_close is not sent within 10 s of the previous one; the writes are still issued (rules_engine.c:1640-1657; same on 2.1.3)")
    ctx.note("s4-13, tm-22: T24 auto_close row: trigger-mask condition, repeats, no event without a valve since 2.1.4, cause \"reconnect\" form (mask not checked), sensor_id/valve_id, rmleak_asserted, location only with sensor_meta")

    # R2 auto_close_blocked_override
    set_cell(ctx.cell(T, 2, 1, "Leak during 24h override window."),
             '''A wet report that would have closed the valve (auto_close_enabled true, the source's trigger-mask bit set) while a 24-hour override window is open. The valve stays open. At most one per 60 s for the whole hub, so a second leak within 60 s sends leak_detected only.''')
    set_cell(ctx.cell(T, 2, 2, "source_type, device_id"),
             '''source_type, sensor_id or valve_id, override_remaining_s (seconds left in the window; treat it as optional: a 2.1.3 hub leaves it out before its clock is set)''')
    ctx.note("s4-14, tm-23: T24 auto_close_blocked_override row: only for a leak that would have closed the valve, 60 s per hub, sensor_id/valve_id, override_remaining_s")

    # R3 water_access_override_enabled
    set_cell(ctx.cell(T, 3, 1, "Valve opened after an auto-close"),
             '''A 24-hour override window started: the user cleared RMLEAK with the button on the valve, which also opens it (trigger "button"), or the app sent override_enable (trigger "c2d_command"). Also trigger "button" when the valve links again and the hub finds it open with RMLEAK cleared while a leak incident is latched: a button press the hub could not see. If a leak source is still wet at that moment and auto_close_enabled is true, the hub closes the valve instead (auto_close with cause "reconnect") and no window starts. An override_enable while a window is open restarts the window at 24 hours and sends this event again; the valve button does not.''')
    set_cell(ctx.cell(T, 3, 2, "trigger, expires_ts, remaining_s"),
             '''trigger, expires_ts (end of the window, Unix epoch; since 2.1.4 left out when the window started before the hub's clock was set, where a 2.1.3 hub sends a time in 1970), remaining_s (always 86400)''')
    set_cell(ctx.cell(T, 3, 3, "Show 24h override banner."),
             '''Show 24h override banner. Send push (APP-FR-116). Use remaining_s when expires_ts is missing or earlier than 2024 (1704067200).''')
    for c in (1, 2):
        for p in ctx.cell(T, 3, c).paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)
    ctx.note("global fix INT-19: T24 water_access_override_enabled row, cells 1-2: 9 pt like the rest of the table (V4 gave them no size)")
    ctx.note("s4-15: T24 water_access_override_enabled row: trigger values, reconnect inference, refresh only by override_enable, expires_ts left out before clock sync since 2.1.4")

    # R4 water_access_override_expired
    set_cell(ctx.cell(T, 4, 1, "24h override window timed out."),
             '''The 24-hour override window ended. If any provisioned leak source is still wet and auto_close_enabled is true, the hub asserts RMLEAK and closes the valve at once, even when that source's trigger-mask bit is clear (4.8.1). No auto_close is sent for that close; valve_state_changed (closed, rmleak true) follows when the valve closes. If the valve is not linked, the close waits until it links again.''')
    set_cell(ctx.cell(T, 4, 2, "(none)"),
             '''auto_close_resumed (true when a leak source was still wet when the window ended; it does not say the valve closed: with auto_close_enabled false nothing closes), active_leak_count (wet leak sources when the window ended)''')
    set_cell(ctx.cell(T, 4, 3, "Remove override banner. Auto-close re-enabled."),
             '''Remove override banner. When auto_close_resumed is true, send push (APP-FR-118). If auto_close_enabled is also true (data.rules in the latest snapshot), expect the valve to close with RMLEAK set. Show the valve from valve_state_changed and the snapshot, not from this event.''')
    ctx.note("s4-16, tm-24: T24 water_access_override_expired row: immediate re-close with no auto_close (mask not checked), auto_close_resumed and active_leak_count")

    # R5 auto_close_reenabled
    set_cell(ctx.cell(T, 5, 1, "Override cancelled by C2D command."),
             '''override_cancel ended an open override window. If any provisioned leak source is still wet and auto_close_enabled is true, the hub asserts RMLEAK and closes the valve at once, even when that source's trigger-mask bit is clear (4.8.1). No auto_close is sent for that close; valve_state_changed follows when the valve closes. Not sent when no window was open (the cmd_ack is still "ok").''')
    set_cell(ctx.cell(T, 5, 2, 'reason: "c2d_command"'),
             '''previous_remaining_s (seconds that were left in the window), reason (always "c2d_command")''')
    set_cell(ctx.cell(T, 5, 3, "Remove override banner."),
             '''Remove override banner. This event can be missing (a later rules event can replace it before it is sent), so also remove the banner on cmd_ack "ok" for override_cancel, and confirm it with override_active false in the next snapshot. If a leak is still active and auto_close_enabled is true, expect the valve to close with RMLEAK set; confirm from valve_state_changed and the snapshot.''')
    ctx.note("s4-17, tm-25: T24 auto_close_reenabled row: immediate re-close with no auto_close (mask not checked), previous_remaining_s, not sent without an open window")

    # R6 rmleak_cleared
    set_cell(ctx.cell(T, 6, 1, "Leak reset command cleared RMLEAK."),
             '''leak_reset cleared a latched leak incident, RMLEAK on the valve, or an open override window (leak_reset ends the window too). It does not open the valve. Not sent when there was nothing to clear (the cmd_ack is still "ok"). While any provisioned leak source is still wet, including one whose trigger-mask bit is clear, leak_reset is refused with a cmd_ack error and this event is not sent.''')
    set_cell(ctx.cell(T, 6, 2, "(none)"),
             '''valve_id (left out on a hub with no valve), override_cancelled (true; present only when the leak_reset ended an override window)''')
    set_cell(ctx.cell(T, 6, 3, "Update valve card. Enable Open Valve."),
             '''Update valve card: RMLEAK cleared, the valve does not move; take the lock icon from valve.rmleak in the snapshot that follows. Enable Open Valve. Remove the override banner when override_cancelled is true.''')
    ctx.note("s4-18: T24 rmleak_cleared row: valve_id, override_cancelled, ends the override window, refused while any source is wet, no event when nothing to clear")

    # R7 rmleak_auto_cleared
    set_cell(ctx.cell(T, 7, 1, "RMLEAK auto-cleared after 30s"),
             '''The hub cleared RMLEAK by itself: every provisioned leak source, including any whose trigger-mask bit is clear, had been dry for 10 s. In practice it comes 10 to 12 s after the last source reports dry. Up to 2.1.3 the wait was 30 s and the clear came 30 to 60 s after the last dry report. Sent only after a latched leak incident, also on a hub with leak sensors and no valve. The valve stays closed: the user then opens it, and Leak Reset is not needed.''')
    set_cell(ctx.cell(T, 7, 2, "(none)"),
             '''valve_id (left out on a hub with no valve), clear_after_seconds (the wait in force: 10; 30 on 2.1.3)''')
    set_cell(ctx.cell(T, 7, 3, "Update valve card."),
             '''Update valve card: closed, RMLEAK cleared (confirm from valve.rmleak in the next snapshot). Enable Open Valve.''')
    ctx.note("s4-19, tm-26: T24 rmleak_auto_cleared row: 10 s since 2.1.4 (30 s on 2.1.3), every provisioned source counts, valve_id and clear_after_seconds, valve stays closed")

    # After T24 and V4's spacer P328: RMLEAK and override changes with no rules event,
    # closed by a blank, so it reads T24, blank, intro, bullets, blank, auto_close example.
    ctx.T(24, "Event Name")
    if ctx.orig(328) != "" or ctx.orig(329) != "Example: auto_close event":
        raise EditError("P328 is expected to be V4's empty spacer after T24")
    cur = insert_after(ctx.P(328), '''**RMLEAK and override changes that send no rules event.** Take RMLEAK from data.valve.rmleak (present only while the valve is linked) and the window from data.override_active in the snapshot; do not rely on the event stream alone:''',
                       like=normal, rich=True)
    cur = insert_lines_after(cur, [
        '''decommission target "all" ends the leak incident and the override window with no rmleak_cleared and no water_access_override_expired, on 2.1.3 and 2.1.4. Since 2.1.4, removing the hub's last device does the same (a 2.1.3 hub emptied by a decommission falls silent).''',
        '''Since 2.1.4, replacing or removing the valve while no other leak source is wet ends the leak incident at once, with no rmleak_auto_cleared. A 2.1.3 hub kept the incident until the automatic clear or a leak_reset.''',
        '''An override window that ends, at expiry or with override_cancel, while a leak source is still wet and auto_close_enabled is true closes the valve with no auto_close (see water_access_override_expired and auto_close_reenabled above).''',
        '''The hub holds one rules event at a time. In a rare race a rules event is replaced by a later one before it is sent.''',
        '''If the hub cannot re-read a rules or health event it has built (for example when it is short of memory), it sends data.event "rules_engine" or "health_engine" with the original JSON as a string in data.raw. Log it; the next snapshot carries the state.''',
    ], style="List Bullet")
    _blank(ctx, cur)
    ctx.note("s4-20: new list after T24 of RMLEAK and override changes that send no rules event (decommission all / hub emptied, valve replaced, silent re-close, single pending slot, raw fallback)")

    _gateway(ctx, 333, 1073)
    p338 = ctx.P(338, '"device_id": "00:80:E1:27:B4:96"')
    replace(p338, '"device_id"', '"sensor_id"')
    ctx.code_after(p338, ['''    "rmleak_asserted": true,'''])
    ctx.note(f"s4-21, tm-27: auto_close example: sensor_id (not device_id), rmleak_asserted added, gateway.short_id, fw {ctx.FW}")

    if ctx.orig(341) != "}" or not ctx.orig(342).startswith("5.2.6"):
        raise EditError("P341 is expected to close the auto_close example")
    cur = _blank(ctx, ctx.P(341))
    cur = insert_after(cur, '''Example: auto_close with cause "reconnect", sent when the valve linked again while a leak was still active. It has cause and active_leak_count, and no source_type or location:''',
                       like=ctx.P(329, "Example: auto_close event"))
    ctx.code_after(cur, [
        '''{''',
        '''  "schema": "eflostop.v2",''',
        '''  "ts": 1770589720,''',
        f'''  "gateway": {{ "id": "GW-34B7DA6AAD54", "short_id": "AD54", "fw": "{ctx.FW}", "uptime_s": 1290 }},''',
        '''  "type": "event",''',
        '''  "data": {''',
        '''    "event": "auto_close",''',
        '''    "cause": "reconnect",''',
        '''    "sensor_id": "00:80:E1:27:B4:96",''',
        '''    "rmleak_asserted": true,''',
        '''    "active_leak_count": 1''',
        '''  }''',
        '''}''',
    ])
    ctx.note('s4-22: new example of the auto_close cause "reconnect" form')

    # ------------------------------------------------------------------ 5.2.6 Health events
    set_text(ctx.P(342, "5.2.6 Health Alert Event"), "5.2.6 Health Events: device_offline and device_recovered")
    ctx.note("s4-23: heading 5.2.6 renamed 'Health Events: device_offline and device_recovered'")

    p343 = ctx.P(343, "Sent when a device health rating changes.")
    set_text(p343, '''Sent only when a device's reachability changes. device_offline: a device the hub has heard since it started (or since it was added) loses its link: a BLE or LoRa leak sensor silent for 10 minutes (checked every 30 s, so 10 to 10.5 minutes), or the valve unlinked for 3 minutes. device_recovered: that device is heard again; it comes only after a device_offline. No device_offline is sent for a device not heard since the hub started or since it was added (the snapshot shows it with last_seen_age_s null), or for a device that was wet when it went silent (it stays critical with leak_state true in the snapshot). No health event is sent for a leak, for a battery at any level (including a valve at 10 % or less) or for a weak signal: those reach the cloud only in the snapshot (device rating and system_health). At most one health event per device per 60 s. Since 2.1.4 a change that falls inside those 60 s is sent when they end; 2.1.3 dropped it, so a 2.1.3 hub's events could leave a device offline after it came back: take the device's state from the snapshot.''')
    ctx.note("s4-24: 5.2.6 intro: health events report reachability only (canon C6); never-heard and wet-when-silent devices send none; held-back change sent since 2.1.4")
    insert_after(p343, '''There is no event named health_alert and no field new_rating. A health event is an event with data.category "health" and data.event "device_offline" or "device_recovered"; its new rating is data.rating, and device_offline always carries "critical". Battery and signal ratings, including a valve battery at 10 % or less, come only from the snapshot: a notification driven by them, such as APP-FR-117, compares each snapshot with the one before. Leaks come from leak_detected and leak_cleared (5.2.4).''')
    ctx.note("s4-25: new paragraph: no health_alert event, no new_rating field; battery and signal ratings only in the snapshot")

    _gateway(ctx, 347, 1200)
    replace(ctx.P(352, '"dev_type": "ble_leak_sensor"'), '"dev_type"', '"source_type"')
    replace(ctx.P(353, '"device_id": "00:80:E1:27:B6:A5"'), '"device_id"', '"sensor_id"')
    replace(ctx.P(355, '"prev_rating": "good"'), '"good"', '"warning"')
    ctx.note(f"s4-26: device_offline example: source_type (not dev_type), sensor_id (not device_id), prev_rating warning (battery 15 % and RSSI -92 rate warning), gateway.short_id, fw {ctx.FW}")

    T = 25
    ctx.T(T, "Field")
    set_cell(ctx.cell(T, 3, 0, "data.dev_type"), "data.source_type")
    set_cell(ctx.cell(T, 3, 2, "Same vocabulary as source_type"),
             '''"valve", "ble_leak_sensor" or "lora". The same key and values as on leak and auto_close events, so one lookup table covers all of them.''')
    ctx.note("s4-27: T25 data.dev_type -> data.source_type; 1.8.0 history dropped")
    set_cell(ctx.cell(T, 4, 0, "data.device_id"), "data.sensor_id or data.valve_id")
    set_cell(ctx.cell(T, 4, 2, "Device identifier, uppercase"),
             '''The device's id: sensor_id for a BLE or LoRa leak sensor, valve_id for the valve, the same key as in the snapshot and leak events. Upper case: a BLE MAC like "00:80:E1:27:B6:A5", a LoRa id like "0x1A2B3C4D". A health event held on the hub while it could not send can arrive after the cmd_ack of the decommission that removed its device: ignore it (APP-NF-017).''')
    ctx.note("s4-28: T25 data.device_id -> data.sensor_id or data.valve_id; late health event after a decommission is ignored (APP-NF-017)")
    set_cell(ctx.cell(T, 5, 2, "New health rating."),
             '''The device's rating after the change. device_offline: always "critical". device_recovered: any rating, including "critical" when the device came back wet or, since 2.1.4, is a valve at 10 % battery or less. A 2.1.3 hub never sends device_recovered with "critical": for a device that came back wet it sends none. Tell offline from recovered by data.event, never by comparing ratings.''')
    ctx.note("s4-29: T25 data.rating: values per event; branch on data.event")
    set_cell(ctx.cell(T, 6, 2, "Previous health rating."),
             '''The rating before the change. For the valve's device_offline it is normally "warning" (its first 3 minutes unlinked). Since 2.1.4 it can equal data.rating, when the event was held back by the 60 s limit and sent afterwards.''')
    ctx.note("s4-30: T25 data.prev_rating: valve normally warning; can equal rating since 2.1.4")
    set_cell(ctx.cell(T, 8, 2, "Last known RSSI (dBm). Omitted when unknown."),
             '''A leak sensor's last known RSSI (dBm). Omitted when unknown; never present for the valve.''')
    ctx.note("s4-31: T25 data.rssi never present for the valve")
    set_cell(ctx.cell(T, 9, 2, "Seconds since last contact."),
             '''device_offline only. A leak sensor: seconds since its last packet, normally 600 to 630. The valve: seconds since its link dropped, normally 180 to 210 (since 2.1.4; 2.1.3 counted from the valve's last changed reading and could show hours). Omitted when 0 or unknown.''')
    ctx.note("s4-32: T25 data.offline_duration_s: device_offline only, expected values, valve measured from the link drop since 2.1.4")
    zebra(ctx.T(T, "Field"), ctx.body_shd)
    ctx.note("global fix INT-06: T25 (health event fields) re-banded; V4's own banding was broken (rows 1 and 2 both shaded)")

    # ------------------------------------------------------------------ 5.2.7 cmd_ack
    set_text(ctx.P(363, "The hub sends cmd_ack as an event after processing"),
             '''The hub sends one cmd_ack, as an event, for every command that arrives as an eflostop.cmd envelope: schema "eflostop.cmd" (or the legacy "eflostop.cmd.v1") and a non-empty string cmd, with or without an id. data.id is left out when the command had no id. Always set one, a UUID or another unique string of at most 63 characters: a longer id comes back cut and will not match. A message that is not a valid envelope (not JSON, another or misspelt schema, no cmd) gets no cmd_ack. It is not always ignored: the hub can run a legacy command found anywhere in its text, or apply it as a provision payload, with no cmd_ack (5.3). So the backend builds every command with a JSON library as an exact eflostop.cmd envelope. A C2D message over 8,192 bytes is refused with an error cmd_ack.''')
    ctx.note("s4-33: 5.2.7 intro: an ack for every envelope command, with or without id; id at most 63 characters; no ack for a non-envelope message, which can still run a legacy command (5.3); 8,192-byte limit")

    _gateway(ctx, 368, 1080)
    ctx.note(f"s4-34: cmd_ack success example: gateway.short_id, fw {ctx.FW}")
    _gateway(ctx, 382, 1081)
    ctx.note(f"s4-35: cmd_ack error example: gateway.short_id, fw {ctx.FW}")

    T = 26
    ctx.T(T, "Field")
    set_cell(ctx.cell(T, 2, 2, "The correlation id from the original command."),
             '''The correlation id from the command, at most 63 characters: a longer id comes back cut and will not match. Left out, not empty, when the command had none. Match on this to resolve pending commands; the first ack for an id resolves it.''')
    ctx.note("s4-36: T26 data.id: left out when absent, cut to 63 characters, first ack resolves")
    set_cell(ctx.cell(T, 3, 2, "The command name that was executed."),
             '''The cmd value as the hub received it (up to 31 characters), also when the command was refused or unknown. "unknown" when the hub refused an oversized message before it could read the name.''')
    ctx.note("s4-37: T26 data.cmd: echo of the received cmd, also for refused and unknown commands")
    set_cell(ctx.cell(T, 4, 2, '"ok" on success, "error" on failure.'),
             '''"ok": the hub accepted the command. For valve_open, valve_close and valve_set_state it means queued, not that the valve moved. If the valve is not linked, the hub holds the command and sends it when the valve links again, with no time limit (a newer valve command replaces it; a hub restart drops it). valve_state_changed or a snapshot confirms the move (APP-FR-055); a command to the position the valve already has sends no valve_state_changed.
"error": the hub refused the command; error.detail says why.''')
    ctx.note("s4-38: T26 data.status: ok means accepted, and queued for valve commands (canon C12, C23)")
    set_cell(ctx.cell(T, 5, 2, "The command name (for programmatic matching)."),
             '''Only present when status is "error". Always the same as data.cmd: it does not identify the error. To tell errors apart, match error.detail exactly (table below).''')
    ctx.note("s4-39: T26 data.error.code is always the command name; match error.detail")
    set_cell(ctx.cell(T, 6, 2, "Human-readable error message."),
             '''Why the command was refused, as fixed text (table below). "Show" texts are written for users and are shown as they are (APP-FR-054); "Log" texts are technical: log them and show the app's own message.''')
    ctx.note("s4-40: T26 data.error.detail: fixed texts, Show or Log per the table below")

    cur = _blank(ctx, ctx.T(26, "Field"))
    cur = insert_after(cur, '''**cmd_ack error.detail texts.** Exact, from hub firmware 2.1.4; match them exactly. Use: "Show" = written for users, shown as is (APP-FR-054); "Log" = technical: log it and show the app's own message.''',
                       like=normal, rich=True)
    err = insert_table_after(cur, ctx.T(24, "Event Name"), ["Command", "error.detail (exact)", "When", "Use"], [
        ["valve_open, valve_close, valve_set_state", "No valve is set up for this hub.", '''No valve provisioned. Since 2.1.4 (2.1.3 acked "ok").''', "Show"],
        ['valve_open, valve_set_state "open"', "Valve is locked after a leak (RMLEAK). Clear it with leak_reset first, or use override to open the valve during a leak.", '''The valve reports RMLEAK set, or a leak incident is latched with no override window open, also while the valve is disconnected. 2.1.3 checked only the valve's own RMLEAK.''', "Show"],
        ['valve_open, valve_set_state "open"', f"Valve battery critical ({LE}10 %): the valve will not open. Replace the batteries.", '''The valve's last battery reading is 10 % or less. Since 2.1.4 (2.1.3 acked "ok" and the valve stayed shut). Closing is never refused for battery.''', "Show"],
        ["valve_open, valve_close, valve_set_state", "The valve command could not be queued. Try again.", '''The hub's valve command queue is full. Since 2.1.4 (2.1.3 acked "ok").''', "Show"],
        ["valve_set_state", '''missing 'state' field (expected "open" or "closed")''', "No string state in the payload.", "Log"],
        ["valve_set_state", '''invalid state value (expected "open" or "closed")''', '''state is not "open" or "closed".''', "Log"],
        ["leak_reset", "A leak is still active. Fix the leak first, or use override to open the valve during a leak.", '''A provisioned leak source is wet (including one whose trigger-mask bit is clear); rarely also when the hub was busy for 1 s: resend after a few seconds.''', "Show"],
        ["override_enable", "No active leak to override. Use the normal Open Valve control.", "No leak incident latched, no valve RMLEAK and no window open.", "Show"],
        ["override_enable", "Water detected at the valve. It can't be opened remotely until the valve area is dry.", "The valve flood probe is wet.", "Show"],
        ["override_enable", "The valve isn't responding. Check its power and connection, then try again.", "The valve did not link within about 10 s.", "Show"],
        ["override_enable", "No valve is set up for this hub.", "No valve provisioned.", "Show"],
        ["override_enable", "Something went wrong applying the override. Your water state is unchanged. Try again.", "Internal error.", "Show"],
        ["override_cancel", "override cancel failed", '''The hub was busy. No open window is not an error: it acks "ok".''', "Log"],
        ["rules_config", "rules config update failed", "Payload missing or null, or the save failed.", "Log"],
        ["sensor_meta", "sensor metadata update failed", "Payload missing or rejected, or the save failed.", "Log"],
        ["provision", "provisioning failed", "valve_id not a valid MAC, nothing usable in the payload, or the save failed.", "Log"],
        ["decommission", "missing decommission target", "No target.", "Log"],
        ["decommission", "unknown decommission target", "target is not valve, lora, ble_leak_sensor (or ble, ble_leak) or all.", "Log"],
        ["decommission", "valve decommission failed", '''No valve provisioned, or the save failed. No valve: since 2.1.4 (2.1.3 acked "ok").''', "Log"],
        ["decommission", "ble sensor decommission failed", "sensor_id missing or not provisioned, or the save failed.", "Log"],
        ["decommission", "lora sensor decommission failed", "That LoRa sensor is not provisioned, or the save failed.", "Log"],
        ["decommission", "full decommission failed", "The erase failed.", "Log"],
        ["set_hub_name", "missing 'name' field", "No string name.", "Log"],
        ["set_hub_name", "name too long (max 31 chars)", "name longer than 31 bytes (UTF-8).", "Log"],
        ["any", "unknown command", "cmd not recognised.", "Log"],
        ["any", "exceeds the reassembly limit: <N> bytes exceeds the 8192 byte limit", '''The C2D message is over 8,192 bytes; N is its size. data.cmd is "unknown" when the hub could not read it.''', "Log"],
        ["any", "out of memory: <N> bytes exceeds the 8192 byte limit", "The hub had no memory to receive a large message; the text reads exactly so. Resend.", "Log"],
    ])
    _band(err)
    # T24's grid (2589/2123/2098/2046) starves the long texts; keep its 8,856-twip total.
    _widths(err, [1750, 3200, 3106, 800])
    ctx.note("s4-41: new table after T26 with every cmd_ack error.detail text (canon C12), with the 2.1.4 refusals and their 2.1.3 behaviour")

    cur = _blank(ctx, err)
    insert_after(cur, '''**When the cmd_ack arrives.** The hub sends the cmd_ack as soon as it has handled the command, normally ahead of the events the command causes. It comes before the snapshot the command causes and, since 2.1.4, before the twin report it causes. 2.1.3 sent the twin report first for provision, rules_config and set_hub_name, and sent none after a single-device decommission. Do not depend on this order: the snapshot is the reference. Older events can still arrive after the ack. After every "ok" the hub sends a snapshot, normally within about 5 s, that shows its state after the command; for a valve command, valve_state_changed and the snapshot that follows it confirm the move. Most acks leave the hub within a second or two of the command; override_enable can take about 10 s more while the hub links the valve. If the hub's send queue is full, the ack is held and sent later, so it can arrive after the app's 20 s timeout (APP-FR-053): the backend records it against its command, and the app takes the state from the next snapshot (APP-NF-018). An ack can also be lost, for example when the connection drops as it is sent: a missing ack means the result is unknown, not that the command failed. The hub does not de-duplicate commands by id: a command delivered twice runs twice and is acked twice; the first ack for an id resolves it, and later ones are only logged.''',
                 like=normal, rich=True)
    ctx.note("s4-42: new 'When the cmd_ack arrives' paragraph: order (ack before the twin report since 2.1.4), snapshot after ok, late or lost acks (APP-NF-018), no id de-duplication")
