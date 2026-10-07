# ruff: noqa: E501
"""V5 edits: 9 Edge Cases, 10 Acceptance Tests, 11 Traceability Matrix, Appendix A.1-A.4: V4 P518-end, tables T37-T44.

Addresses V4's original numbering through ctx (see common.py). Every edit: ctx.note(...).

Findings applied: s7-01..s7-66 (s7-23, s7-27, s7-45, s7-59 do not exist), s7-v01..s7-v07,
tm-36..tm-41, tm-v04, tm-v07. Facts follow final/canon.md (C3 trigger mask, C4 battery,
C5 devices not heard yet, C6 health events, C7 RMLEAK clear, C8 empty hub, C9 data.reason,
C10 dedupe, C11 ordering, C12 error texts, C13 offline buffer, C14 offline gate, C15 C2D
expiry, C16 hub name, C19 legacy keywords, C20 LoRa, C21 Wi-Fi setup, C22 no pre-2.1.3
history, C23 valve confirmation, C24 banner). IDs follow final/id_map.json: TC-013..TC-022,
TC-N09, TC-N11..TC-N18 (no TC-N10: V4's revision history cites it), APP-NF-013/014/017/018
in the trace matrix, no APP-FR-106.

Review fixes R7-01..R7-12 applied: content corrections checked against the code, TC-N04 per
canon C17, T40 R4 per C20, expect= guards on every ctx.row(), and table formatting (_format:
V4's zebra shading across inserted rows, 9 pt runs in the V4 rows that had no size).
"""
import copy

from docx.oxml.ns import qn

from .common import (EditError, add_row, delete_para, delete_row, insert_after,  # noqa: F401
                     insert_lines_after, insert_table_after, replace, replace_in_cell,
                     set_cell, set_text, unique_cells)


def _set_row(ctx, k, r, expect, values):
    """Set the distinct cells of V4's row R<r> of T<k>; a None value keeps that cell."""
    ctx.row(k, r, expect)
    for c, v in enumerate(values):
        if v is not None:
            set_cell(ctx.cell(k, r, c), v)


def _add_rows(table, after, like, rows):
    """Insert rows in order after `after` (a _Row), copying `like`'s formatting."""
    cur = after
    for vals in rows:
        cur = add_row(table, vals, like_row=like, after_row=cur)
    return cur


# Elements that may precede w:shd inside w:tcPr (ECMA-376 order).
_SHD_BEFORE = ("w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders")


def _restripe(table):
    """V4 zebra: body rows 1, 3, 5 ... carry R1's F5F5F5 shading, rows 2, 4, 6 ... none.

    Run after every row of the table is in place: inserted rows copy R1 (shaded), so they
    would otherwise break the alternation."""
    tpl = table.rows[1]._tr.find(qn("w:tc") + "/" + qn("w:tcPr") + "/" + qn("w:shd"))
    if tpl is None:
        raise EditError("restripe: R1 has no shading")
    tpl = copy.deepcopy(tpl)
    for i, row in enumerate(table.rows[1:], start=1):
        for tc in row._tr.findall(qn("w:tc")):
            tcpr = tc.get_or_add_tcPr()
            for s in tcpr.findall(qn("w:shd")):
                tcpr.remove(s)
            if i % 2:
                prev = [el for el in tcpr if el.tag in {qn(t) for t in _SHD_BEFORE}]
                new = copy.deepcopy(tpl)
                if prev:
                    prev[-1].addnext(new)
                else:
                    tcpr.insert(0, new)


def _match_size(row, ref_row):
    """Give every run in `row` that has no size of its own the size of `ref_row`'s runs.

    V4 rows T38 R12, T39 R7-R8 and T40 R12 carry no w:sz, so they render at the 11 pt
    default beside the 9 pt rows around them."""
    ref = next((r.font.size for c in unique_cells(ref_row) for p in c.paragraphs
                for r in p.runs if r.font.size is not None), None)
    if ref is None:
        raise EditError("match_size: reference row has no run size")
    for c in unique_cells(row):
        for p in c.paragraphs:
            for r in p.runs:
                if r.font.size is None:
                    r.font.size = ref


# --------------------------------------------------------------------------- 9 Edge Cases

def _edge_cases(ctx):
    set_text(ctx.P(519, "The app must handle these situations"),
             '''The app and backend must handle these situations without crashing or showing incorrect state. Hubs in the field may still run 2.1.3: read gateway.fw and handle both versions. Where a row says "since 2.1.4", it also gives the 2.1.3 behaviour. A leak sensor is identified by sensor_id, the valve by valve_id.''')
    ctx.note("s7-01: 9 intro names both firmware versions, gateway.fw and the identity keys sensor_id / valve_id")

    t = ctx.T(37, "Scenario")
    body = ctx.row(37, 1, "Hub offline")

    # R1 Hub offline (s7-02; offline gate per canon C14, banner per C24, local protection per C21)
    _set_row(ctx, 37, 1, "Hub offline", [
        None,
        '''IoT Hub reports the device disconnected, or the backend has received neither a snapshot nor a lifecycle message for longer than the larger of 2 x snapshot_interval_s and 5 minutes (10 minutes at the default 300 s). The hub's MQTT keep-alive is 60 s, so after a power cut or a silent network loss IoT Hub reports the hub disconnected only after about 90 s. On a 2.1.3 hub with no devices, which sends no snapshot and no lifecycle, only the IoT Hub connection state applies.''',
        '''Show "Hub Offline". Grey out data. Block new commands (controls disabled). If the last known state has a leak, keep the "Leak Detected" banner, marked "Hub offline - last known state". Ignore a disconnect followed by a reconnect within 60 s: every hub reconnects about every 18 hours to renew its 24-hour SAS token. While the router or the internet is down and Wi-Fi is still saved, the hub keeps watching its leak sensors and its valve on its own, closes the valve on a leak when auto-close is on, and keeps up to 16 events to send when it reconnects (see "Events replayed after an outage"). Since 2.1.4 this also holds when the hub restarts while the router is down; a 2.1.3 hub that restarts while Wi-Fi is down does not act on leaks until it is back on Wi-Fi and its cloud start-up has run. When the hub reconnects, its lifecycle restarts the snapshot-age gate; after a restart the first snapshot can follow up to about 150 s later.''',
    ])
    ctx.note("s7-02: T37 'Hub offline' uses the C14 gate (larger of 2 x snapshot_interval_s and 5 min, snapshot or lifecycle), the 90 s keep-alive delay, the 60 s reconnect debounce, the C24 leak banner, local protection while offline (C21) and the lifecycle restarting the gate")

    # New row after R1: Wi-Fi setup (s7-03, reduced to canon C21 wording)
    _add_rows(t, ctx.row(37, 1, "Hub offline"), body, [[
        "Hub in Wi-Fi setup",
        '''The hub's Wi-Fi setup network is open: first setup, or after the 10-second Wi-Fi reset''',
        '''The hub is offline to the cloud: it sends nothing and runs no commands. It reaches the cloud only after it has joined the router and, since 2.1.4, after its setup network has closed, 15 to 60 seconds after it joins (at most 75 s). Show "Hub Offline".''',
    ]])
    ctx.note("s7-03: new T37 row 'Hub in Wi-Fi setup' in canon C21 wording only (offline to the cloud, show Hub Offline); the BLE-scanning statements are left out")

    # R3 Valve BLE disconnected (s7-04; C23 confirmation, C12 refusals, C3.8 reconnect close)
    _set_row(ctx, 37, 3, "Valve BLE disconnected", [
        None,
        '''data.valve.connected=false and state="disconnected". The valve entry then holds only valve_id, state, connected, rating and last_seen_age_s: battery, leak_state, rmleak and fw_version are absent. last_seen_age_s counts the seconds since the link dropped (since 2.1.4; 2.1.3 counts from the valve's last change of value) and is null if the valve has not linked since the hub started. When a linked valve drops, its rating is "warning" for 3 minutes (system_health.reason "Valve disconnected"), then "critical" ("Valve offline"), and the hub sends device_offline, unless the valve flood probe was wet when the link dropped: the valve then reads "critical" at once as a leak (system_health.reason "Leak detected: valve, Valve offline") and gets no device_offline. A valve not linked since the hub started gets no device_offline (see "Hub just restarted").''',
        '''Show "Disconnected" on the valve card, "--" for the battery and the last known lock state, greyed (APP-FR-022). Keep the valve's last known leak state in the leak banner (APP-FR-036). valve_open and valve_close still get cmd_ack "ok": the hub holds the newest one and sends it when the valve links again (a hub restart drops it). Show "Waiting for valve" until valve_state_changed or a snapshot shows the new state (APP-FR-055); if neither does within 60 s of the ack, show the state from the latest snapshot with "The valve has not confirmed yet. Check that it has power and is in range." Since 2.1.4, valve_open is refused with the RMLEAK error while a leak incident is latched and no override window is open, even with the valve disconnected; a 2.1.3 hub acks it "ok" and holds the open for the valve's reconnect. override_enable waits up to 10 s for the valve, then fails with "The valve isn't responding. Check its power and connection, then try again." If a leak source is still wet when the valve links again, auto-close is on and no override window is active, the hub locks and closes the valve, whatever the trigger mask, and sends auto_close with "cause":"reconnect" (no source_type, no location) unless the valve was already closed and locked. During an override window a reconnect closes nothing.''',
    ])
    ctx.note("s7-04: T37 'Valve BLE disconnected' lists the keys a disconnected valve carries, the warning-then-critical rating and device_offline, held commands with 'Waiting for valve' and the 60 s confirmation rule (C23), the 2.1.4 latched-incident refusal, the 10 s override_enable wait and the reconnect auto_close")
    ctx.note("global fix completeness-03, completeness-06: 'Valve BLE disconnected': no device_offline for a valve that drops while its flood probe is wet (app_ble_valve.c:1902-1923); the card shows \"--\" for the battery as APP-FR-022 says, and the banner keeps the last known leak state")
    ctx.note("review R7-01: the reconnect close in T37 'Valve BLE disconnected', T37 'Valve closes for an unchecked source' and A.3 'Trigger mask check' needs no override window active (rules_engine.c:1567-1580), as 4.8.1 states")

    # New rows after R3: valve readings not in yet (s7-05), valve battery critical (s7-06)
    _add_rows(t, ctx.row(37, 3, "Valve BLE disconnected"), body, [
        [
            "Valve readings not in yet",
            '''data.valve.connected=true and state="unknown", battery=null, fw_version=null: the link is up but the hub has not read the valve yet (battery is null since 2.1.4; a 2.1.3 hub also sends state "unknown" here, but battery 0, and leaves "unknown" as soon as it has read the position, before the battery and rmleak). leak_state and rmleak then read false, even on a locked valve.''',
            '''Show "Connecting..." on the valve card. Ignore leak_state and rmleak until state is "open" or "closed", and keep the last known lock state. From a 2.1.3 hub, treat a valve battery of 0 as unknown.''',
        ],
        [
            "Valve battery critical",
            '''Since 2.1.4, a valve battery at 10 % or less makes valve.rating and system_health.rating "critical" (reason "Valve battery critical"), and the hub sends a snapshot within seconds. No health event is sent for it. valve_open and valve_set_state "open" are refused with "Valve battery critical (≤10 %): the valve will not open. Replace the batteries." Closing is never refused for battery. At 11-20 % the valve reads "warning" ("Valve battery low"). A 2.1.3 hub never rates the valve critical for its battery (20 % or less is "warning", 21-35 % "good") and acks valve_open "ok". On any hub version the valve itself closes at 10 % or less and will not open: valve_state_changed "closed" arrives with no command, no auto_close and no leak (if the valve is not linked, the next snapshot shows it closed).''',
            '''Raise the battery alert and its push from the snapshot (valve.battery and valve.rating; APP-FR-117), not from health events. Show the error detail verbatim, including the "≤" character (U+2264, sent as UTF-8). Show the valve closed with the battery message, not as a leak or an auto-close, and disable "Open with 24h Override" (APP-FR-043).''',
        ],
    ])
    ctx.note("s7-05: new T37 row 'Valve readings not in yet' (state unknown, battery and fw_version null since 2.1.4; 2.1.3 battery 0)")
    ctx.note("s7-06: new T37 row 'Valve battery critical' (2.1.4 bands, snapshot not health event, refusal text with U+2264; 2.1.3 bands and ok ack)")
    ctx.note("global fix completeness-04: 'Valve battery critical': the valve closes itself at 10 % or less on any hub version (valve FW 2.2.0 app_main.c readBatteryADC), so a close with no command or leak is expected; the override is disabled (APP-FR-043)")
    ctx.note("review R7-02: 'Valve readings not in yet': a 2.1.3 hub also sends state unknown there, the only 2.1.3 difference is battery 0 (master app_ble_valve.c:1409-1412)")

    # R4 Sensor not heard yet (s7-07; canon C5)
    _set_row(ctx, 37, 4, "Sensor never seen", [
        "Sensor not heard yet",
        '''The sensor has not reported since the hub started or since a provision added it: connected=false, last_seen_age_s=null, battery=null, rssi=null, fw_version=null, leak_state=false, rating="critical". The hub leaves it out of system_health.rating for its first 10 minutes. Meanwhile system_health.reason reads "Syncing - waiting for N device(s)" when nothing else is wrong, or ends with "syncing N device(s)". After that the sensor counts as offline ("N sensor(s) offline"). No device_offline is sent for a sensor that was never heard.''',
        '''Show "--" for live data and a grey "Not heard yet" badge, not "Offline", during its first 10 minutes (APP-FR-086). Always show sensor_id, location.code and location.label.''',
    ])
    ctx.note("s7-07: T37 'Sensor never seen' becomes 'Sensor not heard yet': sensor_id (not device_id), every value such a sensor carries, the 10-minute roll-up exclusion, the Syncing reasons, no device_offline, the 'Not heard yet' badge (APP-FR-086)")

    # R5 Sensor removed (s7-08; C5 2.1.3 survivor reset, C11 no 2.1.3 twin report)
    _set_row(ctx, 37, 5, "Sensor removed", [
        None,
        '''After the decommission cmd_ack "ok", the next snapshot (within about 5 s) no longer lists the sensor. Since 2.1.4 a twin report follows with ble_leak_sensor_count one lower, and the other devices keep their values. A 2.1.3 hub sends no twin report after a single-device decommission, and every remaining leak sensor reads null battery, rssi and last_seen_age_s, with "Syncing" in system_health.reason, until it is next heard. A device_offline or device_recovered held on the hub for the removed sensor can still arrive after the ack.''',
        '''Remove it from the UI. No error. Ignore health events for a device the hub no longer lists (APP-NF-017). On a 2.1.3 hub, keep showing the other sensors' last known values while they read null after a removal.''',
    ])
    ctx.note("s7-08: T37 'Sensor removed' gives the snapshot and twin confirmation, the 2.1.3 survivor wipe and missing twin report, and late health events (APP-NF-017)")
    ctx.note("review R7-03: the 2.1.3 survivor reset after a single decommission covers the other leak sensors, not the valve (master app_iothub.c reseed_valve_health_if_connected); T37 'Sensor removed' and TC-018")

    # New rows after R5: no valve / no devices (s7-09, canon C8), last device removed (s7-10, C3.7)
    _add_rows(t, ctx.row(37, 5, "Sensor removed"), body, [
        [
            "Hub with no valve, or no devices",
            '''With no valve provisioned, data.valve is {}; a 2.1.3 hub sends {"state":"disconnected","connected":false} instead or, while it is linked to a nearby valve (section 2.4), that valve's valve_id and readings. With no devices at all, ble_leak_sensors and lora_sensors are [] and system_health reads "excellent" with reason "No devices provisioned"; since 2.1.4 such a hub still sends a lifecycle at each connect, twin reports and snapshots. A 2.1.3 hub that is not provisioned (provisioned false), such as one emptied by a decommission, sends no lifecycle and no snapshot, only twin reported. A 2.1.3 hub left with no devices by a provision stays provisioned and keeps sending lifecycle and snapshots, with "valve": {"state": "disconnected", "connected": false} and empty sensor arrays. Since 2.1.4 a leak on a hub with no valve sends leak_detected and no auto_close; 2.1.3 also sent auto_close with rmleak_asserted=false.''',
            '''Treat the valve as present only when data.valve has a valve_id and, on a 2.1.3 hub, twin reported valve_id is not null; hide the valve controls otherwise. Decide "no devices" from the device lists (no valve_id and both sensor arrays empty; in lifecycle and twin reported, no valve_id or valve_id null and both counts 0), not from the provisioned flag. Show "No devices set up yet", not "All Clear". For a 2.1.3 hub that sends no snapshot or lifecycle, use only the IoT Hub connection state for online and offline: do not show "Hub Offline" just because no snapshot or lifecycle arrives.''',
        ],
        [
            "Last device removed",
            '''Since 2.1.4, when a decommission or a provision leaves the hub with no device, the hub resets its rules to auto_close_enabled=true and trigger_mask=7 and clears any leak latch and override window, before the cmd_ack and with no event. Up to 2.1.3 such a hub kept its last rules: until its next restart when a decommission emptied it, and across restarts when a provision emptied it (it stays provisioned). It also kept any leak latch and override window, across restarts.''',
            '''Refresh the settings screen and the override banner from the next snapshot (data.rules, override_active). A provision sent afterwards is applied on top of these defaults. To reuse a 2.1.3 hub for a new installation or owner, send decommission target "all" (section 6.6).''',
        ],
    ])
    ctx.note("s7-09: new T37 row 'Hub with no valve, or no devices' (valve {}, the empty-hub snapshot, 2.1.3 silence, no auto_close without a valve; 'no devices' decided from the device lists, not the provisioned flag, per canon C8)")
    ctx.note("review R7-07: 'Hub with no valve, or no devices': a 2.1.3 empty hub is not shown offline just because no snapshot or lifecycle arrives")
    ctx.note("s7-10: new T37 row 'Last device removed' (2.1.4 rules reset to 7, latch and override cleared; 2.1.3 kept its rules until restart)")
    ctx.note("global fix INT-09/XC-03, XC-11, XC-05, XC-12: 'Hub with no valve, or no devices' and 'Last device removed': 2.1.3 goes silent when not provisioned, not when empty; it kept its rules until restart after a decommission and across restarts after a provision, with its latch and window; the 2.1.3 nearby-valve guard; reuse a 2.1.3 hub with decommission all")

    # R6 Duplicates (s7-11; canon C10, C12; APP-FR-106 renamed to APP-NF-013)
    _set_row(ctx, 37, 6, "Duplicate cmd_ack", [
        "Duplicate messages and duplicate cmd_ack",
        '''The hub delivers every message at least once, so the same message can arrive twice. Since 2.1.4 an event whose publish meets a dropping connection is kept and sent again after the reconnect (a 2.1.3 hub did not keep it), and the MQTT client can deliver its own copy too, so the second copy can arrive after newer messages. The copies are byte-identical, ts included. Separately, the hub does not de-duplicate commands by id: if IoT Hub delivers a C2D message again, the hub runs it again and sends a second cmd_ack with the same id, whose ts or status can differ.''',
        '''Backend: drop a message only when its body is byte-identical to one already received (compare the raw body, or its SHA-256), keep the first copy, and never de-duplicate on a subset of fields such as gateway.id + ts + event (APP-NF-013). App: the first cmd_ack for a correlation id resolves the command; ignore later acks with that id and take the state from the next snapshot.''',
    ])
    ctx.note("s7-11: T37 'Duplicate cmd_ack' becomes 'Duplicate messages and duplicate cmd_ack': at-least-once delivery since 2.1.4, the byte-identical dedupe rule (APP-NF-013), commands not de-duplicated by id")

    # New row after R6: replayed and out-of-order events (s7-12; canon C11, C18)
    _add_rows(t, ctx.row(37, 6, "Duplicate cmd_ack"), body, [[
        "Events replayed after an outage, or out of order",
        '''After a reconnect the hub first sends the events it kept while offline (up to 16, oldest first, each with its original ts), then the lifecycle, then new events. A cmd_ack for a command that arrives during that replay is sent at once and can come before the lifecycle. Since 2.1.4 the kept events include those raised before the first clock sync after a power cut (2.1.3 discarded them). Since 2.1.4 an auto_close can also arrive after the same sensor's leak_cleared (a leak report the hub evaluated a moment late): it reports that the valve was closed, not that the leak came back.''',
        '''Order history by ts, never by arrival (APP-NF-014). Match an auto_close to its leak by sensor_id or valve_id and ts, not by position. An event older than the latest snapshot goes into history only and does not change the displayed state. When an event's ts is more than 5 minutes older than its arrival, its push adds the time, for example "Water leak detected in Kitchen at 14:05" (APP-FR-111). Take the current state from the first snapshot after the reconnect.''',
    ]])
    ctx.note("s7-12: new T37 row 'Events replayed after an outage, or out of order' (replay before the lifecycle, pre-sync events since 2.1.4, late auto_close, order by ts, late-push wording per canon C18)")

    # R7 Command retry (s7-13 with canon C15: expiry on the four valve-actuating commands only)
    _set_row(ctx, 37, 7, "Command retry", [
        None,
        "The user sends the same command again, for example after a timeout",
        '''Use a new correlation id for every send. The hub does not check a command's age and does not de-duplicate commands by id: it runs every C2D message it receives, when it receives it, so a message still queued in IoT Hub runs when the hub reconnects. The backend therefore sets a 30-second expiry on valve_open, valve_close, valve_set_state and override_enable (the app's 20 s ack timeout plus 10 s for the backend and IoT Hub; section 5.5); a command the app reported as timed out can still run for at most about 10 s after that. Every other command keeps the IoT Hub default time-to-live. valve_open or valve_close on a valve already in that state gets cmd_ack "ok" and changes nothing. override_enable during an active window starts the window again at 24 hours.''',
    ])
    ctx.note("s7-13: T37 'Command retry' states that the hub runs queued C2D messages late and gives the C15 rule: a 30 s expiry on valve_open, valve_close, valve_set_state and override_enable only")

    # New row after R7: malformed command (s7-v04 in canon C19 wording)
    _add_rows(t, ctx.row(37, 7, "Command retry"), body, [[
        "Malformed command",
        '''A C2D message that is not a valid command envelope: not a JSON object, schema not exactly "eflostop.cmd" (or the legacy "eflostop.cmd.v1"), or no cmd string''',
        '''The hub sends no cmd_ack, so the app times out after 20 s (APP-FR-053). The message is not always ignored: the hub searches its whole text, in any letter case, for the old plain-text commands DECOMMISSION_VALVE, DECOMMISSION_LORA:, DECOMMISSION_BLE:, DECOMMISSION_ALL, VALVE_OPEN, VALVE_CLOSE, RULES_CONFIG:, SENSOR_META:, LEAK_RESET and OVERRIDE_CANCEL, and runs the one it finds, again with no cmd_ack. So a command with a misspelt schema and cmd "valve_open" opens the valve, and DECOMMISSION_ALL anywhere in the text (an id, a label) resets the hub to its factory state. A valid JSON object with no cmd string and none of these words is applied as a provision payload. The backend shall build every command with a JSON library as an exact eflostop.cmd envelope (section 5.3), and shall never send user text, or anything else, as a C2D body outside such an envelope. Hub firmware 2.1.3 and 2.1.4 behave the same.''',
    ]])
    ctx.note("s7-v04: new T37 row 'Malformed command' (no cmd_ack, legacy keyword scan, provision fallback, strict envelope rule) in canon C19 wording")

    # R8 Hub just restarted (s7-14; canon C5, C9)
    _set_row(ctx, 37, 8, "Boot sync window", [
        "Hub just restarted",
        '''A lifecycle shows that the hub restarted (its ts minus gateway.uptime_s is at or after the previous lifecycle's ts, section 5.2.1), and data.reset_reason says why. Devices not heard yet read connected=false, last_seen_age_s=null and rating="critical" in their own entries, but the hub leaves them out of system_health.rating for up to 10 minutes (a sensor) or 3 minutes (the valve); system_health.reason then reads "Syncing - waiting for N device(s)" or ends with "syncing N device(s)". The first snapshot is "boot" once every device has been heard, or "fast" as soon as the valve link is ready (at the latest about 150 s after the connect); a "boot" snapshot then follows once every device has been heard or 3 minutes after the restart. No device_offline is sent for a device not heard since the restart.''',
        '''Show a "Hub starting up - checking devices" note while gateway.uptime_s is below 600. Show unheard devices as "Not heard yet" (APP-FR-086), not "Offline". Drive the overall badge from system_health.rating, not from the per-device ratings.''',
    ])
    ctx.note("s7-14: T37 'Boot sync window' becomes 'Hub just restarted': 10 min (sensor) / 3 min (valve) roll-up exclusion instead of 2 min, the fast and boot snapshots, no device_offline for unheard devices, a start-up note while uptime_s < 600")
    ctx.note("review R7-11, global fix INT-08: 'Hub just restarted' tells a restart from a reconnect with ts minus gateway.uptime_s (5.2.1), not a small uptime_s; reset_reason is in every lifecycle")

    # R9 Clock not synced (s7-15; canon C10 exception, C13)
    _set_row(ctx, 37, 9, "Time not synced", [
        "Clock not yet synced on the hub",
        '''The hub never publishes a ts below 1704067200 (2024-01-01 UTC): it connects to the cloud only after its clock is set. Since 2.1.4, events raised before that (for example a leak_detected, its auto_close or a health event) are kept and sent after the first connect, in order, with ts corrected from the hub's uptime, so they arrive with a ts earlier than that connection's lifecycle. 2.1.3 discarded them. If writing the corrected ts to flash failed, two copies of such an event can differ in ts by about 1 s and in nothing else.''',
        '''Order history by ts, not arrival, and accept events whose ts is earlier than the lifecycle they follow. Two messages identical except for a ts at most 1 s apart may be treated as one (APP-NF-013). If a ts below 1704067200 ever arrives (possible only from a hub moved back from 2.1.4 to 2.1.3 with such an event still kept), store the event with its arrival time marked "time unknown" and do not use it to change the current state.''',
    ])
    ctx.note("s7-15: T37 'Time not synced' becomes 'Clock not yet synced on the hub': no ts below 1704067200 is published; since 2.1.4 pre-sync events arrive late with a corrected ts (2.1.3 discarded them)")

    # R10 Override + reboot (s7-16)
    _set_row(ctx, 37, 10, "Override + reboot", [
        None,
        "Hub restarts during the 24-hour override",
        '''The window is stored on the hub and survives a restart: the first snapshot after the restart has override_active=true and the correct override_remaining_s. Since 2.1.4, a window started before the hub's clock synced carries on through the sync if the hub has not restarted since, and snapshots leave out expires_ts until up to about 30 s after the sync; a 2.1.3 hub ended such a window at the clock sync. If the hub restarts before its clock has synced, such a window ends at the next clock sync, on 2.1.4 as on 2.1.3. A window that ran out while the hub was restarting can end without a water_access_override_expired event. Since 2.1.4 a hub that restarts with no devices clears the window; a 2.1.3 hub keeps it. Drive the banner and countdown from the snapshot's override_active and override_remaining_s, not only from events.''',
    ])
    ctx.note("s7-16: T37 'Override + reboot' adds the pre-sync window (since 2.1.4; 2.1.3 ended it at the sync), the silent expiry during a restart and the empty-hub clear; the snapshot drives the banner")
    ctx.note("global fix XC-06: 'Override + reboot': the boot-time empty-hub clear is since 2.1.4 (HEAD app_iothub.c:3718-3750); 2.1.3 reloads the window at every boot (master rules_engine.c:534)")
    ctx.note("review R7-04: 'Override + reboot': a pre-sync window carries on through the sync only if the hub has not restarted; one restored before a sync ends at the next sync on both versions (rules_engine.c:1803-1825)")

    # R11 Large sensor count (s7-17; LoRa read-only per canon C20)
    _set_row(ctx, 37, 11, "Large sensor count", [
        None,
        "Up to 16 BLE leak sensors and up to 16 LoRa leak sensors provisioned",
        '''The UI must handle up to 16 BLE leak sensors and up to 16 LoRa leak sensors (shown read-only) in the list, plus the valve. data.lora_sensors is always present ([] when there are none). A snapshot of a full hub is about 8 KB, and can reach about 10 KB with long labels.''',
    ])
    ctx.note("s7-17: T37 'Large sensor count' adds up to 16 LoRa leak sensors (read-only, C20), data.lora_sensors always present, a full snapshot about 8 KB")
    ctx.note("integrator: 'Large sensor count' gives the same snapshot size as 7.1: about 8 KB, up to about 10 KB with long labels (app_iothub.c:2244-2246; HANDOFF.md section 15 worst case)")

    # New rows after R11: rules edge cases (tm-36)
    _add_rows(t, ctx.row(37, 11, "Large sensor count"), body, [
        [
            "Settings changed during a leak",
            "The user changes auto-close or a leak source while a source is wet",
            '''Show the new values from the next snapshot. Do not expect the valve to move at once. A source switched on closes the valve at its next leak report (within about 5 minutes for a BLE leak sensor) or at the next valve reconnect. Switching a source or auto-close off never opens the valve or clears RMLEAK.''',
        ],
        [
            "Valve closes for an unchecked source",
            '''A source whose checkbox is clear is wet, auto-close is on, and the valve reconnects outside an override window, the override window ends or override_cancel is sent''',
            '''The hub asserts RMLEAK and closes the valve, whatever the trigger mask (section 4.8.1). The reconnect close sends auto_close with cause "reconnect"; the other two send no auto_close. Show the valve locked as for any auto-close.''',
        ],
        [
            "Rules missing from a snapshot",
            "data.rules absent (the hub could not read its rules within 1 s)",
            '''Keep the last known auto-close and trigger mask values. Do not show them as off.''',
        ],
    ])
    ctx.note("tm-36: new T37 rows 'Settings changed during a leak', 'Valve closes for an unchecked source' (with the reconnect auto_close per canon C3.8) and 'Rules missing from a snapshot'")


# --------------------------------------------------------------------------- 10.1 End-to-end tests

def _tests(ctx):
    t = ctx.T(38, "Test ID")
    body = ctx.row(38, 1, "TC-001")

    set_cell(ctx.cell(38, 1, 2, "Trigger leak on BLE sensor"),
             '''With the valve open and linked, auto-close on and the BLE leak sensors trigger checked (data.rules.trigger_mask has bit 0 set; default 7), wet a BLE leak sensor.''')
    set_cell(ctx.cell(38, 1, 3, "App shows leak banner"),
             '''leak_detected (source_type "ble_leak_sensor", that sensor_id), then auto_close with rmleak_asserted=true, and valve_state_changed with valve_state "closed" and rmleak=true. The next snapshot shows valve.state "closed", valve.rmleak=true, and system_health "critical" with a reason starting "Leak detected: " and the sensor's label (its sensor_id when it has no label). App shows the leak banner and "Closed (Locked)". Push received within 10 s.''')
    ctx.note("tm-v04, s7-18: TC-001 steps require the BLE trigger bit, auto-close on and an open, linked valve; the expected result names the events and snapshot fields")

    _set_row(ctx, 38, 2, "TC-002", [
        None, None,
        "With the valve open and connected, send valve_close from the app.",
        '''cmd_ack "ok" (the command was queued) and the card shows "Closing...". Then valve_state_changed with valve_state "closed" and a snapshot with valve.state "closed". Valve card shows "Closed" within 20 s.''',
    ])
    ctx.note("s7-19: TC-002 treats cmd_ack ok as queued and confirms the close with valve_state_changed and the snapshot")

    _set_row(ctx, 38, 3, "TC-003", [
        None, None,
        '''With the valve closed and connected, no leak latched (valve.rmleak=false) and the valve battery above 10 %, send valve_open from the app.''',
        '''cmd_ack "ok" (queued) and the card shows "Opening...". Then valve_state_changed with valve_state "open" and a snapshot with valve.state "open". Valve card shows "Open" within 20 s.''',
    ])
    ctx.note("s7-20: TC-003 preconditions exclude the 2.1.4 refusals (latched leak, battery at 10 % or less); confirmation as TC-002")

    _set_row(ctx, 38, 4, "TC-004", [
        None, None,
        '''After TC-001, dry the sensor and send leak_reset within 10 s of its leak_cleared (2.1.3: within 30 s); use a backend test tool if the app is too slow. Then send valve_open.''',
        '''leak_reset: cmd_ack "ok", rmleak_cleared event, snapshot with valve.rmleak=false; RMLEAK icon gone, Open Valve enabled, valve still closed. valve_open: cmd_ack "ok", valve_state_changed "open"; water restored. If the hub clears RMLEAK by itself first (TC-016), leak_reset still returns "ok" and changes nothing.''',
    ])
    ctx.note("s7-21: TC-004 sends leak_reset after the sensor dries and before the automatic clear (10 s; 30 s on 2.1.3); expected events named")

    _set_row(ctx, 38, 5, "TC-005", [
        None, None,
        "After TC-001, with the sensor still wet, open the valve with the button on the valve. Then feed the backend's ingestion a water_access_override_expired with auto_close_resumed true.",
        '''water_access_override_enabled with trigger "button", remaining_s 86400 and expires_ts; valve_state_changed "open" with rmleak=false; snapshot with override_active=true and override_remaining_s counting down. App shows "24h override active" with countdown; the push uses the APP-FR-116 "button" copy. When the wet sensor next reports (within about 5 minutes), the hub sends auto_close_blocked_override (at most one a minute for the whole hub) and the valve stays open. The backend sends the "Override Ended" push for the fed water_access_override_expired (APP-FR-118).''',
    ])
    ctx.note("s7-22: TC-005 names the button trigger, the override events and snapshot fields, the APP-FR-116 push and auto_close_blocked_override")

    set_cell(ctx.cell(38, 6, 3, "Override banner disappears"),
             '''cmd_ack "ok". auto_close_reenabled event with reason "c2d_command" and previous_remaining_s. Override banner disappears. If any leak source is still wet and auto-close is on, the valve also closes and locks (valve_state_changed closed, valve.rmleak true), whatever the trigger mask; no auto_close event is sent for this close.''')
    ctx.note("tm-v07: TC-006 expected result names auto_close_reenabled and the close-and-lock when a source is still wet")

    _set_row(ctx, 38, 7, "TC-007", [
        None, None,
        '''With a sensor that has reported since the hub last started, remove its battery and wait up to 13 minutes.''',
        '''A device_offline event (category "health", source_type "ble_leak_sensor", that sensor_id, rating "critical", offline_duration_s about 600) arrives 10 to 10.5 minutes after the hub last heard the sensor (a dry sensor reports about every 100 s). The next snapshot shows connected=false, rating "critical" and a reason including "1 sensor offline". Push received. A sensor never heard since the hub started produces no device_offline.''',
    ])
    ctx.note("s7-24: TC-007 waits up to 13 min, expects device_offline (not a health alert) with its fields, and the snapshot reason")

    _set_row(ctx, 38, 8, "TC-008", [
        None, None,
        "Unplug the hub. Wait more than 10 minutes, then plug it back in.",
        '''"Hub Offline" once IoT Hub has reported the device disconnected for more than 60 s (IoT Hub reports it about 90 s after the power cut, so about 150 s after it), and in any case once neither a snapshot nor a lifecycle has arrived for 10 minutes (the larger of 2 x snapshot_interval_s and 5 minutes, at the default 300 s). Data greyed out and controls disabled (APP-NF-011). After the hub is back: lifecycle, then snapshots; "Hub Offline" clears when IoT Hub reports the hub connected and the lifecycle arrives, without waiting for a snapshot.''',
    ])
    ctx.note("s7-25: TC-008 checks the IoT Hub connection state first, the C14 snapshot-or-lifecycle gate (10 min at the default), the command block and the recovery")
    ctx.note("integrator (s6 cross-module, s6-27): TC-008 shows Hub Offline once IoT Hub has reported the hub disconnected for more than 60 s (the APP-NF-010 debounce)")

    _set_row(ctx, 38, 10, "TC-010", [
        None, None,
        "Change the sensor's location and label from the app.",
        '''cmd_ack "ok"; a snapshot within about 5 s shows the new location.code and location.label; later leak events for that sensor carry them.''',
    ])
    ctx.note("s7-26: TC-010 gives the ack and the 5 s snapshot, and checks later leak events")

    _set_row(ctx, 38, 11, "TC-011", [
        None, None,
        '''Turn "Automatic shutoff" off in Settings. Keep the valve linked. Wet a BLE leak sensor.''',
        '''Command payload {"auto_close_enabled":false}. cmd_ack "ok". Next snapshot: data.rules.auto_close_enabled=false, trigger_mask unchanged. leak_detected and its push arrive. No auto_close. Valve stays open.''',
    ])
    ctx.note("tm-37: TC-011 keeps the valve linked and checks the payload, data.rules, the unchanged mask and the reported leak")

    set_cell(ctx.cell(38, 12, 3, "water_access_override_enabled event"),
             '''cmd_ack "ok" (up to about 10 s if the hub must first reconnect the valve; it can arrive before or after the override event); water_access_override_enabled event with trigger="c2d_command" and remaining_s=86400; valve_state_changed to open with rmleak=false; snapshot override_active=true. Banner and countdown shown.''')
    ctx.note("s7-28: TC-012 allows the 10 s valve reconnect wait and either order of ack and override event")

    _add_rows(t, ctx.row(38, 12, "TC-012"), body, [
        ["TC-013", "Settings: BLE leak sensors trigger off",
         '''On a hub at trigger_mask 7 with auto-close on, uncheck "BLE leak sensors". Then add one more BLE leak sensor through the backend's add-device provision (the complete sensor list, no auto_close_enabled and no rules). Keep the valve linked for the whole test. Wet a BLE leak sensor.''',
         '''Command payload {"trigger_ble_leak":false}. cmd_ack "ok". Next snapshot and twin reported: trigger_mask=6. After the provision, data.rules still shows auto_close_enabled=true and trigger_mask=6 (APP-FR-097). leak_detected with source_type "ble_leak_sensor" and its push arrive. No auto_close. Valve stays open.'''],
        ["TC-014", "Settings: valve flood probe trigger off",
         '''On a hub at trigger_mask 7 with auto-close on, uncheck "Valve flood probe". Keep the valve linked. Wet the valve's flood probe.''',
         '''Command payload {"trigger_valve_flood":false}. trigger_mask=3. leak_detected with source_type "valve" arrives. The hub sends no auto_close. Whether the valve acts on its own probe is valve firmware behaviour and is not checked here.'''],
        ["TC-015", "Settings: hidden LoRa bit kept",
         '''On a hub at trigger_mask 7, uncheck and re-check "BLE leak sensors", then uncheck and re-check "Valve flood probe".''',
         '''Each command carries only trigger_ble_leak or trigger_valve_flood. Snapshots show 6, 7, 3, 7. Bit 1 stays set throughout.'''],
        ["TC-016", "Interlock auto-clear",
         "After TC-001, dry the sensor and wait.",
         '''leak_cleared, then about 10-12 s later rmleak_auto_cleared with clear_after_seconds 10, and a snapshot with valve.rmleak=false. system_health leaves "warning" ("Leak interlock latched"). The valve stays closed; Open Valve becomes available without Leak Reset. (2.1.3: 30-60 s, clear_after_seconds 30.)'''],
        ["TC-017", "Sensor back online",
         "After TC-007, put the battery back.",
         '''On the sensor's first report: device_recovered (category "health", that sensor_id, prev_rating "critical"), and a snapshot with connected=true and its battery and rssi.'''],
        ["TC-018", "Remove one sensor",
         "On a hub with a valve and two or more sensors, decommission one sensor.",
         '''cmd_ack "ok"; a snapshot within about 5 s without that sensor; twin reported ble_leak_sensor_count one lower. The other sensors and the valve keep their battery, rssi and last_seen_age_s (on 2.1.3 the other sensors read null and "Syncing" until next heard, and no twin report follows the decommission). Any later health event for the removed sensor_id is ignored.'''],
        ["TC-019", "Remove the valve, then the last sensor",
         "Decommission the valve, then every sensor.",
         '''After the valve: "valve":{}, twin reported valve_id null, valve controls hidden. After the last sensor: ble_leak_sensors [], system_health "excellent" with reason "No devices provisioned", data.rules back to auto_close_enabled=true and trigger_mask=7, override_active=false; heartbeat snapshots continue. App shows "No devices set up yet". (Since 2.1.4. On 2.1.3 the valve entry reads {"state":"disconnected","connected":false}, no twin report follows a single decommission, and the emptied hub falls silent.)'''],
        ["TC-020", "Leak while the hub is offline",
         "Turn the router off. Wet a sensor. After 1 minute dry it, then turn the router back on.",
         '''The valve closes at once (check at the valve). After the hub reconnects, the events of the episode (leak_detected, auto_close, valve_state_changed, leak_cleared, rmleak_auto_cleared) arrive before the lifecycle with their original ts, then snapshots. History shows them at their ts. The backend stores each message once. (On 2.1.3 the hub does not rejoin the router on its own after this outage: unplug it and plug it back in once the router is back, then check the same messages.)'''],
        ["TC-021", "Hub restart during override",
         "During TC-005's override, power-cycle the hub.",
         '''After it reconnects: lifecycle with reset_reason "power_on"; the first snapshot has override_active=true, override_remaining_s continuing from before (not 86400) and the same expires_ts. Unheard devices show "Not heard yet" and do not raise system_health for up to 10 minutes; no device_offline is sent for them.'''],
        ["TC-022", "Snapshot interval",
         "Set the twin desired snapshot_interval_s to 600, then to 30, then restart the hub.",
         '''Twin reported snapshot_interval_s becomes 600, and heartbeat snapshots (data.reason "heartbeat") come about 600 s after the last snapshot when nothing else happens. 30 is rejected: reported stays 600. After the restart, reported is still 600.'''],
    ])
    ctx.note("tm-38: new tests TC-013 (BLE leak sensors trigger off; 'BLE leak sensor' per canon C2 terminology), TC-014 (valve flood probe trigger off), TC-015 (LoRa bit kept)")
    ctx.note("s7-29: new tests TC-016 (interlock auto-clear, 10 s; 30 s on 2.1.3) and TC-017 (device_recovered)")
    ctx.note("s7-30: new tests TC-018 (remove one sensor; 2.1.3 survivor wipe and no twin report) and TC-019 (remove the valve and the last sensor; 2.1.3 behaviour per canon C8)")
    ctx.note("global fix INT-07, XC-08, XC-14: TC-005 checks the APP-FR-118 push; TC-013 checks that an add-device provision keeps the user's mask (APP-FR-097, P0); TC-020 says a 2.1.3 hub must be power-cycled after the outage (2.6 Wi-Fi setup)")
    ctx.note("s7-31: new tests TC-020 (leak while offline, replay), TC-021 (restart during override) and TC-022 (snapshot interval)")


# --------------------------------------------------------------------------- 10.2 Negative tests

def _negative_tests(ctx):
    t = ctx.T(39, "Test ID")
    body = ctx.row(39, 1, "TC-N01")

    _set_row(ctx, 39, 1, "TC-N01", [
        None, None,
        '''Send valve_set_state (a) without a "state" field, (b) with "state":"half".''',
        '''cmd_ack error with detail (a) "missing 'state' field (expected "open" or "closed")", (b) "invalid state value (expected "open" or "closed")". Valve unchanged.''',
    ])
    ctx.note("s7-32: TC-N01 uses the exact error.detail texts and adds the invalid-value case")

    _set_row(ctx, 39, 2, "TC-N02", [
        None, None,
        'Send decommission with target "all".',
        '''cmd_ack "ok", then a final snapshot with data.reason "decommission" showing no devices ("valve":{}, reason "No devices provisioned"). The hub restarts, normally 3 to 5 s after the ack and at most about 30 s after it. After it reconnects: lifecycle with reset_reason "software", provisioned=false and ble_leak_sensor_count 0, then snapshots of the empty hub; the app shows "No devices set up yet". gateway.name is gone and the snapshot interval is back to 300 s, provided the backend set the twin's desired hub_name and snapshot_interval_s to null: the hub applies desired properties again at every connect. On 2.1.3 no lifecycle or snapshot follows the restart, because a 2.1.3 hub with no devices sends none; its twin reported shows provisioned false, valve_id null and both sensor counts 0.''',
    ])
    ctx.note("s7-33: TC-N02 adds the final 'decommission' snapshot, the 3-30 s restart, the lifecycle of the empty hub and the 'No devices set up yet' state (as APP-FR-123), the desired hub_name / snapshot_interval_s reset (canon C16) and the 2.1.3 behaviour")

    _set_row(ctx, 39, 3, "TC-N03", [
        None, None,
        '''(a) With the hub shown offline, try a valve command. (b) With the hub online, send valve_close and cut the hub's power within a second.''',
        '''(a) Controls are disabled and no command is sent (APP-NF-011). (b) After 20 s the APP-FR-053 timeout message is shown and the UI keeps the previous state. The backend set a 30 s expiry on the command (section 5.5), so a hub that reconnects later does not run it. If the hub ran it before the power cut, or reconnects within those 30 s, its cmd_ack arrives late: the backend logs it against the command and the app UI does not change on it (APP-NF-018).''',
    ])
    ctx.note("s7-34: TC-N03 no longer sends a command to an offline hub (APP-NF-011); it tests the timeout, the 30 s valve-command expiry (canon C15) and the late ack (APP-NF-018)")

    _set_row(ctx, 39, 4, "TC-N04", [
        None, None,
        '''Scan (a) a valve or leak sensor label (type=valve or type=sensor), (b) a code whose id is not GW- followed by 12 hex characters.''',
        '''(a) "This is not a hub QR code." (b) "Invalid hub ID format." No hub is linked.''',
    ])
    ctx.note("review R7-08: TC-N04 matches the QR rules of canon C17 (valve and sensor labels rejected with 'This is not a hub QR code.')")

    _set_row(ctx, 39, 6, "TC-N06", [
        None,
        "Null and missing values",
        '''View a sensor not heard since the hub started (connected=false, rating "critical", last_seen_age_s, battery, rssi and fw_version null, leak_state false). Also view the valve with its link down (battery, leak_state, rmleak and fw_version absent) and while it reads state "unknown" (battery null).''',
        '''Live fields show "--"; the sensor shows "Not heard yet" with its sensor_id and location; the valve card shows "Disconnected" or "Connecting...". No crash.''',
    ])
    ctx.note("s7-35: TC-N06 lists the real values of an unheard sensor (not 'all null') and adds the valve's absent keys and its 'unknown' state")

    _set_row(ctx, 39, 7, "TC-N07", [
        None, None,
        '''With the valve connected, send override_enable (a) with no leak incident latched, the valve's RMLEAK clear and no override window open, (b) with water at the valve's own flood probe (valve flood bit set in trigger_mask, the default). (c) Switch the valve off and send override_enable.''',
        '''cmd_ack error with the exact detail: (a) "No active leak to override. Use the normal Open Valve control." (b) "Water detected at the valve. It can't be opened remotely until the valve area is dry." (c) after about 10 s, "The valve isn't responding. Check its power and connection, then try again." UI unchanged.''',
    ])
    ctx.note("s7-36: TC-N07 keeps the valve connected for (a) and (b), states their preconditions and adds (c), the unreachable valve")

    _set_row(ctx, 39, 8, "TC-N08", [
        None, None,
        '''With a sensor leaking and the valve auto-closed (rmleak=true): (a) send leak_reset; (b) send valve_open; (c) switch the valve off and send valve_open again.''',
        '''(a) cmd_ack error "A leak is still active. Fix the leak first, or use override to open the valve during a leak." (b) and (c) cmd_ack error "Valve is locked after a leak (RMLEAK). Clear it with leak_reset first, or use override to open the valve during a leak." Valve stays closed and locked; no valve_state_changed transient, and nothing opens the valve when it is switched back on. (c) is since 2.1.4; a 2.1.3 hub acks "ok" and holds the open for the valve's reconnect.''',
    ])
    ctx.note("s7-37: TC-N08 adds (c), valve_open refused while the incident is latched and the valve is disconnected (since 2.1.4)")

    _add_rows(t, ctx.row(39, 8, "TC-N08"), body, [
        ["TC-N09", "rules_config without payload",
         'Send rules_config with "payload": null.',
         '''cmd_ack status "error", error.code "rules_config", error.detail "rules config update failed". data.rules unchanged.'''],
        ["TC-N11", "Valve battery critical",
         '''With the valve open and no leak latched, lower the valve's supply to 10 % battery or less (bench supply). Then send valve_open.''',
         '''The valve closes itself: valve_state_changed "closed" with no command, no auto_close and no leak. A snapshot within seconds shows valve.rating "critical" and system_health "critical" with reason "Valve battery critical"; no device_offline. The backend sends one "Valve Battery Critical" push; a later snapshot still at 10 % or less, or a valve reconnect, sends no second one (APP-FR-117, on 2.1.3 too). valve_open: cmd_ack error "Valve battery critical (≤10 %): the valve will not open. Replace the batteries.", shown verbatim including the "≤" character. (Since 2.1.4; a 2.1.3 hub rates it "warning" and acks "ok", and the valve stays shut.)'''],
        ["TC-N12", "No valve set up",
         '''On a hub with sensors and no valve, send valve_open, valve_close and override_enable, then wet a sensor.''',
         '''Each command: cmd_ack error "No valve is set up for this hub." The leak sends leak_detected and no auto_close. The snapshot shows "valve":{}. Valve controls are hidden. (Since 2.1.4. A 2.1.3 hub acks valve_open and valve_close "ok", sends auto_close with rmleak_asserted=false and shows a disconnected valve entry.)'''],
        ["TC-N13", "Command for a disconnected valve",
         '''With the valve open, switch it off and wait until the snapshot shows valve.connected=false. Send valve_close. Switch the valve back on.''',
         '''cmd_ack "ok" (queued). The valve card shows "Waiting for valve". When the valve links again, valve_state_changed "closed" arrives and the card shows "Closed".'''],
        ["TC-N14", "Unparseable command",
         '''From a backend test tool, send the C2D text {not json (it must not contain a command name such as VALVE_OPEN, which the hub would run).''',
         '''No cmd_ack is sent. The app times out after 20 s and shows the APP-FR-053 message. Nothing changes on the hub.'''],
        ["TC-N15", "Duplicate message",
         '''Feed the backend's ingestion a captured leak_detected message twice, byte for byte, with a newer snapshot between the two copies. Then feed two cmd_ack messages with the same ts that differ only in data.id and data.cmd.''',
         '''The leak_detected gives one history row and one push. Both cmd_acks are kept, and each resolves its own command (APP-NF-013).'''],
        ["TC-N16", "Late and out-of-order events",
         '''With the latest snapshot showing valve.state "open", feed the backend's ingestion a valve_state_changed "closed" whose ts is older than that snapshot's ts. Then feed a leak_detected whose ts is two hours old.''',
         '''The valve card stays "Open" and the dashboard banner does not change. Both events appear in history at their own ts, not at arrival time. The leak push states the time of the event (APP-FR-111, APP-NF-014).'''],
        ["TC-N17", "Health event for a removed device",
         '''Decommission a BLE leak sensor (cmd_ack "ok"). Then feed the backend's ingestion a device_offline for that sensor_id with a ts from before the decommission.''',
         '''No push. The sensor does not reappear in the app. The backend logs the event (APP-NF-017).'''],
        ["TC-N18", "Late cmd_ack",
         '''Send valve_close from the app to a hub that does not answer (for example a simulated hub). 30 s after the send, feed the matching cmd_ack with status "ok" to the backend's ingestion.''',
         '''After 20 s the app shows the APP-FR-053 timeout message. The late ack is logged with its command, and the app UI does not change on it (APP-NF-018).'''],
    ])
    ctx.note("tm-38: new test TC-N09 (rules_config without payload)")
    ctx.note("global fix completeness-04, INT-07: TC-N11 expects the valve's own close at 10 % or less and checks the APP-FR-117 push")
    ctx.note("s7-38: new tests TC-N11 (valve battery critical), TC-N12 (no valve set up; 2.1.3 behaviour added), TC-N13 (command for a disconnected valve), TC-N14 (unparseable command); renumbered from TC-N10..N13 because V4's revision history cites TC-N10")
    ctx.note("id_map (decision 25): new tests TC-N15 (APP-NF-013 dedupe), TC-N16 (APP-NF-014 order), TC-N17 (APP-NF-017 removed device), TC-N18 (APP-NF-018 late ack)")


# --------------------------------------------------------------------------- 11 Traceability

def _trace(ctx):
    t = ctx.T(40, "App Screen/Feature")

    rows = {
        1: ("Dashboard", ["Dashboard", "APP-FR-021..027", "snapshot (all fields), lifecycle", "(none)",
                          "TC-008, TC-009, TC-019, TC-021"], "s7-39"),
        2: ("Valve Control", ["Valve Control", "APP-FR-051..058", "snapshot (valve), cmd_ack, valve_state_changed",
                              "valve_open, valve_close, valve_set_state",
                              "TC-002, TC-003, TC-N01, TC-N03, TC-N08, TC-N11, TC-N12, TC-N13, TC-N14"], "s7-40"),
        3: ("Leak Handling", ["Leak Handling", "APP-FR-036..042",
                              "leak_detected, leak_cleared, auto_close, rmleak_cleared, rmleak_auto_cleared, auto_close_blocked_override, water_access_override_enabled, water_access_override_expired, auto_close_reenabled; snapshot (valve.rmleak, override_active, override_remaining_s)",
                              "leak_reset, override_cancel",
                              "TC-001, TC-004, TC-005, TC-006, TC-016, TC-020, TC-021"], "s7-41"),
        5: ("Health / Battery", ["Health / Battery", "APP-FR-081..086",
                                 "snapshot (system_health, per-device rating, battery), device_offline / device_recovered events under the health category; a valve battery that turns critical comes only in the snapshot",
                                 "(none)", "TC-007, TC-017, TC-N11"], "s7-42"),
        6: ("Settings", ["Settings", "APP-FR-091..097",
                         "snapshot (data.rules, override), twin reported (auto_close_enabled, trigger_mask, snapshot_interval_s)",
                         "rules_config, provision (setup opt-in), Twin desired snapshot_interval_s",
                         "TC-011, TC-013, TC-014, TC-015, TC-022, TC-N09"], "s7-43, tm-39"),
        7: ("Event History", ["Event History", "APP-FR-101..105, APP-NF-013, APP-NF-014",
                              "All event types, ordered by ts; byte-identical duplicates dropped", "(none)",
                              "TC-001 (verify history entry), TC-020, TC-N15, TC-N16"], "s7-v05"),
        8: ("Push Notifications", ["Push Notifications", "APP-FR-111..115, APP-FR-117, APP-FR-118",
                                   "leak_detected, auto_close, water_access_override_expired, device_offline (health category); snapshot valve.battery for the valve battery push (no event is sent for it)",
                                   "(none)", "TC-001, TC-005, TC-007, TC-020, TC-N11"], "s7-44"),
        11: ("Device Mgmt", ["Device Mgmt", "APP-FR-121..124",
                             'cmd_ack, snapshot (device lists, "valve":{} when no valve), Device Twin reported (ble_leak_sensor_count, valve_id), lifecycle (after decommission all)',
                             "decommission", "TC-018, TC-019, TC-N02"], "s7-47"),
        12: ("Remote Override", ["Remote Override", "APP-FR-043..045, APP-FR-116",
                                 "water_access_override_enabled (+trigger), cmd_ack, snapshot (override_active)",
                                 "override_enable", "TC-012, TC-N07, TC-N12"], "s7-v06"),
    }
    for r, (expect, vals, ids) in rows.items():
        _set_row(ctx, 40, r, expect, vals)
        ctx.note(f"{ids}: T40 '{expect}' row per id_map trace_matrix (new requirement and test IDs, named events)")

    set_cell(ctx.cell(40, 9, 1, "APP-FR-011..016"), "APP-FR-004..009")
    ctx.note("id_map (s2-01): T40 'Onboarding' requirement range APP-FR-011..016 corrected to APP-FR-004..009")
    set_cell(ctx.cell(40, 10, 1, "APP-FR-001..006"), "APP-FR-001..003")
    ctx.note("s7-46: T40 'Auth / Login' requirement range APP-FR-001..006 corrected to APP-FR-001..003")

    rel = add_row(t, ["Offline and Reliability", "APP-NF-010..014, APP-NF-017, APP-NF-018",
                "IoT Hub device-connection state, snapshot age, lifecycle, events replayed from the hub's offline buffer, duplicate and late messages",
                "(none)",
                "TC-008, TC-N03, TC-N06, TC-020, TC-N15, TC-N16, TC-N17, TC-N18"],
            like_row=ctx.row(40, 1, "Dashboard"), after_row=ctx.row(40, 12, "Remote Override"))
    ctx.note("s7-48: new T40 row 'Offline and Reliability' traces APP-NF-010..014, APP-NF-017, APP-NF-018 to their tests")
    add_row(t, ["Security and Logging", "APP-NF-001..006, APP-NF-015, APP-NF-016",
                "All D2C messages (logged with the body hash, APP-NF-016)",
                "All commands (logged with the correlation id, APP-NF-015)",
                "(code review, privacy audit and log review)"],
            like_row=ctx.row(40, 1, "Dashboard"), after_row=rel)
    ctx.note("global fix INT-07: new T40 row 'Security and Logging' traces APP-NF-001..006, 015 and 016, so every requirement is in the matrix; the push row adds TC-005 (APP-FR-118)")

    set_cell(ctx.cell(40, 4, 2, "snapshot (ble_leak_sensors)"), "snapshot (ble_leak_sensors; lora_sensors read-only)")
    ctx.note("review R7-12: T40 'Sensor Detail' traces lora_sensors (read-only), as APP-FR-066 and canon C20 state; IDs and tests unchanged")


# --------------------------------------------------------------------------- Appendix

def _appendix(ctx):
    # A.1 Location codes (s7-49, s7-50)
    _set_row(ctx, 41, 0, "Code String", [
        "Code String (on the wire)", "Hub Enum Value (internal only)", None,
    ])
    ctx.note("s7-49: A.1 header says the code string is the wire value and the enum value is hub-internal")
    insert_after(ctx.T(41, "Code String"),
                 '''The cloud only ever sends and receives the code string; the enum value never appears on the wire. The hub matches the string without regard to case. A location_code it does not recognise is ignored: the stored location is kept and the cmd_ack is still "ok", so send only the strings above. A label is limited to 31 bytes of UTF-8 (an accented letter takes 2 bytes, an emoji 4). The hub cuts a longer label at 31 bytes, which can split a character, so the app shall limit labels to 31 bytes. The valve has no location: leak events from the valve flood probe always carry {"code":"unknown","label":""}, and the snapshot's valve entry has no location.''',
                 like=ctx.P(519, "The app must handle"))
    ctx.note("s7-50: new paragraph after A.1 (string only on the wire, case-insensitive, unknown code ignored with ok, 31-byte label, the valve has no location)")

    # A.2 Health thresholds
    t42 = ctx.T(42, "Threshold")
    b42 = ctx.row(42, 1, "BLE sensor offline timeout")
    _set_row(ctx, 42, 1, "BLE sensor offline timeout", [
        "Sensor offline timeout", None,
        '''A BLE or LoRa leak sensor not heard for 10 minutes becomes critical. The hub checks every 30 s, so this happens 10 to 10.5 minutes after the last report. The hub sends device_offline only if it had heard the sensor since it started or since it was added. A sensor that was wet when it went silent stays critical as a leak and gets no device_offline.''',
    ])
    ctx.note("s7-v02: A.2 offline timeout covers LoRa sensors, the 30 s check and when device_offline is sent")
    _set_row(ctx, 42, 2, "Valve disconnect timeout", [
        None, None,
        '''After the valve's BLE link drops it reads "warning" ("Valve disconnected") for 3 minutes, then "critical" ("Valve offline"), and the hub sends device_offline (only for a valve linked since the hub started or since it was added). A valve whose flood probe was wet when the link dropped reads "critical" at once as a leak ("Leak detected: valve, Valve offline") and gets no device_offline.''',
    ])
    ctx.note("s7-v03: A.2 valve timeout gives the warning grace, the reasons and device_offline")
    _set_row(ctx, 42, 3, "Boot sync window", [
        "First-snapshot wait", "3 minutes",
        '''After a restart the hub holds its first complete snapshot (data.reason "boot") until every device has been heard, or for 3 minutes; 2.5 minutes after a provision that adds devices. A "fast" snapshot can come earlier, as soon as the valve link is ready, and at the latest about 150 s after the hub connects.''',
    ])
    ctx.note("s7-51: A.2 'Boot sync window 2 minutes' becomes 'First-snapshot wait 3 minutes' (gates the boot snapshot, not offline marking); the pre-2.1.3 value is not mentioned (canon C22)")
    _add_rows(t42, ctx.row(42, 3, "Boot sync window"), b42, [[
        "Not-heard-yet grace", "10 minutes (sensor), 3 minutes (valve)",
        '''A device not heard since the hub started reads rating "critical" and connected=false in its own entry, but is left out of system_health.rating until this time has passed (reason "Syncing - waiting for N device(s)", or ending with "syncing N device(s)" after other causes). After a provision the new devices get 10 minutes (sensor) or 2.5 minutes (valve). After that they count as offline. No device_offline is sent for a device never heard since the hub started or since it was added. Up to 2.1.3, every provision or decommission also returned the remaining devices to this state until each was heard again; since 2.1.4 the remaining devices keep their state.''',
    ]])
    ctx.note("s7-52: new A.2 row 'Not-heard-yet grace' (10 min sensor, 3 min valve; provision 10 / 2.5 min; 2.1.3 reset the other devices on every provision or decommission)")
    _set_row(ctx, 42, 4, "Battery warning threshold", [
        "Sensor battery warning", None,
        '''A sensor battery at or below 20% gives warning ("N sensor(s) battery low"). A sensor battery never makes a device critical.''',
    ])
    ctx.note("s7-53: A.2 battery warning row is sensor-only")
    _set_row(ctx, 42, 5, "Battery good threshold", [
        "Sensor battery good", None,
        '''A sensor battery at or below 35% (and above 20%) gives good (not excellent).''',
    ])
    ctx.note("s7-54: A.2 battery good row is sensor-only")
    _add_rows(t42, ctx.row(42, 5, "Battery good threshold"), b42, [
        ["Valve battery critical", "10%",
         '''Since 2.1.4, a valve battery at or below 10% gives critical ("Valve battery critical"). No health event is sent; the hub sends a snapshot within seconds, and valve_open is refused. 2.1.3 rates the valve with the sensor bands (20% warning, 35% good) and never makes it critical.'''],
        ["Valve battery low", "20%",
         '''Since 2.1.4, a valve battery from 11% to 20% gives warning ("Valve battery low"). Above 20% the valve battery reads excellent: the valve has no good band. An unknown valve battery (null) is not rated.'''],
    ])
    ctx.note("s7-55: new A.2 rows for the 2.1.4 valve battery bands (critical at 10 % or less, warning 11-20 %, no good band; 2.1.3 used the sensor bands)")
    _set_row(ctx, 42, 6, "RSSI warning threshold", [
        None, None,
        '''A sensor RSSI at or below -90 dBm gives warning ("N sensor(s) signal weak"). The valve's signal is not rated.''',
    ])
    ctx.note("s7-56: A.2 RSSI warning row is sensor-only")
    _set_row(ctx, 42, 7, "RSSI good threshold", [
        None, None,
        '''A sensor RSSI from -89 to -80 dBm gives good (not excellent). The valve's signal is not rated.''',
    ])
    ctx.note("s7-v01: A.2 RSSI good row is sensor-only, band -89 to -80 dBm")
    _set_row(ctx, 42, 8, "Alert debounce", [
        None, None,
        '''At most one device_offline or device_recovered per device per 60 s. Since 2.1.4 a change that falls inside those 60 s is sent when they end (within about 90 s), so prev_rating can equal rating; 2.1.3 dropped it.''',
    ])
    ctx.note("s7-57: A.2 alert debounce names the health events and the 2.1.4 trailing edge (2.1.3 dropped the alert)")
    _add_rows(t42, ctx.row(42, 9, "Health evaluation cycle"), b42, [
        ["Leak", "Immediate",
         '''A wet leak sensor or a wet valve flood probe is critical: system_health.reason "Leak detected: " and the sensor's label (its sensor_id when it has no label, "valve" for the valve flood probe), or "N leaks detected". No health event is sent for it.'''],
        ["Leak interlock latched", "While latched",
         '''While the hub holds a leak incident latched, system_health.rating is at least warning ("Leak interlock latched"), also after every source is dry, until RMLEAK is cleared: by the automatic clear 10 s after every source is dry (30 s up to 2.1.3), by leak_reset or by an override.'''],
        ["Rating-change snapshot", "Within about 5 s",
         '''Since 2.1.4, a change of system_health.rating to, from or within warning and critical, and every change of the valve's battery band, sends a snapshot (data.reason "event"). A change between excellent and good waits for the next heartbeat. 2.1.3 showed a change that raised no health event only in the next heartbeat or event snapshot.'''],
        ["Health events", "Reachability only",
         '''device_offline when a device heard since the hub started, or since it was added, loses its link while dry (a sensor that was wet when it went silent stays critical as a leak and gets none); device_recovered when it is heard again. Since 2.1.4 its rating can still be "critical", for example when it came back wet; a 2.1.3 hub sends no device_recovered for a device that came back wet. A leak, a battery at any level and a weak signal never raise a health event.'''],
    ])
    ctx.note("s7-58: new A.2 rows 'Leak', 'Leak interlock latched', 'Rating-change snapshot' (since 2.1.4) and 'Health events' (reachability only, canon C6)")
    ctx.note("review R7-05: A.2 'Sensor offline timeout' and 'Health events': a sensor wet when it went silent gets no device_offline; device_recovered while still critical is since 2.1.4 (2.1.3 sends none for a device back wet)")

    # A.3 Rules engine
    t43 = ctx.T(43, "Rule")
    b43 = ctx.row(43, 1, "Auto-close cooldown")
    _set_row(ctx, 43, 2, "RMLEAK auto-clear timeout", [
        None, "10 seconds (30 seconds up to 2.1.3)",
        '''When every provisioned leak source, including sources whose trigger-mask bit is clear, has been dry for 10 s, the hub clears RMLEAK and sends rmleak_auto_cleared (in practice 10 to 12 s after the last dry report; 30 to 60 s up to 2.1.3). The valve stays closed. rmleak_auto_cleared.clear_after_seconds carries the value in force.''',
    ])
    ctx.note("tm-40: A.3 RMLEAK auto-clear is 10 s since 2.1.4 (30 s up to 2.1.3), covers masked sources, and leaves the valve closed")
    _set_row(ctx, 43, 3, "RMLEAK grace period", [
        None, None,
        '''Internal. For 5 s after the hub locks the valve, or after the valve reconnects, the hub does not read a cleared RMLEAK as a button override. No app action.''',
    ])
    ctx.note("s7-60: A.3 RMLEAK grace period marked internal, no app action")
    _set_row(ctx, 43, 4, "Override window duration", [
        None, None,
        '''Duration of the water access override window. An override_enable sent while a window is active starts it again at 24 hours, with a new water_access_override_enabled (remaining_s 86400).''',
    ])
    ctx.note("s7-v07: A.3 override duration says a second override_enable restarts the window at 24 h")
    _set_row(ctx, 43, 5, "Override blocked cooldown", [
        None, None,
        '''At most one auto_close_blocked_override per 60 s for the whole hub, whichever sensor is wet.''',
    ])
    ctx.note("s7-61: A.3 blocked-override cooldown is hub-wide")
    _add_rows(t43, ctx.row(43, 5, "Override blocked cooldown"), b43, [
        ["Default rules", "Auto-close on, trigger mask 7",
         '''Used on a new hub, after decommission "all", and, since 2.1.4, when a decommission or a provision leaves the hub with no device. Up to 2.1.3 such a hub kept its last rules: until its next restart when a decommission emptied it, and across restarts when a provision emptied it (it stays provisioned).'''],
        ["Trigger mask check", "Each leak report",
         '''The mask is checked when a leak report arrives from a provisioned source. A valve reconnect outside an override window, an override_cancel and the end of an override window close the valve for any wet provisioned source while auto-close is on, without checking the mask (section 4.8.1).'''],
        ["Override reconnect wait", "10 seconds",
         '''override_enable waits up to 10 s for the valve to link. If it does not, the command fails with "The valve isn't responding. Check its power and connection, then try again." Its cmd_ack can therefore take about 10 s.'''],
        ["Override cancel or expiry during a leak", "Immediate",
         '''If a leak source is still wet and auto-close is on when the window is cancelled or ends, the hub asserts RMLEAK and closes the valve at once, whatever the trigger mask. override_cancel sends auto_close_reenabled; the end of the window sends water_access_override_expired, whose auto_close_resumed is true whenever a leak source is still wet, also when auto-close is off. Neither sends auto_close.'''],
        ["Hub without a valve", "No auto_close",
         '''A leak latches the incident and sends leak_detected, but no auto_close (since 2.1.4; 2.1.3 sent auto_close with rmleak_asserted=false).'''],
    ])
    ctx.note("global fix INT-09/XC-03, completeness-03: A.3 'Default rules' gives the 2.1.3 rules after an emptying decommission and provision; A.2 'Valve disconnect timeout' has no device_offline for a valve that dropped while wet")
    ctx.note("tm-41: new A.3 rows 'Default rules' (with the 2.1.4 empty-hub reset) and 'Trigger mask check'")
    ctx.note("s7-62: new A.3 rows 'Override reconnect wait', 'Override cancel or expiry during a leak' (auto_close_resumed as the firmware sets it) and 'Hub without a valve'")

    # A.4 Offline buffer
    t44 = ctx.T(44, "Limit")
    b44 = ctx.row(44, 1, "Maximum buffered events")
    _set_row(ctx, 44, 1, "Maximum buffered events", [
        None, None,
        '''Events kept in flash while the hub cannot send; they survive a restart. When full, the oldest is overwritten. They are sent oldest first at the next connect, before the lifecycle, and, since 2.1.4, every 10 s while connected if the hub could not send some (2.1.3 sent them only at the next connect). Since 2.1.4 decommission "all" clears them; a 2.1.3 hub keeps them through the reset (section 5.5).''',
    ])
    ctx.note("s7-63: A.4 buffered events: flash, survive a restart, oldest overwritten, replay order, 10 s in-session replay since 2.1.4, cleared by decommission all")
    ctx.note("integrator: A.4 'cleared by decommission all' is since 2.1.4, as 5.5 and 5.3 say (master app_iothub.c has no offline_buffer_clear; HEAD app_iothub.c:3635)")
    _set_row(ctx, 44, 2, "Maximum event size", [
        None, None,
        '''Since 2.1.4 an event longer than 512 bytes is not kept; it is lost if it cannot be sent at once. A 2.1.3 hub cut such an event to 512 bytes and later sent it as invalid JSON: the backend logs a message it cannot parse and moves on. Current event types normally fit.''',
    ])
    ctx.note("s7-64: A.4 event size: not kept since 2.1.4; 2.1.3 truncated to invalid JSON, which the backend logs (canon C13)")
    _add_rows(t44, ctx.row(44, 2, "Maximum event size"), b44, [
        ["What is kept", "Events only",
         '''Leak, valve, rules and health events, and cmd_ack. Snapshots, lifecycle messages and twin reports are not kept; the hub builds them again after it reconnects.'''],
        ["Events before the first clock sync", "Kept",
         '''Since 2.1.4, kept and sent after the first connect with ts corrected from the hub's uptime; one left over from a restart before the sync is dropped. 2.1.3 discarded every event raised before the clock sync.'''],
        ["Duplicates", "Possible",
         '''Since 2.1.4 an event whose publish meets a dropping connection is kept and sent again after the reconnect, so it can arrive twice, byte-identical, possibly after newer messages (APP-NF-013).'''],
        ["Send queue", "24 messages",
         '''Since 2.1.4 the hub queues up to 24 messages in RAM for sending. Health events and twin reports take at most 8 of the places; beyond that they wait on the hub and go later. If all 24 places are full, a further leak, valve or rules event is lost, and the next snapshot carries the state. 2.1.3 sent each message directly; one the MQTT client refused was lost.'''],
    ])
    ctx.note("s7-65: new A.4 rows 'What is kept', 'Events before the first clock sync', 'Duplicates' and 'Send queue' (health events and twin reports wait rather than drop)")

    # A.5 Message timing (s7-66)
    head = insert_after(t44, "A.5 Message Timing", like=ctx.P(531, "A.4 Offline Buffer Limits"))
    t45 = insert_table_after(head, t44, ["Item", "Value", "Description"], [
        ["Heartbeat snapshot", "300 s by default",
         '''Set by the twin desired snapshot_interval_s (60 to 3600 s); kept across restarts and shown in twin reported. Counted from the last snapshot of any kind. data.reason "heartbeat".'''],
        ["Snapshot after an event or a command", "About 0.3 to 5 s",
         '''Every event and every command that succeeded is followed by a snapshot (data.reason "event"). Snapshots are at least 5 s apart, so a burst of events gives one snapshot.'''],
        ["First snapshot after a connect", "Up to about 150 s",
         '''After a reconnect without a restart: the "boot" snapshot, within seconds of the lifecycle. After a restart: a "fast" snapshot as soon as the valve link is ready, at the latest about 150 s after the connect, then a "boot" snapshot once every device has been heard or 3 minutes after the restart.'''],
        ["After a provision that adds devices", "5 minutes",
         '''A "boot" snapshot once every new device has been heard or after 2.5 minutes; a snapshot every 30 s and on each sensor report for 5 minutes (data.reason "event"); and, within 6 minutes, one each time a new device is heard for the first time (data.reason "commission").'''],
        ["Twin reported", "At connect and after changes",
         '''Sent at every connect, after every provision and decommission, after rules_config and set_hub_name, and after every desired-property change. Since 2.1.4 it follows the command's cmd_ack, normally within a second. 2.1.3 sent it before the cmd_ack for provision, rules_config and set_hub_name, and sent none after a decommission.'''],
        ["Replay of kept events", "At connect; also every 10 s since 2.1.4",
         '''Kept events go out before the lifecycle at each connect and, since 2.1.4, every 10 s while connected if the hub could not send some.'''],
        ["MQTT keep-alive", "60 s",
         '''After a power cut or a silent network loss, IoT Hub reports the hub disconnected only after about 90 s.'''],
        ["Token renewal reconnect", "About every 18 hours",
         '''The hub closes and reopens its connection to renew its 24-hour SAS token, and sends a lifecycle after it. The backend ignores a disconnect followed by a reconnect within 60 s.'''],
        ["Decommission all", "About 3 to 30 s",
         '''The cmd_ack, a twin report and a final snapshot (data.reason "decommission") go out, then the hub restarts, normally 3 to 5 s after the ack and at most about 30 s after it. 2.1.3 sends no twin report before the restart.'''],
    ])
    ctx.note("s7-66: new Appendix A.5 Message Timing (heartbeat, event snapshot, first snapshot, provision pulse, twin reports with the 2.1.3 order, replay, keep-alive, token renewal per canon C14, decommission all)")
    _restripe(t45)


def _format(ctx):
    """Table formatting after every row is in place (review R7-06, R7-09)."""
    for k, exp in ((37, "Scenario"), (38, "Test ID"), (39, "Test ID"), (40, "App Screen/Feature"),
                   (42, "Threshold"), (43, "Rule"), (44, "Limit")):
        _restripe(ctx.T(k, exp))
    ctx.note("review R7-06: T37-T40, T42-T44 and the new A.5 table keep V4's zebra shading (odd body rows F5F5F5, even rows plain), also across inserted rows")
    for k, r, exp, ref_exp in ((38, 12, "TC-012", "TC-001"), (39, 7, "TC-N07", "TC-N01"),
                               (39, 8, "TC-N08", "TC-N01"), (40, 12, "Remote Override", "Dashboard")):
        _match_size(ctx.row(k, r, exp), ctx.row(k, 1, ref_exp))
    ctx.note("review R7-09: the new T40 row copies R1's 9 pt formatting; V4 rows T38 R12, T39 R7-R8 and T40 R12, which had no font size (11 pt default), now use the table's 9 pt")


def apply(ctx):
    _edge_cases(ctx)
    _tests(ctx)
    _negative_tests(ctx)
    _trace(ctx)
    _appendix(ctx)
    _format(ctx)
